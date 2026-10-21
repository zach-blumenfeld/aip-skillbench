"""Extract the case background (PDF text) and profile every sheet of the workbook.

stdin: {"state": {"data_path", "background_path"}, ...}
stdout: {"background_text": str, "workbook_summary": str}
"""
import json
import logging
import sys

import openpyxl
from pypdf import PdfReader

logging.getLogger("pypdf").setLevel(logging.ERROR)  # pypdf warns on benign xref damage
MAX_TEXT = 30000


def pdf_text(path):
    try:
        reader = PdfReader(path)
    except Exception as e:  # unreadable PDF: say so, let the agent fall back
        return f"[could not read PDF {path}: {e}]"
    parts = []
    for i, page in enumerate(reader.pages, 1):
        try:
            parts.append(f"--- page {i} ---\n{page.extract_text() or ''}")
        except Exception as e:
            parts.append(f"--- page {i} --- [extract failed: {e}]")
    text = "\n".join(parts)
    if not text.replace("-", "").replace("page", "").strip():
        text += "\n[no extractable text: PDF may be scanned; see references/pdf-reference.md OCR notes]"
    return text[:MAX_TEXT]


def fmt(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def profile_sheet(ws_vals, ws_form):
    rows = [list(r) for r in ws_vals.iter_rows(values_only=True)]
    nonempty = [(i + 1, r) for i, r in enumerate(rows) if any(c is not None for c in r)]
    formulas = sum(
        1 for r in ws_form.iter_rows(values_only=True) for c in r
        if isinstance(c, str) and c.startswith("=")
    )
    out = [f"## Sheet '{ws_vals.title}': {len(rows)} rows x {max((len(r) for r in rows), default=0)} cols, "
           f"{len(nonempty)} non-empty rows, {formulas} formula cells"]
    # Header row = first row with >=2 text cells whose next non-empty row is mostly numeric.
    header = None
    for k, (rn, r) in enumerate(nonempty[:-1]):
        texts = [c for c in r if isinstance(c, str)]
        nxt = [c for c in nonempty[k + 1][1] if c is not None]
        if len(texts) >= 2 and nxt and sum(isinstance(c, (int, float)) for c in nxt) >= max(2, len(nxt) // 2):
            header = (k, rn, r)
            break
    pre = nonempty[: header[0]] if header else nonempty[:25]
    for rn, r in pre[:25]:
        cells = [f"{openpyxl.utils.get_column_letter(j + 1)}{rn}={fmt(c)}" for j, c in enumerate(r) if c is not None]
        out.append("  " + " | ".join(cells))
    if header:
        k, rn, r = header
        cols = [(openpyxl.utils.get_column_letter(j + 1), c) for j, c in enumerate(r) if c is not None]
        data = nonempty[k + 1:]
        out.append(f"  HEADER at Excel row {rn}: " + ", ".join(f"{col}='{c}'" for col, c in cols))
        out.append(f"  {len(data)} data rows below header (Excel rows {data[0][0]}..{data[-1][0]})")
        idx = [openpyxl.utils.column_index_from_string(col) - 1 for col, _ in cols]
        for rn2, r2 in data[:5] + ([("...", None)] if len(data) > 8 else []) + data[-3:]:
            if r2 is None:
                out.append("  ...")
                continue
            out.append(f"  row {rn2}: " + ", ".join(fmt(r2[j]) if j < len(r2) else "" for j in idx))
        gaps = [i for i in range(data[0][0], data[-1][0] + 1) if not any(c is not None for c in rows[i - 1])]
        if gaps:
            out.append(f"  NOTE: {len(gaps)} blank rows inside the table (Excel rows {gaps[:10]}); a record may have been moved elsewhere")
        partial = [rn2 for rn2, r2 in data if any((r2[j] if j < len(r2) else None) is None for j in idx)]
        if partial:
            out.append(f"  NOTE: {len(partial)} table rows have blank cells in header columns (first at row {partial[0]})")
        extra = [rn2 for rn2, r2 in data if any(c is not None for j, c in enumerate(r2) if j not in idx)]
        for rn2 in extra[:10]:
            r2 = rows[rn2 - 1]
            cells = [f"{openpyxl.utils.get_column_letter(j + 1)}={fmt(c)}" for j, c in enumerate(r2) if c is not None and j not in idx]
            out.append(f"  NOTE: row {rn2} has values outside the header columns: " + ", ".join(cells))
        if len(extra) > 10:
            out.append(f"  NOTE: {len(extra) - 10} more rows have values outside the header columns")
    return "\n".join(out)


def main():
    state = json.load(sys.stdin)["state"]
    wb_vals = openpyxl.load_workbook(state["data_path"], data_only=True, read_only=True)
    wb_form = openpyxl.load_workbook(state["data_path"], read_only=True)
    summary = [f"# Workbook {state['data_path']}: sheets {wb_vals.sheetnames}"]
    for name in wb_vals.sheetnames:
        summary.append(profile_sheet(wb_vals[name], wb_form[name]))
    json.dump({"background_text": pdf_text(state["background_path"]),
               "workbook_summary": "\n".join(summary)}, sys.stdout)


if __name__ == "__main__":
    main()
