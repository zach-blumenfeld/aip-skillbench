#!/usr/bin/env python3
"""Write lookup / statistics / fold-change FORMULAS into the Task sheet, recalc, verify.

stdin: {"state": {"workbook_path", "output_path", "detected_layout", "layout_overrides",
                  "missing_value_policy"}, ...}
stdout: {"verdict": "verified" | "needs_attention", "fill_report": {...}, "fill_issues": [...]}

Formulas only, never hard-coded numbers. Cell styles are left untouched (openpyxl keeps
fills/fonts when only .value changes). After saving, LibreOffice recalculates the file
(scripts/recalc.py) so cached values exist, and every target value is checked against
an independent Python computation from the Data sheet.
"""
import json
import math
import os
import re
import shutil
import statistics
import sys
import tempfile

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from recalc import recalc

TOL = 1e-6


def merge(base, over):
    out = dict(base)
    for k, v in (over or {}).items():
        if isinstance(v, dict) and isinstance(out.get(k), dict):
            out[k] = {**out[k], **v}
        else:
            out[k] = v
    return out


def qs(name):
    return name if re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", name) else "'" + name.replace("'", "''") + "'"


def ref_list(cells):
    """Compact reference text for cells in one row/column: 'C11:G11' when contiguous."""
    if not cells:
        return None
    parsed = [re.fullmatch(r"([A-Z]+)(\d+)", c).groups() for c in cells]
    cols = [column_index_from_string(p[0]) for p in parsed]
    rows = [int(p[1]) for p in parsed]
    if len(set(rows)) == 1 and cols == list(range(cols[0], cols[0] + len(cols))):
        return f"{cells[0]}:{cells[-1]}" if len(cells) > 1 else cells[0]
    if len(set(cols)) == 1 and rows == list(range(rows[0], rows[0] + len(rows))):
        return f"{cells[0]}:{cells[-1]}" if len(cells) > 1 else cells[0]
    return ",".join(cells)


def is_blank(v):
    return v is None or (isinstance(v, str) and v.strip() == "")


def style_sig(cell):
    f, fo = cell.fill, cell.font
    fill = str(f.fgColor.rgb)[-6:] if f is not None and f.fill_type and f.fgColor is not None and f.fgColor.type == "rgb" else None
    color = str(fo.color.rgb)[-6:] if fo is not None and fo.color is not None and fo.color.type == "rgb" else None
    return (fill, fo.name if fo else None, float(fo.sz) if fo and fo.sz else None,
            bool(fo.b) if fo else False, bool(fo.i) if fo else False, color)


def snapshot_styles(path):
    wb = load_workbook(path)
    snap = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                sig = style_sig(c)
                if any(sig[i] for i in (0, 5)) or sig[3] or sig[4]:
                    snap[f"{ws.title}!{c.coordinate}"] = sig
    return snap


def run(state):
    issues = []
    L = merge(state["detected_layout"], state.get("layout_overrides") or {})
    policy = state.get("missing_value_policy") or "exclude"
    if policy not in ("exclude", "zero"):
        raise SystemExit(f"missing_value_policy must be 'exclude' or 'zero', got {policy!r}")
    src, dst = state["workbook_path"], state["output_path"]

    backup = os.path.join(tempfile.mkdtemp(prefix="protein_xlsx_"), "original_" + os.path.basename(src))
    shutil.copy2(src, backup)

    wb = load_workbook(src)
    ws, ds = wb[L["task_sheet"]], wb[L["data_sheet"]]
    D, K = L["data"], L["lookup"]
    dsn = qs(L["data_sheet"])
    hr, fr, lr = D["header_row"], D["first_row"], D["last_row"]
    idc, gc, lc = D["id_col"], D.get("gene_col"), D["last_col"]
    hdr_rng = f"{dsn}!$A${hr}:${lc}${hr}"
    body_rng = f"{dsn}!$A${fr}:${lc}${lr}"
    id_rng = f"{dsn}!${idc}${fr}:${idc}${lr}"
    gene_rng = f"{dsn}!${gc}${fr}:${gc}${lr}" if gc else None

    headers = {}
    for c in ds[hr]:
        if c.value is not None:
            headers.setdefault(str(c.value), c.column)

    sample_cols, target_rows = K["sample_cols"], K["target_rows"]
    groups = L.get("groups") or {}
    sres, pres = L["sample_resolution"], L["protein_resolution"]
    t_id, t_gene = K["id_col"], K.get("gene_col")
    written = {"lookup": 0, "stats": 0, "fold_change": 0}
    expected = {}
    examples = {}

    # ---- step 1: two-way lookups -------------------------------------------------
    for r in target_rows:
        p = pres[str(r)]
        if p["by"] == "gene" and gene_rng and t_gene:
            rm = f"MATCH(${t_gene}{r},{gene_rng},0)"
        else:
            rm = f"MATCH(${t_id}{r},{id_rng},0)"
            if p["by"] is None:
                issues.append(f"Protein {p['id']} (row {r}) is not in the Data sheet; its lookups will show #N/A.")
        for col in sample_cols:
            s = sres[col]
            head = f"{col}${K['header_row']}"
            if s["by"] == "contains":
                cm = f'MATCH("*"&{head}&"*",{hdr_rng},0)'
            elif s["by"] == "trim":
                cm = f"MATCH(TRIM({head}),{hdr_rng},0)"
            else:
                cm = f"MATCH({head},{hdr_rng},0)"
                if s["by"] is None and p["by"] is not None:
                    issues.append(f"Sample '{s['name']}' (column {col}) has no unique Data header; its lookups will show #N/A.")
            idx = f"INDEX({body_rng},{rm},{cm})"
            ws[f"{col}{r}"] = f'=IF({idx}="","",{idx})' if policy == "exclude" else f"={idx}"
            written["lookup"] += 1
            if p["data_row"] is not None and s["data_header"] is not None:
                v = ds.cell(p["data_row"], headers[s["data_header"]]).value
                if is_blank(v):
                    v = "" if policy == "exclude" else 0
                expected[f"{col}{r}"] = v
    examples["lookup"] = ws[f"{sample_cols[0]}{target_rows[0]}"].value

    # ---- step 2: group statistics --------------------------------------------------
    S = L["stats_cells"]
    ctrl = [c for c in sample_cols if groups.get(c) == "Control"]
    trt = [c for c in sample_cols if groups.get(c) == "Treated"]
    if not ctrl or not trt:
        issues.append(f"Group assignment incomplete: {len(ctrl)} Control, {len(trt)} Treated columns.")
    for i, r in enumerate(target_rows):
        for g, cols in (("control", ctrl), ("treated", trt)):
            rng = ref_list([f"{c}{r}" for c in cols])
            if rng is None:
                continue
            mcell, scell = S[f"{g}_mean"][i], S[f"{g}_sd"][i]
            ws[mcell] = f'=IF(COUNT({rng})=0,"",AVERAGE({rng}))'
            ws[scell] = f'=IF(COUNT({rng})<2,"",STDEV({rng}))'
            written["stats"] += 2
            vals = [expected.get(f"{c}{r}") for c in cols]
            nums = [v for v in vals if isinstance(v, (int, float)) and not isinstance(v, bool)]
            if all(f"{c}{r}" in expected for c in cols):
                expected[mcell] = statistics.fmean(nums) if nums else ""
                expected[scell] = statistics.stdev(nums) if len(nums) >= 2 else ""
    examples["control_mean"] = ws[S["control_mean"][0]].value
    examples["control_sd"] = ws[S["control_sd"][0]].value

    # ---- step 3: fold change ----------------------------------------------------------
    F = L["fc_cells"]
    for i in range(len(target_rows)):
        cm, tm = S["control_mean"][i], S["treated_mean"][i]
        lg, fc = F["log2_fc"][i], F["fold_change"][i]
        ws[lg] = f'=IF(COUNT({tm},{cm})<2,"",{tm}-{cm})'
        ws[fc] = f'=IF({lg}="","",2^{lg})'
        written["fold_change"] += 2
        if cm in expected and tm in expected:
            a, b = expected[cm], expected[tm]
            if isinstance(a, float) and isinstance(b, float):
                expected[lg] = b - a
                expected[fc] = 2 ** (b - a)
            else:
                expected[lg] = expected[fc] = ""
    examples["log2_fc"] = ws[F["log2_fc"][0]].value
    examples["fold_change"] = ws[F["fold_change"][0]].value

    os.makedirs(os.path.dirname(os.path.abspath(dst)), exist_ok=True)
    wb.save(dst)
    unrecalculated = os.path.join(os.path.dirname(backup), "formulas_only_" + os.path.basename(dst))
    shutil.copy2(dst, unrecalculated)

    # every yellow Task cell should now hold something
    wb2 = load_workbook(dst)
    t2 = wb2[L["task_sheet"]]
    empty_yellow = [c.coordinate for row in t2.iter_rows() for c in row
                    if style_sig(c)[0] and style_sig(c)[0].upper() == "FFFF00" and c.value is None]
    if empty_yellow:
        issues.append(f"Yellow input cells still empty: {empty_yellow[:20]}")

    # ---- recalc + verify ----------------------------------------------------------------
    pre = snapshot_styles(unrecalculated)
    rc = recalc(dst)
    mismatches = []
    style_drift = []
    if "error" in rc:
        issues.append(f"Recalculation failed: {rc['error']}. The file at output_path holds the formulas but no cached values.")
    else:
        if rc.get("total_errors"):
            issues.append(f"Formula errors after recalculation: {rc['error_summary']}")
        vals = load_workbook(dst, data_only=True)[L["task_sheet"]]
        forms = load_workbook(dst)[L["task_sheet"]]
        for addr, exp in expected.items():
            got = vals[addr].value
            if not (isinstance(forms[addr].value, str) and forms[addr].value.startswith("=")):
                mismatches.append({"cell": addr, "problem": "formula lost after recalculation"})
                continue
            if exp == "" or exp is None:
                ok = is_blank(got)
            elif isinstance(exp, (int, float)):
                ok = isinstance(got, (int, float)) and math.isclose(got, exp, rel_tol=TOL, abs_tol=TOL)
            else:
                ok = got == exp
            if not ok:
                mismatches.append({"cell": addr, "expected": exp, "got": got})
        if mismatches:
            issues.append(f"{len(mismatches)} recalculated values differ from the independent Python check.")
        post = snapshot_styles(dst)
        for k in sorted(set(pre) | set(post)):
            if pre.get(k) != post.get(k):
                style_drift.append({"cell": k, "before": pre.get(k), "after": post.get(k)})
        if style_drift:
            issues.append(f"LibreOffice changed fills/fonts on {len(style_drift)} cells; "
                          f"the untouched-style formulas-only copy is at {unrecalculated}.")

    blanks = sorted(a for a, v in expected.items() if v == "")
    report = {
        "output_path": dst,
        "original_backup": backup,
        "formulas_only_copy": unrecalculated,
        "missing_value_policy": policy,
        "cells_written": written,
        "formula_examples": examples,
        "recalc": rc,
        "cells_checked": len(expected),
        "blank_by_design": blanks,
        "mismatches": mismatches[:20],
        "style_drift": style_drift[:20],
    }
    verdict = "verified" if not issues else "needs_attention"
    return {"verdict": verdict, "fill_report": report, "fill_issues": issues}


def main():
    payload = json.load(sys.stdin)
    print(json.dumps(run(payload.get("state", payload)), default=str))


if __name__ == "__main__":
    main()
