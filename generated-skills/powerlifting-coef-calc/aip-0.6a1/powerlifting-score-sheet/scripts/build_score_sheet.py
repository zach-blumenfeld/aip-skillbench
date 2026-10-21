#!/usr/bin/env python3
"""Build a powerlifting score sheet (Dots, Wilks or IPF GL) with live Excel formulas.

AIP mode (no arguments): reads {"state": {workbook_path, score_system, decimals},
"assets": {"coefficients": <json text>}} on stdin, writes the score sheet into the
workbook in place, recalculates it with LibreOffice (or gnumeric), verifies every
cached value against an independent Python computation, and prints one JSON object.

Verify-only mode (for re-checking after a manual recalc):
    python build_score_sheet.py --verify <xlsx> --system dots --decimals 3
"""

import json
import math
import os
import shutil
import subprocess
import sys
import tempfile
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
BASE_COLUMNS = ["Name", "Sex", "BodyweightKg", "Best3SquatKg", "Best3BenchKg", "Best3DeadliftKg"]
EXTRA_COLUMNS = {"ipf_goodlift": ["Event", "Equipment"]}
LIFTS = ["Best3SquatKg", "Best3BenchKg", "Best3DeadliftKg"]
TOTAL_COLUMN = "TotalKg"
EXCEL_ERRORS = ["#VALUE!", "#DIV/0!", "#REF!", "#NAME?", "#NULL!", "#NUM!", "#N/A"]


def fail(msg):
    print(msg, file=sys.stderr)
    sys.exit(1)


def norm(h):
    return "".join(str(h).split()).lower() if h is not None else ""


def num(v):
    """Excel N(): numbers pass through, everything else (blank, text) is 0."""
    if isinstance(v, bool):
        return float(v)
    if isinstance(v, (int, float)):
        return float(v)
    return 0.0


def xround(v, d):
    return float(Decimal(repr(v)).quantize(Decimal(1).scaleb(-d), rounding=ROUND_HALF_UP))


def lit(v):
    s = format(Decimal(repr(abs(float(v)))), "f")
    return (s.rstrip("0").rstrip(".") if "." in s else s) or "0"


def poly_xl(coefs, x):
    n = len(coefs) - 1
    out = ""
    for i, c in enumerate(coefs):
        p = n - i
        term = lit(c) + (f"*{x}^{p}" if p > 1 else f"*{x}" if p == 1 else "")
        out += ("-" if c < 0 else ("+" if out else "")) + term
    return f"({out})"


def poly_py(coefs, x):
    acc = 0.0
    for c in coefs:
        acc = acc * x + c
    return acc


def sex_key(sex):
    return "F" if str(sex or "").strip().upper() == "F" else "M"


def equip_group(eq, groups):
    for g, members in groups.items():
        if str(eq or "").strip().lower() in [m.lower() for m in members]:
            return g
    return None


# ---------- formulas ----------

def total_formula(cols, r, d):
    cells = ",".join(f"{cols[c]}{r}" for c in LIFTS)
    return f"=ROUND(SUM({cells}),{d})"


def score_formula(system, coef, cols, r, d):
    sx, bw, tot = f"{cols['Sex']}{r}", f"{cols['BodyweightKg']}{r}", f"{cols[TOTAL_COLUMN]}{r}"
    if system in ("dots", "wilks"):
        m, w = coef[system]["men"], coef[system]["women"]
        xm = f"MAX({lit(m['min'])},MIN({lit(m['max'])},{bw}))"
        xw = f"MAX({lit(w['min'])},MIN({lit(w['max'])},{bw}))"
        denom = f'IF({sx}="F",{poly_xl(w["poly_high_to_low"], xw)},{poly_xl(m["poly_high_to_low"], xm)})'
        return f"=IF(OR(N({bw})<=0,N({tot})<=0),0,ROUND({tot}*500/{denom},{d}))"
    g = coef["ipf_goodlift"]
    ev, eq = f"{cols['Event']}{r}", f"{cols['Equipment']}{r}"
    grp_test = {k: "OR(" + ",".join(f'{eq}="{m}"' for m in v) + ")" for k, v in g["equipment_groups"].items()}
    expr = "0"
    for key, (A, B, C) in reversed(list(g["params"].items())):
        e, s, q = key.split("|")
        sex_test = f'{sx}="F"' if s == "F" else f'{sx}<>"F"'
        pts = f"ROUND({tot}*MAX(0,100/({lit(A)}-{lit(B)}*EXP(-{lit(C)}*{bw}))),{d})"
        expr = f'IF(AND({ev}="{e}",{sex_test},{grp_test[q]}),{pts},{expr})'
    return f"=IF(OR(N({bw})<35,N({tot})<=0),0,{expr})"


# ---------- independent Python reference ----------

def expected(system, coef, row, d):
    total = xround(sum(num(row[c]) for c in LIFTS), d)
    bw = num(row["BodyweightKg"])
    if system in ("dots", "wilks"):
        if bw <= 0 or total <= 0:
            return total, 0.0
        p = coef[system]["women" if sex_key(row["Sex"]) == "F" else "men"]
        x = min(max(bw, p["min"]), p["max"])
        return total, xround(total * 500 / poly_py(p["poly_high_to_low"], x), d)
    g = coef["ipf_goodlift"]
    if bw < 35 or total <= 0:
        return total, 0.0
    key = f"{str(row['Event'] or '').strip()}|{sex_key(row['Sex'])}|{equip_group(row['Equipment'], g['equipment_groups'])}"
    if key not in g["params"]:
        return total, 0.0
    A, B, C = g["params"][key]
    return total, xround(total * max(0.0, 100 / (A - B * math.exp(-C * bw))), d)


# ---------- workbook inspection ----------

def find_data_sheet(wb, needed):
    order = sorted(wb.worksheets, key=lambda ws: ws.title.strip().lower() != "data")
    for ws in order:
        hdr = {norm(c.value): i for i, c in enumerate(ws[1], start=1) if c.value is not None}
        if all(norm(n) in hdr for n in needed):
            return ws, hdr
    found = {ws.title: [c.value for c in ws[1]] for ws in wb.worksheets}
    fail(f"No sheet has all required columns {needed}. Headers found: {found}")


def read_rows(ws, hdr, needed):
    rows = []
    for r in range(2, ws.max_row + 1):
        vals = {n: ws.cell(row=r, column=hdr[norm(n)]).value for n in needed}
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in vals.values()):
            continue
        rows.append(vals)
    return rows


def coerce_numbers(rows):
    """Numbers stored as text would score 0 under N(); write them as numbers instead."""
    out = []
    for i, row in enumerate(rows, start=2):
        for c in ["BodyweightKg"] + LIFTS:
            v = row[c]
            if isinstance(v, str) and v.strip():
                try:
                    row[c] = float(v.strip().replace(",", "."))
                    out.append(f"row {i} ({row['Name']}): {c} '{v}' was text; copied as the number {row[c]}")
                except ValueError:
                    pass
    return out


def anomalies_for(rows, system, coef):
    out = []
    for i, row in enumerate(rows, start=2):
        who = f"row {i} ({row['Name']})"
        sx = str(row["Sex"] or "").strip()
        if sx not in ("M", "F", "Mx"):
            out.append(f"{who}: Sex '{row['Sex']}' is not M/F/Mx; scored with the men's parameters")
        if num(row["BodyweightKg"]) <= 0:
            out.append(f"{who}: BodyweightKg '{row['BodyweightKg']}' missing or non-numeric; score is 0")
        for c in LIFTS:
            v = row[c]
            if v is None or (isinstance(v, str) and not v.strip()):
                out.append(f"{who}: {c} is blank; counted as 0 in TotalKg")
            elif not isinstance(v, (int, float)):
                out.append(f"{who}: {c} '{v}' is text; SUM ignores it")
            elif v < 0:
                out.append(f"{who}: {c} is negative ({v}), a failed-attempt marker; it lowers TotalKg")
        if system == "ipf_goodlift":
            g = coef["ipf_goodlift"]
            if str(row["Event"] or "").strip() not in ("SBD", "B"):
                out.append(f"{who}: Event '{row['Event']}' has no IPF GL parameters; score is 0")
            if equip_group(row["Equipment"], g["equipment_groups"]) is None:
                out.append(f"{who}: Equipment '{row['Equipment']}' has no IPF GL parameters; score is 0")
    return out


# ---------- recalculation ----------

def has_cached(path, sheet, col_letter, n):
    if n == 0:
        return True
    ws = load_workbook(path, data_only=True)[sheet]
    return all(ws[f"{col_letter}{r}"].value is not None for r in range(2, n + 2))


def recalc(path, sheet, col_letter, n):
    """Return the engine name that left cached values in `path`, or None."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice:
        with tempfile.TemporaryDirectory() as tmp:
            prof = Path(tmp, "profile").as_uri()
            outdir = Path(tmp, "out")
            cmd = [soffice, f"-env:UserInstallation={prof}", "--headless", "--norestore", "--nolockcheck",
                   "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(outdir), str(path)]
            try:
                subprocess.run(cmd, capture_output=True, text=True, timeout=240)
                conv = outdir / Path(path).with_suffix(".xlsx").name
                if conv.exists() and has_cached(conv, sheet, col_letter, n):
                    shutil.copyfile(conv, path)
                    return "libreoffice"
            except subprocess.TimeoutExpired:
                pass
        try:  # the xlsx skill's macro-based recalc as a second attempt
            sys.path.insert(0, str(HERE))
            import recalc as rc
            res = rc.recalc(str(path), 120)
            if "error" not in res and has_cached(path, sheet, col_letter, n):
                return "libreoffice-macro"
        except Exception:
            pass
    ssconvert = shutil.which("ssconvert")
    if ssconvert:
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp, "recalc.xlsx")
            try:
                subprocess.run([ssconvert, "--recalc", str(path), str(out)], capture_output=True, timeout=240)
                if out.exists() and has_cached(out, sheet, col_letter, n):
                    shutil.copyfile(out, path)
                    return "gnumeric"
            except subprocess.TimeoutExpired:
                pass
    return None


# ---------- verification ----------

def verify(path, system, coef, d):
    wb_f = load_workbook(path)
    wb_v = load_workbook(path, data_only=True)
    sheet = coef[system]["score_column"]
    sname = next((ws.title for ws in wb_f.worksheets if ws.title.strip().lower() == sheet.lower()), None)
    if sname is None:
        return {"status": "mismatch", "detail": f"sheet '{sheet}' not found"}
    wf, wv = wb_f[sname], wb_v[sname]
    hdr = {c.value: i for i, c in enumerate(wf[1], start=1) if c.value is not None}
    need = [n for n in BASE_COLUMNS + EXTRA_COLUMNS.get(system, [])]
    missing = [n for n in need + [TOTAL_COLUMN, sheet] if n not in hdr]
    if missing:
        return {"status": "mismatch", "detail": f"columns missing from '{sname}': {missing}"}
    errors, mismatches, uncached, not_formula, max_diff, n = {}, [], 0, 0, 0.0, 0
    tol = 10 ** -d + 1e-9
    for r in range(2, wf.max_row + 1):
        row = {k: wv.cell(row=r, column=hdr[k]).value for k in need}
        if row["Name"] is None:
            continue
        n += 1
        exp_t, exp_s = expected(system, coef, row, d)
        for col, exp in ((TOTAL_COLUMN, exp_t), (sheet, exp_s)):
            f = wf.cell(row=r, column=hdr[col]).value
            v = wv.cell(row=r, column=hdr[col]).value
            ref = f"{sname}!{get_column_letter(hdr[col])}{r}"
            if not (isinstance(f, str) and f.startswith("=")):
                not_formula += 1
            if isinstance(v, str) and any(e in v for e in EXCEL_ERRORS):
                errors.setdefault(v, []).append(ref)
            elif v is None:
                uncached += 1
            else:
                diff = abs(float(v) - exp)
                max_diff = max(max_diff, diff)
                if diff > tol:
                    mismatches.append({"cell": ref, "cached": v, "expected": exp})
    if errors:
        status = "formula_errors"
    elif uncached:
        status = "not_recalculated"
    elif mismatches or not_formula:
        status = "mismatch"
    else:
        status = "verified"
    return {"status": status, "rows_checked": n, "formula_errors": errors, "uncached_cells": uncached,
            "non_formula_cells": not_formula, "mismatches": mismatches[:20], "max_abs_diff": round(max_diff, 9)}


# ---------- main ----------

def build(state, coef):
    path = Path(state["workbook_path"]).expanduser()
    system, d = state["score_system"], int(state["decimals"])
    if system not in ("dots", "wilks", "ipf_goodlift"):
        fail(f"score_system '{system}' is not supported by this script (dots, wilks, ipf_goodlift)")
    if not path.exists():
        fail(f"workbook not found: {path}")
    if d < 0:
        fail("decimals must be >= 0")

    wb = load_workbook(path)
    needed = BASE_COLUMNS + EXTRA_COLUMNS.get(system, [])
    data_ws, hdr = find_data_sheet(wb, needed)
    copied = [n for n in sorted(needed, key=lambda n: hdr[norm(n)])]  # keep the Data sheet's order
    data_names = {norm(n): data_ws.cell(row=1, column=hdr[norm(n)]).value for n in copied}
    rows = read_rows(data_ws, hdr, needed)
    coerced = coerce_numbers(rows)

    score_col = coef[system]["score_column"]
    notes = []
    target = next((ws for ws in wb.worksheets if ws.title.strip().lower() == score_col.lower()), None)
    if target is None:
        target = wb.create_sheet(score_col)
        notes.append(f"created sheet '{score_col}'")
    elif target is data_ws:
        fail(f"the data sheet is named '{target.title}'; refusing to overwrite it")
    elif any(v is not None for r in target.iter_rows(values_only=True) for v in r):
        notes.append(f"cleared existing content of sheet '{target.title}' ({target.max_row} rows)")
        target.delete_rows(1, target.max_row)
    if norm(TOTAL_COLUMN) in hdr:
        notes.append("Data already has a TotalKg column; not copied, recomputed by formula instead")

    headers = [data_names[norm(n)] for n in copied] + [TOTAL_COLUMN, score_col]
    cols = {n: get_column_letter(i) for i, n in enumerate(copied + [TOTAL_COLUMN, score_col], start=1)}
    # explicit coordinates: append() starts after an empty sheet's phantom A1 row
    for c, h in enumerate(headers, start=1):
        target.cell(row=1, column=c, value=h)
    for i, row in enumerate(rows, start=2):
        vals = [row[n] for n in copied] + [total_formula(cols, i, d), score_formula(system, coef, cols, i, d)]
        for c, v in enumerate(vals, start=1):
            target.cell(row=i, column=c, value=v)

    backup = Path(tempfile.mkdtemp(prefix="score-sheet-backup-")) / path.name
    shutil.copyfile(path, backup)
    wb.save(path)

    engine = recalc(path, target.title, cols[score_col], len(rows))
    check = verify(path, system, coef, d)
    check.update({"engine": engine, "backup_path": str(backup), "notes": notes,
                  "anomalies": coerced + anomalies_for(rows, system, coef), "data_sheet": data_ws.title})
    vals = load_workbook(path, data_only=True)[target.title]
    check["sample_rows"] = [[vals.cell(row=r, column=c).value for c in range(1, len(headers) + 1)]
                            for r in range(2, min(len(rows), 3) + 2)]
    return {
        "score_sheet": target.title,
        "copied_columns": headers[:-2],
        "formula_examples": {TOTAL_COLUMN: total_formula(cols, 2, d), score_col: score_formula(system, coef, cols, 2, d)},
        "rows_written": len(rows),
        "recalc_status": check["status"],
        "check_report": check,
    }


def main():
    if len(sys.argv) > 1:
        args = sys.argv[1:]
        if args[0] != "--verify" or len(args) < 2:
            fail(__doc__)
        opts = dict(zip(args[2::2], args[3::2]))
        coef = json.loads((HERE.parent / "assets" / "coefficients.json").read_text())
        print(json.dumps(verify(Path(args[1]), opts.get("--system", "dots"), coef, int(opts.get("--decimals", 3))), indent=2))
        return
    payload = json.load(sys.stdin)
    assets = payload.get("assets") or {}
    raw = assets.get("coefficients")
    coef = json.loads(raw) if isinstance(raw, str) else raw
    if not coef:
        coef = json.loads((HERE.parent / "assets" / "coefficients.json").read_text())
    print(json.dumps(build(payload["state"], coef), default=str))


if __name__ == "__main__":
    main()
