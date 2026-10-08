#!/usr/bin/env python3
"""Score every turn and every game of a dice-scoring case study (ModelOff "Roll The Dice").

stdin:  {"currentState": {"xlsx_path", "dice_rules", ["work_dir"], ["data_sheet"]}, ...}
stdout: {"dice_summary", "data_issues", "turns_csv", "games_csv", "work_dir"}

Loader: finds the header row (Turn / Game / Roll 1..n) on any sheet, reads the main
block, then sweeps the whole sheet for records displaced outside it (a turn moved to
another row/column range) so no turn is silently lost. Every anomaly is reported in
data_issues; nothing is guessed silently.

Scoring follows dice_rules (assets/dice_rules.json after the agent checked it against
the PDF). A game's score is the best total over its turns with each turn in a
different category (when categories_unique_per_game). Needs only openpyxl.
"""
import csv
import itertools
import json
import math
import os
import re
import statistics
import sys
import tempfile
import warnings
from collections import Counter, defaultdict

warnings.filterwarnings("ignore")


def fail(msg):
    print(json.dumps({"error": msg}))
    sys.exit(1)


def as_int(v):
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, int):
        return v
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str) and re.fullmatch(r"\s*-?\d+(\.0+)?\s*", v):
        return int(float(v))
    return None


def resolve(path):
    for c in (path, os.path.join("/root", os.path.basename(path or ""))):
        if c and os.path.isabs(c) and os.path.exists(c):
            return c
    fail(f"xlsx_path not found: {path!r} (pass an absolute path)")


# ---------------------------------------------------------------- loading

def find_header(ws, n_rolls):
    for row in ws.iter_rows(max_row=min(ws.max_row, 200)):
        turn = game = None
        rolls = {}
        for c in row:
            if not isinstance(c.value, str):
                continue
            t = c.value.strip().lower()
            m = re.search(r"roll\s*#?\s*(\d+)", t)
            if m:
                rolls[int(m.group(1))] = c.column
            elif "turn" in t and turn is None:
                turn = c.column
            elif "game" in t and game is None:
                game = c.column
        if turn and len(rolls) >= n_rolls:
            return row[0].row, turn, game, [rolls[k] for k in sorted(rolls)][:n_rolls]
    return None


def load_records(wb, rules, sheet_name=None):
    n = rules.get("rolls_per_turn", 6)
    issues, cands = [], []
    sheets = [sheet_name] if sheet_name else wb.sheetnames
    hdr = None
    for name in sheets:
        if name not in wb.sheetnames:
            fail(f"data_sheet {name!r} not in workbook sheets {wb.sheetnames}")
        hdr = find_header(wb[name], n)
        if hdr:
            ws = wb[name]
            break
    if not hdr:
        fail(f"no header row with 'Turn' and 'Roll 1..{n}' found on sheets {sheets}")
    hrow, tcol, gcol, rcols = hdr
    from openpyxl.utils import get_column_letter as L
    main_cols = {tcol, *rcols} | ({gcol} if gcol else set())
    grid = defaultdict(dict)
    for row in ws.iter_rows(min_row=hrow + 1):
        for c in row:
            if c.value is not None:
                grid[c.row][c.column] = c.value
    last = max(grid) if grid else hrow
    for r in range(hrow + 1, last + 1):
        d = grid.get(r, {})
        vals = [d.get(c) for c in [tcol] + ([gcol] if gcol else []) + rcols]
        if all(v is None for v in vals):
            continue
        ints = [as_int(v) for v in vals]
        if all(v is None for v in ints):
            continue  # a text label such as "End Sheet"
        if any(v is None for v in ints):
            issues.append(f"row {r}: incomplete or non-integer record {vals} - skipped")
            continue
        if any(isinstance(v, str) for v in vals):
            issues.append(f"row {r}: numbers stored as text, coerced")
        t, g = ints[0], (ints[1] if gcol else None)
        cands.append({"turn": t, "game": g, "rolls": ints[-n:], "cell": f"{L(tcol)}{r}", "displaced": False})
    # Sweep for displaced records: a run of (turn, [game], n rolls) integers outside the main columns.
    width = len(main_cols)
    for r, d in sorted(grid.items()):
        cols = sorted(c for c, v in d.items() if c not in main_cols and as_int(v) is not None)
        runs, cur = [], []
        for c in cols:
            if cur and c == cur[-1] + 1:
                cur.append(c)
            else:
                if cur:
                    runs.append(cur)
                cur = [c]
        if cur:
            runs.append(cur)
        for run in runs:
            if len(run) < width:
                issues.append(f"row {r}: stray numbers {[d[c] for c in run]} at {L(run[0])}{r} not a full record - ignored")
                continue
            for k in range(0, len(run) - width + 1, width):
                v = [as_int(d[c]) for c in run[k:k + width]]
                cands.append({"turn": v[0], "game": v[1] if gcol else None, "rolls": v[-n:],
                              "cell": f"{L(run[k])}{r}", "displaced": True})
    return cands, issues, ws.title, hrow


def clean(cands, issues, rules):
    tpg = rules.get("turns_per_game", 2)
    faces = rules.get("die_faces", 6)
    by_turn = {}
    for rec in cands:
        t = rec["turn"]
        if rec["displaced"]:
            issues.append(f"turn {t} (game {rec['game']}) found displaced at {rec['cell']} outside the data block - recovered")
        if t in by_turn:
            if by_turn[t]["rolls"] == rec["rolls"]:
                issues.append(f"turn {t}: duplicate identical record at {rec['cell']} - dropped")
            else:
                issues.append(f"turn {t}: conflicting records at {by_turn[t]['cell']} and {rec['cell']} - kept the first")
            continue
        by_turn[t] = rec
    turns = [by_turn[t] for t in sorted(by_turn)]
    for rec in turns:
        bad = [x for x in rec["rolls"] if not 1 <= x <= faces]
        if bad:
            issues.append(f"turn {rec['turn']}: rolls {bad} outside 1..{faces} - scored as-is, check the question")
        if rec["game"] is None:
            rec["game"] = math.ceil(rec["turn"] / tpg)
        elif rec["game"] != math.ceil(rec["turn"] / tpg):
            issues.append(f"turn {rec['turn']}: game number {rec['game']} != ceil(turn/{tpg}) - using the sheet's game number")
    if turns:
        expected = set(range(turns[0]["turn"], turns[-1]["turn"] + 1))
        missing = sorted(expected - set(by_turn))
        if missing:
            issues.append(f"turn numbers missing from the data: {missing[:30]}")
    games = defaultdict(list)
    for rec in turns:
        games[rec["game"]].append(rec)
    for g, recs in games.items():
        if len(recs) != tpg:
            issues.append(f"game {g} has {len(recs)} turns (expected {tpg}): turns {[x['turn'] for x in recs]}")
    return turns, games


# ---------------------------------------------------------------- scoring

def score(rolls, cat):
    rule = cat["rule"]
    hi, lo = max(rolls), min(rolls)
    if rule == "max_times_count":
        return hi * rolls.count(hi)
    if rule == "sum":
        return sum(rolls)
    if rule == "max_min_diff":
        return hi * lo * (hi - lo)
    if rule == "only_two_numbers":
        k = len(set(rolls))
        ok = k == 2 or (k == 1 and cat.get("single_value_counts", False))
        return cat["points"] if ok else 0
    if rule == "all_faces":
        faces = cat.get("faces", 6)
        return cat["points"] if sorted(rolls) == list(range(1, faces + 1)) else 0
    if rule == "ordered_run":
        k = cat.get("run_length", 4)
        for i in range(len(rolls) - k + 1):
            seg = rolls[i:i + k]
            d = [seg[j + 1] - seg[j] for j in range(k - 1)]
            if all(x == 1 for x in d) or all(x == -1 for x in d):
                return cat["points"]
        return 0
    fail(f"unknown rule {rule!r} in dice_rules (allowed: max_times_count, sum, max_min_diff, only_two_numbers, all_faces, ordered_run)")


def stats(vals):
    if not vals:
        return {}
    out = {"count": len(vals), "total": sum(vals), "mean": statistics.fmean(vals),
           "median": statistics.median(vals), "min": min(vals), "max": max(vals),
           "stdev_sample": statistics.stdev(vals) if len(vals) > 1 else 0.0,
           "stdev_population": statistics.pstdev(vals)}
    c = Counter(vals)
    top = max(c.values())
    out["mode"] = sorted(v for v, n in c.items() if n == top)
    out["distribution"] = {str(k): c[k] for k in sorted(c)}
    return out


def main():
    payload = json.load(sys.stdin)
    st = payload.get("currentState", payload)
    rules = st.get("dice_rules")
    if not isinstance(rules, dict) or not rules.get("categories"):
        fail("dice_rules must be an object with a non-empty 'categories' list (start from assets/dice_rules.json)")
    from openpyxl import load_workbook
    xlsx = resolve(st.get("xlsx_path"))
    wb = load_workbook(xlsx, data_only=True)
    cands, issues, sheet, hrow = load_records(wb, rules, st.get("data_sheet"))
    turns, games = clean(cands, issues, rules)
    if not turns:
        fail("no turn records found under the header row")
    cats = rules["categories"]
    ids = [c["id"] for c in cats]
    for rec in turns:
        rec["scores"] = [score(rec["rolls"], c) for c in cats]
        best = max(rec["scores"])
        rec["best"] = best
        rec["best_cats"] = [ids[i] for i, s in enumerate(rec["scores"]) if s == best]

    unique = rules.get("categories_unique_per_game", True)
    game_rows = []
    for g in sorted(games):
        recs = sorted(games[g], key=lambda x: x["turn"])
        k = len(recs)
        if unique and k <= len(cats):
            combos = itertools.permutations(range(len(cats)), k)
        else:
            combos = itertools.product(range(len(cats)), repeat=k)
        best, opts = None, []
        for combo in combos:
            s = sum(recs[i]["scores"][c] for i, c in enumerate(combo))
            if best is None or s > best:
                best, opts = s, [combo]
            elif s == best:
                opts.append(combo)
        game_rows.append({"game": g, "turns": [r["turn"] for r in recs], "score": best,
                          "optimal": [[ids[c] for c in combo] for combo in opts]})

    work = st.get("work_dir") or os.path.join(tempfile.gettempdir(), "dice_model")
    os.makedirs(work, exist_ok=True)
    tpath, gpath = os.path.join(work, "turns.csv"), os.path.join(work, "games.csv")
    n = rules.get("rolls_per_turn", 6)
    with open(tpath, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["turn", "game", *[f"roll{i + 1}" for i in range(n)], "source_cell", *ids,
                    "best_turn_score", "best_turn_categories"])
        for r in turns:
            w.writerow([r["turn"], r["game"], *r["rolls"], r["cell"], *r["scores"], r["best"], "|".join(r["best_cats"])])
    kmax = max(len(x["turns"]) for x in game_rows)
    with open(gpath, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["game", *[f"turn{i + 1}" for i in range(kmax)], *[f"turn{i + 1}_category" for i in range(kmax)],
                    *[f"turn{i + 1}_points" for i in range(kmax)], "game_score", "n_optimal_assignments", "all_optimal_assignments"])
        by_t = {r["turn"]: r for r in turns}
        for x in game_rows:
            first = x["optimal"][0]
            pts = [by_t[t]["scores"][ids.index(c)] for t, c in zip(x["turns"], first)]
            pad = kmax - len(x["turns"])
            w.writerow([x["game"], *x["turns"], *[""] * pad, *first, *[""] * pad, *pts, *[""] * pad,
                        x["score"], len(x["optimal"]), " ; ".join("+".join(o) for o in x["optimal"])])

    per_cat = {}
    for i, c in enumerate(cats):
        vals = [r["scores"][i] for r in turns]
        s = stats(vals)
        s["turns_scoring_nonzero"] = sum(v > 0 for v in vals)
        nz = [v for v in vals if v > 0]
        s["mean_when_nonzero"] = statistics.fmean(nz) if nz else 0.0
        s["name"] = c.get("name", c["id"])
        s["turns_where_best_category"] = sum(ids[i] in r["best_cats"] for r in turns)
        s["turns_where_sole_best_category"] = sum(r["best_cats"] == [ids[i]] for r in turns)
        per_cat[c["id"]] = s
    gscores = [x["score"] for x in game_rows]
    any_opt = Counter()
    sole_opt = Counter()
    pair_sole = Counter()
    for x in game_rows:
        used = {c for o in x["optimal"] for c in o}
        any_opt.update(used)
        if len(x["optimal"]) == 1:
            sole_opt.update(x["optimal"][0])
            pair_sole["+".join(sorted(x["optimal"][0]))] += 1
    gmax = max(gscores)
    summary = {
        "data_sheet": sheet, "header_row": hrow,
        "n_turns": len(turns), "n_games": len(game_rows),
        "rules_used": {c["id"]: {k: v for k, v in c.items() if k in ("rule", "points", "run_length", "single_value_counts")} for c in cats},
        "turn_category_stats": per_cat,
        "turn_best_score_stats": stats([r["best"] for r in turns]),
        "game_score_stats": stats(gscores),
        "games_at_max_score": [x["game"] for x in game_rows if x["score"] == gmax][:50],
        "games_with_tied_optimal_assignments": sum(len(x["optimal"]) > 1 for x in game_rows),
        "category_in_some_optimal_assignment_games": dict(any_opt),
        "category_in_unique_optimal_assignment_games": dict(sole_opt),
        "category_pair_counts_unique_optimal": dict(pair_sole.most_common()),
        "first_turns": [{"turn": r["turn"], "rolls": r["rolls"], "scores": dict(zip(ids, r["scores"]))} for r in turns[:2]],
        "first_games": game_rows[:2],
    }
    print(json.dumps({"dice_summary": summary, "data_issues": issues, "turns_csv": tpath,
                      "games_csv": gpath, "work_dir": work}, default=str))


if __name__ == "__main__":
    main()
