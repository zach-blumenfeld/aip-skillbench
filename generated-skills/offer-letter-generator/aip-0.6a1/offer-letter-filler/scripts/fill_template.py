#!/usr/bin/env python3
"""Fill a Word (.docx) template's {{KEY}} placeholders from a JSON data file.

Handles what breaks naive python-docx fills:
  - placeholders split across several runs (Word splits on spell-check/format edits)
  - paragraphs anywhere: body, nested tables (any depth), text boxes, headers, footers
  - conditional blocks {{IF_X}}...{{END_IF_X}}, in one paragraph or spanning several:
    kept (markers removed) when the data flag for X is yes-like, removed when no-like
  - formatting: only the characters of each placeholder are rewritten, so the runs
    around it keep their own bold/italic/underline

Two ways to run:
  AIP:  JSON on stdin {"state": {template_path, data_path, output_path}, ...}
        -> JSON on stdout {output_file, filled_placeholders, conditionals, issues,
                           letter_text, status}
  CLI:  python fill_template.py --template T.docx --data D.json --output O.docx
        [--condition RELOCATION=yes] [--map PLACEHOLDER=data_key] [--record N]
"""
import argparse
import json
import re
import sys

from docx import Document
from docx.oxml.ns import qn

PLACEHOLDER_RE = re.compile(r"\{\{\s*([A-Za-z0-9_]+)\s*\}\}")
IF_RE = re.compile(r"\{\{\s*IF_([A-Za-z0-9_]+)\s*\}\}")
TRUE_WORDS = {"yes", "y", "true", "t", "1", "on", "included", "include", "eligible"}
FALSE_WORDS = {"no", "n", "false", "f", "0", "off", "none", "n/a", "na", "", "not eligible", "excluded"}
W_P, W_T = qn("w:p"), qn("w:t")
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"


def norm(key):
    return re.sub(r"[^A-Z0-9]+", "_", str(key).strip().upper()).strip("_")


# ---------------------------------------------------------------- data loading
def flatten(obj, prefix="", out=None):
    """Nested dicts -> flat map. Leaf keys are stored both bare and path-joined."""
    out = {} if out is None else out
    for k, v in obj.items():
        nk = norm(k)
        if isinstance(v, dict):
            flatten(v, f"{prefix}{nk}_", out)
        else:
            out.setdefault(nk, v)
            if prefix:
                out.setdefault(prefix + nk, v)
    return out


def load_data(path, record):
    with open(path, encoding="utf-8-sig") as f:
        raw = json.load(f)
    issues = []
    if isinstance(raw, list):
        if record is None and len(raw) != 1:
            issues.append(f"data file holds {len(raw)} records; pick one with --record N (0-based)")
            raw = raw[0] if raw else {}
        else:
            raw = raw[record or 0]
    elif isinstance(raw, dict) and record is not None:
        lists = [v for v in raw.values() if isinstance(v, list) and v and isinstance(v[0], dict)]
        if lists:
            raw = lists[0][record]
    if not isinstance(raw, dict):
        raise SystemExit(f"data file {path} must hold a JSON object of placeholder values")
    return flatten(raw), issues


def fmt(value):
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, (list, tuple)):
        return ", ".join(fmt(v) for v in value)
    return "" if value is None else str(value)


def truthy(value):
    """True / False for yes-like / no-like flags, None when it cannot tell."""
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value != 0
    if value is None:
        return False
    s = str(value).strip().lower()
    if s in TRUE_WORDS:
        return True
    if s in FALSE_WORDS:
        return False
    return None


def condition_key(name, data, cond_overrides):
    """Find the data flag behind {{IF_<name>}}; returns (data_key, value, include|None)."""
    n = norm(name)
    if n in cond_overrides:
        return "(override)", cond_overrides[n], truthy(cond_overrides[n])
    candidates = [n, f"{n}_PACKAGE", f"INCLUDE_{n}", f"HAS_{n}", f"IS_{n}", f"{n}_ELIGIBLE",
                  f"{n}_FLAG", f"{n}_ENABLED", f"{n}_REQUIRED", f"OFFER_{n}"]
    for c in candidates:
        if c in data and truthy(data[c]) is not None:
            return c, data[c], truthy(data[c])
    # any other key built on the name whose value reads as a yes/no flag
    for k, v in data.items():
        if (k.startswith(n + "_") or k.endswith("_" + n)) and truthy(v) is not None \
                and str(v).strip().lower() not in {"", "0", "1"}:
            return k, v, truthy(v)
    for c in candidates:
        if c in data:
            return c, data[c], None
    return None, None, None


# ---------------------------------------------------------------- xml helpers
def own_texts(p):
    """w:t elements belonging to paragraph p itself (not to paragraphs nested in it)."""
    out = []
    for t in p.iter(W_T):
        anc = t.getparent()
        while anc is not None and anc.tag != W_P:
            anc = anc.getparent()
        if anc is p:
            out.append(t)
    return out


def ptext(p):
    return "".join(t.text or "" for t in own_texts(p))


def splice(p, start, end, replacement):
    """Replace characters [start, end) of p's text, touching only the runs they live in.
    The replacement inherits the formatting of the run holding the first character."""
    pos, placed = 0, False
    for t in own_texts(p):
        text = t.text or ""
        a, b = pos, pos + len(text)
        pos = b
        if b <= start or a == b:
            continue
        if a >= end:
            break
        lo, hi = max(start, a) - a, min(end, b) - a
        t.text = text[:lo] + ("" if placed else replacement) + text[hi:]
        t.set(XML_SPACE, "preserve")
        placed = True
    return placed


def remove_paragraph(p):
    parent = p.getparent()
    if parent is None:
        return
    # a table cell / text box must keep at least one paragraph
    if parent.tag in (qn("w:tc"), qn("w:txbxContent")) and len(parent.findall(W_P)) <= 1:
        for t in own_texts(p):
            t.text = ""
        return
    parent.remove(p)


def stories(doc):
    """Root elements of every text story: body plus each header/footer part."""
    roots = [doc.element.body]
    seen = set()
    for rel in doc.part.rels.values():
        if rel.is_external:
            continue
        part = rel.target_part
        name = str(part.partname)
        if re.search(r"/word/(header|footer)\d*\.xml$", name) and name not in seen:
            seen.add(name)
            roots.append(part.element)
    return roots


# ---------------------------------------------------------------- conditionals
def process_conditionals(root, data, cond_overrides, conditionals, issues):
    for _ in range(500):  # each pass resolves one block; bound guards against bad markup
        paras = list(root.iter(W_P))
        hit = None
        for i, p in enumerate(paras):
            m = IF_RE.search(ptext(p))
            if m:
                hit = (i, p, m)
                break
        if not hit:
            return
        i, p, m = hit
        name = m.group(1)
        end_re = re.compile(r"\{\{\s*END_IF_" + re.escape(name) + r"\s*\}\}")
        key, value, include = condition_key(name, data, cond_overrides)
        info = conditionals.setdefault(name, {"data_key": key, "value": fmt(value) if key else None,
                                               "included": include})
        if include is None:
            info["included"] = None
            msg = (f"conditional IF_{name}: no yes/no data flag found" if key is None else
                   f"conditional IF_{name}: data key {key}={value!r} is not a clear yes/no")
            issues.append(msg + "; markers left in place, rerun with --condition " + name + "=yes|no")
            # neutralise this start marker so the loop moves on; restore text afterwards
            splice(p, m.start(), m.end(), "\ue000IF_" + name + "\ue000")
            continue
        # locate END marker: same paragraph after the start, else a later paragraph
        same = end_re.search(ptext(p), m.end())
        if same:
            if include:
                splice(p, same.start(), same.end(), "")
                splice(p, m.start(), m.end(), "")
            else:
                splice(p, m.start(), same.end(), "")
            if not ptext(p).strip():
                remove_paragraph(p)
            continue
        j = next((k for k in range(i + 1, len(paras)) if end_re.search(ptext(paras[k]))), None)
        if j is None:
            issues.append(f"conditional IF_{name}: no matching END_IF_{name}; start marker removed")
            splice(p, m.start(), m.end(), "")
            continue
        q = paras[j]
        em = end_re.search(ptext(q))
        if include:
            splice(q, em.start(), em.end(), "")
            splice(p, m.start(), m.end(), "")
        else:
            splice(q, 0, em.end(), "")
            splice(p, m.start(), len(ptext(p)), "")
            if p.getparent() is q.getparent():
                sib = p.getnext()
                while sib is not None and sib is not q:
                    nxt = sib.getnext()
                    sib.getparent().remove(sib)
                    sib = nxt
            else:
                for mid in paras[i + 1:j]:
                    if mid.getparent() is not None:
                        remove_paragraph(mid)
        for x in (q, p):
            if x.getparent() is not None and not ptext(x).strip():
                remove_paragraph(x)
    issues.append("conditional processing stopped after 500 blocks")


def restore_unresolved(root):
    for p in root.iter(W_P):
        for t in own_texts(p):
            if t.text and "\ue000" in t.text:
                t.text = re.sub("\ue000IF_([A-Za-z0-9_]+)\ue000", r"{{IF_\1}}", t.text)


# ---------------------------------------------------------------- placeholders
def process_placeholders(root, data, key_map, filled, missing):
    for p in list(root.iter(W_P)):
        text = ptext(p)
        matches = list(PLACEHOLDER_RE.finditer(text))
        for m in reversed(matches):  # right-to-left keeps earlier offsets valid
            name = m.group(1)
            if name.upper().startswith(("IF_", "END_IF_")):
                continue  # unresolved conditional markers stay visible
            key = norm(key_map.get(norm(name), name))
            if key not in data:
                missing.add(name)
                continue
            value = fmt(data[key])
            # template already carries the currency sign: avoid "$$185,000"
            if m.start() > 0 and text[m.start() - 1] in "$€£" and value[:1] == text[m.start() - 1]:
                value = value[1:]
            splice(p, m.start(), m.end(), value)
            filled[name] = value


def reread(path):
    """Re-open the saved file: (leftover template markup, plain text of every story)."""
    doc = Document(path)
    found, lines = [], []
    for root in stories(doc):
        label = "body" if root is doc.element.body else root.tag.split("}")[-1]  # hdr / ftr
        for p in root.iter(W_P):
            text = ptext(p)
            found += [m.group(0) for m in re.finditer(r"\{\{[^{}]*\}\}|\{\{|\}\}", text)]
            if text.strip():
                lines.append(text if label == "body" else f"[{label}] {text}")
    return sorted(set(found)), "\n".join(lines)


# ---------------------------------------------------------------- driver
def fill(template_path, data_path, output_path, cond_overrides=None, key_map=None, record=None):
    cond_overrides = {norm(k): v for k, v in (cond_overrides or {}).items()}
    key_map = {norm(k): v for k, v in (key_map or {}).items()}
    data, issues = load_data(data_path, record)
    doc = Document(template_path)
    conditionals, filled, missing = {}, {}, set()
    for root in stories(doc):
        process_conditionals(root, data, cond_overrides, conditionals, issues)
        restore_unresolved(root)
        process_placeholders(root, data, key_map, filled, missing)
    doc.save(output_path)
    for name in sorted(missing):
        issues.append(f"placeholder {{{{{name}}}}} has no value in the data file; left in place")
    left, letter_text = reread(output_path)
    if left:
        issues.append("output still contains template markup: " + ", ".join(left))
    return {
        "output_file": output_path,
        "filled_placeholders": filled,
        "conditionals": conditionals,
        "issues": issues,
        "letter_text": letter_text,
        "status": "needs_attention" if issues else "complete",
    }


def kv_pairs(items):
    out = {}
    for item in items or []:
        if "=" not in item:
            raise SystemExit(f"expected KEY=VALUE, got {item!r}")
        k, v = item.split("=", 1)
        out[k.strip()] = v.strip()
    return out


def main():
    if len(sys.argv) > 1:
        ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
        ap.add_argument("--template", required=True)
        ap.add_argument("--data", required=True)
        ap.add_argument("--output", required=True)
        ap.add_argument("--condition", action="append", help="NAME=yes|no for {{IF_NAME}} blocks")
        ap.add_argument("--map", action="append", help="PLACEHOLDER=data_key when names differ")
        ap.add_argument("--record", type=int, help="0-based record when the data file holds several")
        a = ap.parse_args()
        result = fill(a.template, a.data, a.output, kv_pairs(a.condition), kv_pairs(a.map), a.record)
    else:
        state = json.load(sys.stdin).get("state", {})
        result = fill(state["template_path"], state["data_path"], state["output_path"])
    json.dump(result, sys.stdout, indent=2)
    sys.stdout.write("\n")


if __name__ == "__main__":
    main()
