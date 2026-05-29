#!/usr/bin/env python3
"""
Build a powerlifting-coefficient workbook.

Reads an input `.xlsx`, locates the named source sheet, and writes a target
sheet whose columns are Excel-formula references to the source sheet plus a
`TotalKg` sum formula and a coefficient formula (DOTS / Wilks / IPF
GoodLift). The output is written to the requested path (defaults to the
input — overwrite in place, mimicking the curated task's expectation).

Only `openpyxl` is required at runtime (pre-installed in the curated task
image).

Usage:
    python build_coef_workbook.py \
        --input  /root/data/openipf.xlsx \
        --output /root/data/openipf.xlsx \
        --data-sheet   Data \
        --target-sheet Dots \
        --coef dots \
        --precision 3 \
        --columns 'Name=A,Sex=B,BodyweightKg=I,Best3SquatKg=K,Best3BenchKg=L,Best3DeadliftKg=M'

`--columns` is an ordered comma-separated list of
`<header>=<source-column-letter>` pairs. The target sheet receives them in
the given order, followed by `TotalKg` and a coefficient column named
after `--coef`.

The `Best3SquatKg`, `Best3BenchKg`, `Best3DeadliftKg` headers (in that
order, anywhere in `--columns`) are required for the TotalKg sum.
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path
from typing import Sequence

import openpyxl


# -- formula constants -----------------------------------------------------

# DOTS — (a, b, c, d, e) and (bw_min, bw_max)
DOTS_M = (-0.0000010930, 0.0007391293, -0.1918759221, 24.0900756, -307.75076)
DOTS_F = (-0.0000010706, 0.0005158568, -0.1126655495, 13.6175032, -57.96288)
DOTS_BW_M = (40.0, 210.0)
DOTS_BW_F = (40.0, 150.0)

# Wilks — (a, b, c, d, e, f), reverse-indexed: a + b·x + c·x² + …
WILKS_M = (-216.0475144, 16.2606339, -0.002388645, -0.00113732, 7.01863e-06, -1.291e-08)
WILKS_F = (594.31747775582, -27.23842536447, 0.82112226871, -0.00930733913, 4.731582e-05, -9.054e-08)
WILKS_BW_M = (40.0, 201.9)
WILKS_BW_F = (26.51, 154.53)

# IPF GoodLift — (A, B, C) keyed by (event, sex, equipment)
GOODLIFT: dict[tuple[str, str, str], tuple[float, float, float]] = {
    ("SBD", "M", "Raw"):    (1199.72839, 1025.18162, 0.009210),
    ("SBD", "M", "Single"): (1236.25115, 1449.21864, 0.01644),
    ("SBD", "F", "Raw"):    (610.32796,  1045.59282, 0.03048),
    ("SBD", "F", "Single"): (758.63878,  949.31382,  0.02435),
    ("B",   "M", "Raw"):    (320.98041,  281.40258,  0.01008),
    ("B",   "M", "Single"): (381.22073,  733.79378,  0.02398),
    ("B",   "F", "Raw"):    (142.40398,  442.52671,  0.04724),
    ("B",   "F", "Single"): (221.82209,  357.00377,  0.02937),
}


# -- formula builders ------------------------------------------------------

def _clamp(cell: str, lo: float, hi: float) -> str:
    return f"MAX({lo},MIN({hi},{cell}))"


def _poly4(bw: str, coefs: Sequence[float]) -> str:
    """a·bw⁴ + b·bw³ + c·bw² + d·bw + e."""
    a, b, c, d, e = coefs
    return (
        f"({a}*POWER({bw},4)"
        f"+{b}*POWER({bw},3)"
        f"+{c}*POWER({bw},2)"
        f"+{d}*{bw}"
        f"+{e})"
    )


def _poly5_reverse(bw: str, coefs: Sequence[float]) -> str:
    """a + b·bw + c·bw² + d·bw³ + e·bw⁴ + f·bw⁵ (Wilks convention)."""
    a, b, c, d, e, f = coefs
    return (
        f"({a}"
        f"+{b}*{bw}"
        f"+{c}*POWER({bw},2)"
        f"+{d}*POWER({bw},3)"
        f"+{e}*POWER({bw},4)"
        f"+{f}*POWER({bw},5))"
    )


def dots_formula(sex_cell: str, bw_cell: str, total_cell: str, precision: int) -> str:
    male_bw   = _clamp(bw_cell, *DOTS_BW_M)
    female_bw = _clamp(bw_cell, *DOTS_BW_F)
    male   = f"{total_cell}*(500/{_poly4(male_bw,   DOTS_M)})"
    female = f"{total_cell}*(500/{_poly4(female_bw, DOTS_F)})"
    return f'=ROUND(IF({sex_cell}="F",{female},{male}),{precision})'


def wilks_formula(sex_cell: str, bw_cell: str, total_cell: str, precision: int) -> str:
    male_bw   = _clamp(bw_cell, *WILKS_BW_M)
    female_bw = _clamp(bw_cell, *WILKS_BW_F)
    male   = f"{total_cell}*(500/{_poly5_reverse(male_bw,   WILKS_M)})"
    female = f"{total_cell}*(500/{_poly5_reverse(female_bw, WILKS_F)})"
    return f'=ROUND(IF({sex_cell}="F",{female},{male}),{precision})'


def goodlift_formula(
    sex_cell: str,
    bw_cell: str,
    total_cell: str,
    precision: int,
    *,
    equipment: str,
    event: str,
) -> str:
    """
    Sex × equipment × event pin (A, B, C). The Excel-level IF on the sex
    cell picks the male or female parameter set per row; the inner IF
    guards BW < 35 and a non-positive denominator (per the reference impl).
    """
    m = GOODLIFT.get((event, "M", equipment))
    f = GOODLIFT.get((event, "F", equipment))
    if m is None or f is None:
        raise SystemExit(
            f"IPF GoodLift: no parameters for event={event} equipment={equipment}. "
            f"Known combos: {sorted(GOODLIFT.keys())}"
        )

    def one(params: tuple[float, float, float]) -> str:
        A, B, C = params
        denom = f"({A}-{B}*EXP(-{C}*{bw_cell}))"
        return f"IF(AND({bw_cell}>=35,{denom}>0),{total_cell}*100/{denom},0)"

    return f'=ROUND(IF({sex_cell}="F",{one(f)},{one(m)}),{precision})'


def total_formula(squat_cell: str, bench_cell: str, dead_cell: str) -> str:
    return f"={squat_cell}+{bench_cell}+{dead_cell}"


# -- workbook assembly -----------------------------------------------------

def parse_columns(raw: str) -> list[tuple[str, str]]:
    out: list[tuple[str, str]] = []
    for piece in raw.split(","):
        piece = piece.strip()
        if not piece:
            continue
        if "=" not in piece:
            raise SystemExit(f"Bad --columns entry {piece!r}; want HEADER=COL_LETTER")
        header, col = piece.split("=", 1)
        out.append((header.strip(), col.strip().upper()))
    if not out:
        raise SystemExit("--columns produced no entries")
    return out


def required_cell(header: str, mapping: list[tuple[str, str]]) -> str | None:
    """Return the source-sheet column letter for `header`, or None."""
    for name, col in mapping:
        if name == header:
            return col
    return None


def build(args: argparse.Namespace) -> None:
    in_path = Path(args.input)
    out_path = Path(args.output)
    if not in_path.exists():
        raise SystemExit(f"Input not found: {in_path}")

    wb = openpyxl.load_workbook(in_path)
    if args.data_sheet not in wb.sheetnames:
        raise SystemExit(
            f"Source sheet {args.data_sheet!r} not in workbook. "
            f"Have: {wb.sheetnames}"
        )
    data_ws = wb[args.data_sheet]
    n_rows = data_ws.max_row - 1  # subtract header

    mapping = parse_columns(args.columns)
    headers = [h for h, _ in mapping]

    # Each of these three is needed to compute TotalKg.
    sq = required_cell("Best3SquatKg",    mapping)
    bn = required_cell("Best3BenchKg",    mapping)
    dl = required_cell("Best3DeadliftKg", mapping)
    if not (sq and bn and dl):
        raise SystemExit(
            "TotalKg formula needs Best3SquatKg, Best3BenchKg, Best3DeadliftKg "
            "in --columns mapping."
        )

    sex_col = required_cell("Sex", mapping)
    bw_col  = required_cell("BodyweightKg", mapping)
    if args.coef in {"dots", "wilks", "goodlift"} and not (sex_col and bw_col):
        raise SystemExit("Sex and BodyweightKg columns are required for this coefficient.")

    # Drop & recreate the target sheet so we start clean.
    if args.target_sheet in wb.sheetnames:
        del wb[args.target_sheet]
    ws = wb.create_sheet(args.target_sheet)

    coef_header = {
        "dots":     "Dots",
        "wilks":    "Wilks",
        "goodlift": "Goodlift",
    }[args.coef]

    full_headers = list(headers) + ["TotalKg", coef_header]
    for col_idx, h in enumerate(full_headers, start=1):
        ws.cell(row=1, column=col_idx, value=h)

    n_static = len(headers)
    total_col_letter = openpyxl.utils.get_column_letter(n_static + 1)
    coef_col_letter  = openpyxl.utils.get_column_letter(n_static + 2)

    # Per-row formulas. Excel rows are 1-indexed; row 1 is the header.
    for r in range(2, n_rows + 2):
        # Static column references to the Data sheet.
        for col_idx, (_, src_col) in enumerate(mapping, start=1):
            ws.cell(
                row=r,
                column=col_idx,
                value=f"={args.data_sheet}!{src_col}{r}",
            )

        # Resolve the target-sheet cells that the formulas need to point at.
        def t(header: str) -> str:
            i = headers.index(header) + 1
            return f"{openpyxl.utils.get_column_letter(i)}{r}"

        ws.cell(row=r, column=n_static + 1, value=total_formula(t("Best3SquatKg"), t("Best3BenchKg"), t("Best3DeadliftKg")))

        sex_ref   = t("Sex") if sex_col else None
        bw_ref    = t("BodyweightKg") if bw_col else None
        total_ref = f"{total_col_letter}{r}"

        if args.coef == "dots":
            f = dots_formula(sex_ref, bw_ref, total_ref, args.precision)
        elif args.coef == "wilks":
            f = wilks_formula(sex_ref, bw_ref, total_ref, args.precision)
        elif args.coef == "goodlift":
            f = goodlift_formula(
                sex_ref, bw_ref, total_ref, args.precision,
                equipment=args.equipment,
                event=args.event,
            )
        else:
            raise SystemExit(f"Unknown --coef {args.coef!r}")

        ws.cell(row=r, column=n_static + 2, value=f)

    wb.save(out_path)
    sys.stderr.write(
        f"Wrote {n_rows} rows to {out_path}: sheet {args.target_sheet!r}, "
        f"columns {full_headers}, coef={args.coef}\n"
    )


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--input",  required=True)
    p.add_argument("--output", required=True)
    p.add_argument("--data-sheet",   default="Data")
    p.add_argument("--target-sheet", default="Dots")
    p.add_argument("--coef", choices=["dots", "wilks", "goodlift"], default="dots")
    p.add_argument("--precision", type=int, default=3)
    p.add_argument("--columns", required=True,
                   help="HEADER=COL,HEADER=COL,…  Required headers: Name, Sex, "
                        "BodyweightKg, Best3SquatKg, Best3BenchKg, Best3DeadliftKg.")
    p.add_argument("--equipment", default="Raw",
                   help="IPF GoodLift only. Raw|Single. Wraps/Straps map to Raw; "
                        "Multi/Unlimited map to Single before being passed in.")
    p.add_argument("--event", default="SBD",
                   help="IPF GoodLift only. SBD (full power) or B (bench only).")
    args = p.parse_args(argv)
    build(args)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
