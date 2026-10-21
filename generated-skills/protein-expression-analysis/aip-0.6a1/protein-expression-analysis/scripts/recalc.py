#!/usr/bin/env python3
"""
Excel Formula Recalculation Script
Recalculates all formulas in an Excel file using LibreOffice

Adapted from the xlsx skill's recalc.py (source/xlsx/recalc.py): same macro and
error scan, but the timeout is enforced with subprocess(timeout=...) instead of the
external `timeout`/`gtimeout` command, the first-run profile initialisation gets a
longer timeout and never raises, and the soffice binary is looked up on PATH.
"""

import json
import os
import platform
import shutil
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

EXCEL_ERRORS = ["#VALUE!", "#DIV/0!", "#REF!", "#NAME?", "#NULL!", "#NUM!", "#N/A"]


def soffice_bin():
    return shutil.which("soffice") or shutil.which("libreoffice")


def setup_libreoffice_macro():
    """Setup LibreOffice macro for recalculation if not already configured"""
    if platform.system() == "Darwin":
        macro_dir = os.path.expanduser("~/Library/Application Support/LibreOffice/4/user/basic/Standard")
    else:
        macro_dir = os.path.expanduser("~/.config/libreoffice/4/user/basic/Standard")

    macro_file = os.path.join(macro_dir, "Module1.xba")

    if os.path.exists(macro_file):
        with open(macro_file) as f:
            if "RecalculateAndSave" in f.read():
                return True

    if not os.path.exists(macro_dir):
        try:
            subprocess.run([soffice_bin(), "--headless", "--terminate_after_init"], capture_output=True, timeout=90)
        except (subprocess.TimeoutExpired, OSError):
            pass
        os.makedirs(macro_dir, exist_ok=True)

    macro_content = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
    Sub RecalculateAndSave()
      ThisComponent.calculateAll()
      ThisComponent.store()
      ThisComponent.close(True)
    End Sub
</script:module>"""

    try:
        with open(macro_file, "w") as f:
            f.write(macro_content)
        return True
    except Exception:
        return False


def scan_errors(filename):
    wb = load_workbook(filename, data_only=True)
    details = {err: [] for err in EXCEL_ERRORS}
    total = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    for err in EXCEL_ERRORS:
                        if err in cell.value:
                            details[err].append(f"{ws.title}!{cell.coordinate}")
                            total += 1
                            break
    wb.close()
    return total, {k: {"count": len(v), "locations": v[:20]} for k, v in details.items() if v}


def recalc(filename, timeout=120):
    """Recalculate formulas in an Excel file in place; return a JSON-able report."""
    if not Path(filename).exists():
        return {"error": f"File {filename} does not exist"}
    if not soffice_bin():
        return {"error": "LibreOffice (soffice) not found on PATH"}

    abs_path = str(Path(filename).absolute())
    if not setup_libreoffice_macro():
        return {"error": "Failed to setup LibreOffice macro"}

    cmd = [
        soffice_bin(),
        "--headless",
        "--norestore",
        "vnd.sun.star.script:Standard.Module1.RecalculateAndSave?language=Basic&location=application",
        abs_path,
    ]
    try:
        result = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout)
    except subprocess.TimeoutExpired:
        return {"error": f"LibreOffice recalculation timed out after {timeout}s"}
    if result.returncode != 0:
        return {"error": (result.stderr or "Unknown error during recalculation").strip()[:500]}

    try:
        total, summary = scan_errors(filename)
        wb_f = load_workbook(filename)
        formulas = sum(1 for ws in wb_f.worksheets for row in ws.iter_rows() for c in row
                       if isinstance(c.value, str) and c.value.startswith("="))
        wb_f.close()
        return {"status": "success" if total == 0 else "errors_found", "total_errors": total,
                "error_summary": summary, "total_formulas": formulas}
    except Exception as e:
        return {"error": str(e)}


def main():
    if len(sys.argv) < 2:
        print("Usage: python recalc.py <excel_file> [timeout_seconds]")
        sys.exit(1)
    timeout = int(sys.argv[2]) if len(sys.argv) > 2 else 120
    print(json.dumps(recalc(sys.argv[1], timeout), indent=2))


if __name__ == "__main__":
    main()
