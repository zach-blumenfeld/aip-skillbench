#!/usr/bin/env python3
"""Profile a lifter workbook: sheets, headers, row counts, value sets, anomalies,
and a draft plan for write_scores.py. stdin: {"currentState": {...}}; stdout: JSON."""

import json
import sys

import pl_lib as P

P.ensure_openpyxl()

from openpyxl import load_workbook  # noqa: E402

from pl_lib import CANON, SYSTEMS, find_header, norm, num, resolve_path


def last_row(ws, header_row):
    last = header_row
    for r in range(header_row + 1, ws.max_row + 1):
        if any(c.value not in (None, "") for c in ws[r]):
            last = r
    return last


def main():
    state = json.load(sys.stdin)["currentState"]
    path = resolve_path(state["workbook_path"])
    wb = load_workbook(path)
    sheets = []
    for ws in wb.worksheets:
        headers = [c.value for c in ws[1]] if ws.max_row >= 1 else []
        headers = [h for h in headers if h not in (None, "")]
        sheets.append({
            "name": ws.title,
            "headers": headers,
            "data_rows": max(0, last_row(ws, 1) - 1) if headers else 0,
            "empty": all(c.value in (None, "") for row in ws.iter_rows() for c in row),
            "tables": list(ws.tables.keys()),
        })

    # data sheet = the one whose headers hold sex + bodyweight
    data = next((s for s in sheets if find_header(s["headers"], "sex") and find_header(s["headers"], "bodyweight")), None)
    if data is None:
        print(json.dumps({"profile": {"workbook_path": path, "sheets": sheets,
                                      "error": "no sheet has both a Sex and a Bodyweight header row in row 1"}}))
        return
    ws = wb[data["name"]]
    headers = [c.value for c in ws[1]]
    col = {h: i for i, h in enumerate(headers) if h not in (None, "")}
    fields = {k: find_header(data["headers"], k) for k in CANON}
    rows = list(ws.iter_rows(min_row=2, max_row=data["data_rows"] + 1, values_only=True))

    def values(field):
        h = fields.get(field)
        return [r[col[h]] for r in rows] if h else []

    def counts(field):
        out = {}
        for v in values(field):
            out[str(v)] = out.get(str(v), 0) + 1
        return out

    anomalies = []
    for i, r in enumerate(rows, start=2):
        name = r[col[fields["name"]]] if fields["name"] else f"row {i}"
        bw = r[col[fields["bodyweight"]]]
        if bw in (None, "") or isinstance(bw, str) or num(bw) <= 0:
            anomalies.append(f"row {i} ({name}): bodyweight {bw!r} is blank, text, or <= 0")
        for lift in ("squat", "bench", "deadlift", "total"):
            h = fields.get(lift)
            if not h:
                continue
            v = r[col[h]]
            if isinstance(v, str) and v.strip():
                anomalies.append(f"row {i} ({name}): {h} is text {v!r}")
            elif isinstance(v, (int, float)) and v < 0:
                anomalies.append(f"row {i} ({name}): {h} = {v} (negative = failed lift)")
        sx = r[col[fields["sex"]]]
        if str(sx).strip() not in ("M", "F", "Mx"):
            anomalies.append(f"row {i} ({name}): Sex {sx!r} not M/F/Mx (scored as men)")
        if fields["place"] and str(r[col[fields["place"]]]).strip().upper() in ("DQ", "DD", "NS"):
            anomalies.append(f"row {i} ({name}): Place {r[col[fields['place']]]} (OPL leaves TotalKg empty)")

    others = [s for s in sheets if s["name"] != data["name"]]
    target = next((s for s in others if s["empty"]), None)
    system_hint = None
    for s in sheets:
        n = norm(s["name"])
        for sysname, keys in {"dots": ["dots"], "wilks": ["wilks"], "ipf_gl": ["ipfgl", "goodlift", "gl", "ipfpoints"],
                              "glossbrenner": ["glossbrenner"]}.items():
            if n in keys:
                system_hint = sysname
    copy_fields = ["name", "sex", "bodyweight", "squat", "bench", "deadlift"]
    columns = [{"header": fields[k], "kind": "copy", "source": fields[k]} for k in copy_fields if fields[k]]
    lifts = [fields[k] for k in ("squat", "bench", "deadlift") if fields[k]]
    if fields["total"]:
        columns.append({"header": fields["total"], "kind": "copy", "source": fields["total"]})
    else:
        columns.append({"header": "TotalKg", "kind": "total", "sources": lifts})
    score_header = target["name"] if target else "Score"
    columns.append({"header": score_header, "kind": "score", "system": system_hint or "dots", "round": None})

    profile = {
        "workbook_path": path,
        "sheets": sheets,
        "data_sheet": data["name"],
        "field_headers": fields,
        "sex_counts": counts("sex"),
        "event_counts": counts("event"),
        "equipment_counts": counts("equipment"),
        "bodyweight_range": [min((num(v) for v in values("bodyweight")), default=None),
                             max((num(v) for v in values("bodyweight")), default=None)],
        "first_rows": [dict(zip(data["headers"], r)) for r in rows[:3]],
        "number_formats": {h: ws.cell(row=2, column=col[h] + 1).number_format for h in data["headers"]} if rows else {},
        "anomalies": anomalies[:50],
        "candidate_target_sheet": target["name"] if target else None,
        "system_hint_from_sheet_name": system_hint,
        "supported_systems": list(SYSTEMS),
    }
    draft = {
        "output_path": path,
        "data_sheet": data["name"],
        "target_sheet": target["name"] if target else "Scores",
        "copy_mode": "formula",
        "failed_lift_policy": "opl",
        "clear_target": False,
        "columns": columns,
    }
    print(json.dumps({"profile": profile, "draft_plan": draft}, default=str))


if __name__ == "__main__":
    main()
