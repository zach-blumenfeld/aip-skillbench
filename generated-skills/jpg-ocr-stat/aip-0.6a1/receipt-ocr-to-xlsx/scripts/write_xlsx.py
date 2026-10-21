#!/usr/bin/env python3
"""Merge review corrections into the extracted records, write the results workbook,
then re-open it and verify the exact layout the task requires.

aip mode: stdin {"state": {"records": [...], "output_path": "...", "corrections": [...]?}}.
With `corrections` in the state (reviewed path) it prints {"final_rows": [...]};
without (no review needed) it prints {"final_rows_auto": [...]}.
CLI mode: python3 write_xlsx.py <records.json> <output.xlsx> [corrections.json]
"""
import datetime
import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

SHEET = "results"
HEADER = ["filename", "date", "total_amount"]


def clean_date(v):
    if v in (None, "", "null"):
        return None
    v = str(v).strip()
    try:
        return datetime.date.fromisoformat(v).isoformat() if re.fullmatch(r"\d{4}-\d{2}-\d{2}", v) else None
    except ValueError:
        return None


def clean_total(v):
    if v in (None, "", "null"):
        return None
    v = str(v).strip().replace(",", "")
    v = re.sub(r"^(RM|MYR|\$)\s*", "", v, flags=re.I)
    try:
        return f"{float(v):.2f}"
    except ValueError:
        return None


def build_rows(records, corrections):
    rows = {r["filename"]: {"filename": r["filename"], "date": r.get("date"),
                            "total_amount": r.get("total_amount")} for r in records}
    for c in corrections or []:
        name = c.get("filename")
        if name not in rows:
            raise SystemExit(f"correction for unknown file: {name!r}")
        for k in ("date", "total_amount"):
            if k in c:
                rows[name][k] = c[k]
    out = []
    for name in sorted(rows):
        r = rows[name]
        out.append({"filename": name, "date": clean_date(r["date"]),
                    "total_amount": clean_total(r["total_amount"])})
    return out


def write(rows, path):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append(HEADER)
    for r in rows:
        # strings, not dates/floats: "2018-03-23" and "10.00" must survive as written;
        # a null field is an empty cell
        ws.append([r["filename"], r["date"], r["total_amount"]])
    wb.save(path)


def verify(rows, path):
    wb = load_workbook(path)
    problems = []
    if wb.sheetnames != [SHEET]:
        problems.append(f"sheets are {wb.sheetnames}, expected ['{SHEET}']")
    ws = wb[wb.sheetnames[0]]
    got = [list(r) for r in ws.iter_rows(values_only=True)]
    want = [HEADER] + [[r["filename"], r["date"], r["total_amount"]] for r in rows]
    if got != want:
        problems.append("sheet content differs from the rows written")
    if ws.max_column != 3:
        problems.append(f"{ws.max_column} columns, expected 3")
    names = [r["filename"] for r in rows]
    if names != sorted(names):
        problems.append("rows not ordered by filename")
    if problems:
        raise SystemExit("verification failed: " + "; ".join(problems))


def main():
    if len(sys.argv) > 2:
        records = json.load(open(sys.argv[1]))
        records = records.get("records", records) if isinstance(records, dict) else records
        corr = json.load(open(sys.argv[3])) if len(sys.argv) > 3 else None
        rows = build_rows(records, corr)
        write(rows, sys.argv[2])
        verify(rows, sys.argv[2])
        print(json.dumps(rows, indent=1))
        return
    state = json.load(sys.stdin)["state"]
    reviewed = "corrections" in state
    rows = build_rows(state["records"], state.get("corrections"))
    write(rows, state["output_path"])
    verify(rows, state["output_path"])
    print(json.dumps({"final_rows" if reviewed else "final_rows_auto": rows}))


if __name__ == "__main__":
    main()
