#!/usr/bin/env python3
"""Fill the protein-expression template with live Excel formulas and compute the
same numbers in Python (the shadow values verify_workbook.py checks against).

stdin currentState keys: workbook_path, output_path, layout, data_scale (log2|linear),
stdev_kind (sample|population), add_top_table (bool).
stdout: {"written_cells", "expected", "results", "top_regulated", "build_notes"}
"""
import json
import math
import statistics
import sys

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import column_index_from_string as CI
from openpyxl.utils import get_column_letter as L


def fail(msg):
    print(json.dumps({"error": msg}))
    sys.exit(1)


def num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def ranges(cols, row):
    """Contiguous columns -> 'C11:G11'; otherwise 'C11,E11,G11'."""
    idx = [CI(c) for c in cols]
    if idx == list(range(idx[0], idx[0] + len(idx))):
        return f"{cols[0]}{row}:{cols[-1]}{row}"
    return ",".join(f"{c}{row}" for c in cols)


def main():
    st = json.load(sys.stdin)["currentState"]
    lay = st["layout"]
    scale = st.get("data_scale", "log2")
    sd_fn = "STDEVP" if st.get("stdev_kind") == "population" else "STDEV"  # legacy names: no _xlfn prefix needed
    out = st.get("output_path") or st["workbook_path"]

    wb = load_workbook(st["workbook_path"])
    ws, ds = wb[lay["task_sheet"]], wb[lay["data_sheet"]]
    d = lay["data"]
    q = "'" + ds.title.replace("'", "''") + "'!"
    fr, lr, vc, lc, idc = d["first_row"], d["last_row"], d["first_value_col"], d["last_col"], d["id_col"]
    hr, idcol, gcol = lay["header_row"], lay["id_col"], lay.get("gene_col")
    rows, ctrl, trt = lay["protein_rows"], lay["control_cols"], lay["treated_cols"]

    dhdr = {str(c.value).strip(): c.column for c in ds[d["header_row"]] if c.value is not None}
    drow = {str(ds.cell(r, CI(idc)).value).strip(): r for r in range(lr, fr - 1, -1)}  # first match wins, like MATCH
    written, expected, notes = [], {}, []

    def put(cell, formula, value):
        ws[cell] = formula
        written.append(cell)
        expected[cell] = value

    # Step 1: lookups. Blank source cells stay blank ("") instead of INDEX's silent 0.
    look = {}
    for r in rows:
        pid = str(ws[f"{idcol}{r}"].value).strip()
        for c in ctrl + trt:
            sample = str(ws[f"{c}{hr}"].value).strip()
            ix = (f"INDEX({q}${vc}${fr}:${lc}${lr},MATCH(${idcol}{r},{q}${idc}${fr}:${idc}${lr},0),"
                  f"MATCH({c}${hr},{q}${vc}${d['header_row']}:${lc}${d['header_row']},0))")
            v = ds.cell(drow[pid], dhdr[sample]).value if pid in drow and sample in dhdr else None
            look[(r, c)] = v if num(v) else None
            put(f"{c}{r}", f'=IFERROR(IF({ix}="","",{ix}),"")', look[(r, c)])

    # Step 2: statistics, one column per protein (transposed against the lookup block).
    sr, scols = lay["stats_rows"], lay["stat_cols"]
    sdf = statistics.pstdev if sd_fn == "STDEVP" else statistics.stdev
    min_sd = 2  # an SD of one value (STDEVP gives 0) is not a measurement; leave it blank
    res = []
    hdr_row = lay.get("stats_header_row")
    for i, r in enumerate(rows):
        col = scols[i]
        if hdr_row and gcol and ws[f"{col}{hdr_row}"].value is None:
            ws[f"{col}{hdr_row}"] = f"={gcol}{r}"
            ws[f"{col}{hdr_row}"].font = Font(bold=True)
            written.append(f"{col}{hdr_row}")
        cv = [look[(r, c)] for c in ctrl if look[(r, c)] is not None]
        tv = [look[(r, c)] for c in trt if look[(r, c)] is not None]
        cm = statistics.fmean(cv) if cv else None
        tm = statistics.fmean(tv) if tv else None
        csd = sdf(cv) if len(cv) >= min_sd else None
        tsd = sdf(tv) if len(tv) >= min_sd else None
        cr, tr_ = ranges(ctrl, r), ranges(trt, r)
        if "control_mean" in sr:
            put(f"{col}{sr['control_mean']}", f'=IF(COUNT({cr})=0,"",AVERAGE({cr}))', cm)
        if "control_sd" in sr:
            put(f"{col}{sr['control_sd']}", f'=IF(COUNT({cr})<{min_sd},"",{sd_fn}({cr}))', csd)
        if "treated_mean" in sr:
            put(f"{col}{sr['treated_mean']}", f'=IF(COUNT({tr_})=0,"",AVERAGE({tr_}))', tm)
        if "treated_sd" in sr:
            put(f"{col}{sr['treated_sd']}", f'=IF(COUNT({tr_})<{min_sd},"",{sd_fn}({tr_}))', tsd)
        if scale == "log2":
            l2 = tm - cm if cm is not None and tm is not None else None
            fc = 2 ** l2 if l2 is not None else None
        else:
            fc = tm / cm if cm not in (None, 0) and tm is not None else None
            l2 = math.log2(fc) if fc is not None and fc > 0 else None
        res.append({"row": r, "protein_id": ws[f"{idcol}{r}"].value, "gene": ws[f"{gcol}{r}"].value if gcol else None,
                    "n_control": len(cv), "n_treated": len(tv), "control_mean": cm, "control_sd": csd,
                    "treated_mean": tm, "treated_sd": tsd, "fold_change": fc, "log2_fc": l2, "stat_col": col})
        if cm is None or tm is None:
            notes.append(f"{res[-1]['gene']}: a group has no values, so its fold change is blank")
        elif csd is None or tsd is None:
            notes.append(f"{res[-1]['gene']}: fewer than {min_sd} values in a group, so that StdDev is blank")

    # Step 3: fold change table. Log2 data: Log2 FC = treated mean - control mean, FC = 2^Log2 FC.
    fct = lay.get("fold_change")
    if fct:
        fcc, lgc = fct.get("fc_col"), fct.get("log2_col")
        cmr, tmr = sr.get("control_mean"), sr.get("treated_mean")
        for i, rr in enumerate(fct["rows"]):
            x, col, r = res[i], scols[i], rows[i]
            for tc, src in ((fct.get("id_col"), idcol), (fct.get("gene_col"), gcol)):
                if tc and src and ws[f"{tc}{rr}"].value is None:
                    put(f"{tc}{rr}", f"={src}{r}", ws[f"{src}{r}"].value)
            cm_ref, tm_ref = f"{col}{cmr}", f"{col}{tmr}"
            if scale == "log2":
                if lgc:
                    put(f"{lgc}{rr}", f'=IF(OR({cm_ref}="",{tm_ref}=""),"",{tm_ref}-{cm_ref})', x["log2_fc"])
                if fcc:
                    src = f"{lgc}{rr}" if lgc else f"({tm_ref}-{cm_ref})"
                    put(f"{fcc}{rr}", f'=IF(OR({cm_ref}="",{tm_ref}=""),"",POWER(2,{src}))', x["fold_change"])
            else:
                if fcc:
                    put(f"{fcc}{rr}", f'=IF(OR({cm_ref}="",{tm_ref}=""),"",IF({cm_ref}=0,"",{tm_ref}/{cm_ref}))', x["fold_change"])
                if lgc:
                    src = f"{fcc}{rr}" if fcc else f"({tm_ref}/{cm_ref})"
                    put(f"{lgc}{rr}", f'=IF(ISNUMBER({src}),IF({src}>0,LOG({src},2),""),"")', x["log2_fc"])
            x["fc_row"] = rr

    # Ranking by |log2 FC| (largest first); ties go to the earlier row so ranks are unique; blanks are unranked.
    ranked = sorted([x for x in res if x["log2_fc"] is not None], key=lambda x: -abs(x["log2_fc"]))
    for i, x in enumerate(res):
        a = abs(x["log2_fc"]) if x["log2_fc"] is not None else None
        x["rank"] = (1 + sum(abs(y["log2_fc"]) > a for y in ranked)
                     + sum(y["log2_fc"] is not None and abs(y["log2_fc"]) == a for y in res[:i])) if a is not None else None
        x["direction"] = (None if x["log2_fc"] is None else "Up" if x["log2_fc"] > 0 else "Down" if x["log2_fc"] < 0 else "No change")

    # Step 4 (optional): live ranking table under the fold-change table.
    top = lay.get("top_table")
    if st.get("add_top_table") and top and fct and fct.get("log2_col"):
        lgc = fct["log2_col"]
        t, h, f0 = top["title_row"], top["header_row"], top["first_row"]
        hfill = PatternFill("solid", start_color="ADD8E6")
        if not top.get("existing_template"):
            ws[f"A{t}"] = "Step 4: Top Regulated Proteins (ranked by |Log2 FC|, 1 = strongest; sorted list in H:L)"
            ws[f"A{t}"].font = Font(bold=True)
        for j, name in enumerate(["Protein_ID", "Gene_Symbol", "Log2 FC", "|Log2 FC|", "Rank", "Direction"]):
            cell = ws.cell(h, 1 + j, name)
            cell.font, cell.fill = Font(bold=True), hfill
        n = len(fct["rows"])
        absr = f"$D${f0}:$D${f0 + n - 1}"
        for i, rr in enumerate(fct["rows"]):
            x, r = res[i], f0 + i
            put(f"A{r}", f"={idcol}{rows[i]}", x["protein_id"])
            put(f"B{r}", f"={gcol}{rows[i]}" if gcol else '=""', x["gene"])
            put(f"C{r}", f"={lgc}{rr}", x["log2_fc"])
            put(f"D{r}", f'=IF(ISNUMBER(C{r}),ABS(C{r}),"")', abs(x["log2_fc"]) if x["log2_fc"] is not None else None)
            put(f"E{r}", f'=IF(ISNUMBER(D{r}),COUNTIF({absr},">"&D{r})+COUNTIF($D${f0}:D{r},D{r}),"")', x["rank"])
            put(f"F{r}", f'=IF(ISNUMBER(C{r}),IF(C{r}>0,"Up",IF(C{r}<0,"Down","No change")),"")', x["direction"])

        # Sorted view in H:L (strongest first), looked up from the rank column.
        for j, name in enumerate(["Rank", "Protein_ID", "Gene_Symbol", "Log2 FC", "Direction"]):
            cell = ws.cell(h, 8 + j, name)
            cell.font, cell.fill = Font(bold=True), hfill
        byrank = {x["rank"]: x for x in res if x["rank"] is not None}
        rk = f"$E${f0}:$E${f0 + n - 1}"
        for k in range(1, n + 1):
            r, y = f0 + k - 1, byrank.get(k)
            ws[f"H{r}"] = k
            for col, src, key in (("I", "A", "protein_id"), ("J", "B", "gene"), ("K", "C", "log2_fc"), ("L", "F", "direction")):
                put(f"{col}{r}", f'=IFERROR(INDEX(${src}${f0}:${src}${f0 + n - 1},MATCH(H{r},{rk},0)),"")', y[key] if y else None)

    try:
        wb.save(out)
    except Exception as e:
        fail(f"cannot save {out!r}: {e}")
    rnd = lambda v: round(v, 6) if num(v) else v
    print(json.dumps({
        "output_path": out,
        "written_cells": len(written),
        "expected": expected,
        "results": [{k: rnd(v) for k, v in x.items()} for x in res],
        "top_regulated": [{"rank": x["rank"], "gene": x["gene"], "log2_fc": rnd(x["log2_fc"]),
                           "fold_change": rnd(x["fold_change"]), "direction": x["direction"]} for x in ranked],
        "build_notes": notes,
    }))


if __name__ == "__main__":
    main()
