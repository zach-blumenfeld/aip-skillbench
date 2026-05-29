#!/usr/bin/env python3
# /// script
# requires-python = ">=3.12"
# dependencies = [
#   "openpyxl==3.1.5",
# ]
# ///
"""Verify the structure of a Dots-populated workbook.

Checks (all must pass for ok=true):
  1. Dots sheet exists.
  2. Header row equals ["Name","Sex","BodyweightKg","Best3SquatKg",
     "Best3BenchKg","Best3DeadliftKg","TotalKg","Dots"] in order.
  3. G2 is a formula string (starts with "=") summing D2+E2+F2.
  4. H2 is a formula string (starts with "=") that references B2 (Sex),
     C2 (BodyweightKg), and G2 (TotalKg) and contains ROUND(...,3).

Returns JSON {"ok": bool, "notes": [...]}.

Run via:
  uv run scripts/verify_output.py --output-path /root/data/openipf.xlsx
"""

from __future__ import annotations

import argparse
import json
import sys

import openpyxl

EXPECTED_HEADERS = [
    "Name",
    "Sex",
    "BodyweightKg",
    "Best3SquatKg",
    "Best3BenchKg",
    "Best3DeadliftKg",
    "TotalKg",
    "Dots",
]


def verify(output_path: str) -> dict:
    notes: list[str] = []
    wb = openpyxl.load_workbook(output_path, data_only=False)

    if "Dots" not in wb.sheetnames:
        notes.append("Dots sheet missing")
        return {"ok": False, "notes": notes}
    ws = wb["Dots"]

    headers = [ws.cell(row=1, column=i + 1).value for i in range(len(EXPECTED_HEADERS))]
    if headers != EXPECTED_HEADERS:
        notes.append(f"Header row mismatch: expected {EXPECTED_HEADERS}, got {headers}")

    g2 = ws["G2"].value
    if not (isinstance(g2, str) and g2.startswith("=") and "D2" in g2 and "E2" in g2 and "F2" in g2):
        notes.append(f"G2 should be a sum-formula referencing D2/E2/F2; got {g2!r}")

    h2 = ws["H2"].value
    if not isinstance(h2, str) or not h2.startswith("="):
        notes.append(f"H2 should be a formula string; got {h2!r}")
    else:
        for ref, label in [("B2", "Sex"), ("C2", "BodyweightKg"), ("G2", "TotalKg")]:
            if ref not in h2:
                notes.append(f"H2 missing reference to {ref} ({label}): {h2!r}")
        if "ROUND" not in h2.upper() or ",3)" not in h2.replace(" ", ""):
            notes.append(f"H2 should ROUND(...,3) at the outer level: {h2!r}")

    return {"ok": not notes, "notes": notes}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-path", required=True)
    args = parser.parse_args()
    result = verify(args.output_path)
    json.dump(result, sys.stdout, indent=2)
    sys.stdout.write("\n")
    return 0 if result["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
