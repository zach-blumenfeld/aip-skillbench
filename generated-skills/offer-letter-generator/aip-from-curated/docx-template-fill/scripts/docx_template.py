"""Shared logic for filling {{KEY}} placeholders in a Word (.docx) template.

Works on the raw WordprocessingML of every text-bearing part (body, nested
tables, text boxes, every header/footer, footnotes, endnotes), so nothing the
python-docx API skips (non-default headers, nested tables, text boxes) is missed.

Replacement is done per paragraph on the concatenated <w:t> text, so placeholders
Word split across runs ("{{CANDI" + "DATE_NAME}}") are found. The replacement
value is written into the run where the placeholder starts and the placeholder's
remaining characters are deleted from the following runs, so every run keeps its
own formatting (a bold/underlined split placeholder stays bold/underlined; plain
text around it stays plain).

Requires python-docx (any 1.x). Compatible with Python 3.8+.
"""
import json
import os
import re
import sys

try:
    from docx import Document
    from docx.oxml.ns import qn
except ImportError:  # pragma: no cover
    sys.stdout.write(json.dumps({"error": "python-docx is not installed; pip install python-docx"}))
    sys.exit(1)

from lxml import etree

PLACEHOLDER_RE = re.compile(r"\{\{\s*([A-Za-z0-9_]+)\s*\}\}")
IF_RE = re.compile(r"\{\{\s*IF_([A-Za-z0-9_]+)\s*\}\}")
TRUE_WORDS = {"yes", "y", "true", "t", "1", "include", "included", "on", "eligible"}
FALSE_WORDS = {"no", "n", "false", "f", "0", "none", "exclude", "excluded", "off", "", "n/a", "na", "null"}
PART_RE = re.compile(r"^/word/(document|header\d*|footer\d*|footnotes|endnotes)\.xml$")
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"


# ---------------------------------------------------------------- io helpers

def read_stdin():
    raw = sys.stdin.read()
    payload = json.loads(raw) if raw.strip() else {}
    return payload.get("currentState", payload)


def emit(obj):
    sys.stdout.write(json.dumps(obj, ensure_ascii=False))
    sys.stdout.flush()


def resolve_path(p):
    """Scripts run with cwd = scripts/, so resolve relative paths against the caller's PWD."""
    p = os.path.expanduser(str(p))
    if os.path.isabs(p):
        return p
    base = os.environ.get("PWD") or os.getcwd()
    return os.path.abspath(os.path.join(base, p))


def load_data(state):
    """Data comes from state['data'] (object) if present, else the JSON file at state['data_path']."""
    data = state.get("data")
    if not isinstance(data, dict):
        path = resolve_path(state["data_path"])
        with open(path, encoding="utf-8-sig") as f:
            data = json.load(f)
    if not isinstance(data, dict):
        raise ValueError("employee/template data must be a JSON object of KEY -> value")
    return data


def to_text(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def truthiness(value):
    """True / False for a recognised flag value, None when it cannot be read as a flag."""
    if isinstance(value, bool):
        return value
    if value is None:
        return False
    if isinstance(value, (int, float)):
        return value != 0
    s = str(value).strip().lower()
    if s in TRUE_WORDS:
        return True
    if s in FALSE_WORDS:
        return False
    return None


# ---------------------------------------------------------------- part access

def text_parts(doc):
    """Yield (partname, root_element, part) for every text-bearing XML part, each once."""
    seen = set()
    for part in doc.part.package.iter_parts():
        name = str(part.partname)
        if not PART_RE.match(name) or name in seen:
            continue
        seen.add(name)
        el = getattr(part, "_element", None)
        if el is None:  # generic Part (e.g. footnotes in python-docx): parse the blob
            el = etree.fromstring(part.blob)
            part._aip_parsed = el
        yield name, el, part


def flush_parts(doc):
    """Write back any blob-only parts we parsed and modified."""
    for part in doc.part.package.iter_parts():
        el = getattr(part, "_aip_parsed", None)
        if el is not None:
            part._blob = etree.tostring(el, xml_declaration=True, encoding="UTF-8", standalone=True)


def paragraphs(root):
    return list(root.iter(qn("w:p")))


def own_texts(p):
    """<w:t> elements whose nearest enclosing paragraph is p (skip text-box paragraphs nested inside)."""
    out = []
    for t in p.iter(qn("w:t")):
        anc = t.getparent()
        while anc is not None and anc.tag != qn("w:p"):
            anc = anc.getparent()
        if anc is p:
            out.append(t)
    return out


def para_text(p):
    return "".join(t.text or "" for t in own_texts(p))


def _set(t, s):
    t.text = s
    if s != s.strip() or "  " in s:
        t.set(XML_SPACE, "preserve")


def splice(texts, start, end, value):
    """Replace characters [start, end) of the concatenated text of `texts` with value.

    The value lands in the <w:t> holding `start`; the rest of the span is removed from
    the following <w:t>s. Returns nothing; mutates in place.
    """
    pos = 0
    first = None
    for t in texts:
        s = t.text or ""
        a, b = pos, pos + len(s)
        pos = b
        if b <= start:
            continue
        if a >= end and first is not None:
            break
        lo = max(start, a) - a
        hi = min(end, b) - a
        if first is None:
            _set(t, s[:lo] + value + s[hi:])
            first = t
        else:
            _set(t, s[:lo] + s[hi:])
    if first is None and texts:  # span at the very end
        _set(texts[-1], (texts[-1].text or "") + value)
    # Drop runs the splice emptied (only formatting + an empty <w:t> left).
    for t in texts:
        r = t.getparent()
        if t is first or (t.text or "") or r is None or r.tag != qn("w:r"):
            continue
        if all(c.tag in (qn("w:rPr"), qn("w:t")) and (c.tag != qn("w:t") or not (c.text or "")) for c in r):
            if r.getparent() is not None:
                r.getparent().remove(r)


def split_run_count(p, start, end):
    pos, n = 0, 0
    for t in own_texts(p):
        s = t.text or ""
        a, b = pos, pos + len(s)
        pos = b
        if b > start and a < end:
            n += 1
    return n


# ---------------------------------------------------------------- inspection

def scan(doc):
    """List every placeholder occurrence and conditional marker, part by part."""
    occurrences, conditionals = [], []
    for name, root, _ in text_parts(doc):
        for i, p in enumerate(paragraphs(root)):
            txt = para_text(p)
            for m in PLACEHOLDER_RE.finditer(txt):
                key = m.group(1)
                prev = txt[m.start() - 1] if m.start() > 0 else ""
                occurrences.append({
                    "part": name, "paragraph": i, "key": key,
                    "runs": split_run_count(p, m.start(), m.end()),
                    "prefix_char": prev, "context": txt[max(0, m.start() - 30):m.end() + 30],
                })
    for o in occurrences:
        if o["key"].startswith("IF_") or o["key"].startswith("END_IF_"):
            conditionals.append(o)
    return occurrences, conditionals


def guess_condition_key(cond, data):
    """Which data key governs {{IF_<cond>}}? Exact key, else <cond>_* keys holding a flag-like value."""
    if cond in data and truthiness(data[cond]) is not None:
        return cond, []
    cands = [k for k in data if k == cond or k.startswith(cond + "_") or k.endswith("_" + cond)]
    flags = [k for k in cands if str(data[k]).strip().lower() in TRUE_WORDS | FALSE_WORDS or isinstance(data[k], bool)]
    if len(flags) == 1:
        return flags[0], cands
    if cond in data:
        return cond, cands
    return None, cands


# ---------------------------------------------------------------- conditionals

def apply_conditionals(doc, conditions):
    """Keep (strip markers) or drop (remove content between markers) each IF_ block.

    Handles blocks inside one paragraph and blocks spanning several paragraphs in the
    same container. Paragraphs left empty by a dropped block that spans whole
    paragraphs are removed; a single-paragraph block leaves an empty paragraph, as the
    source pattern does.
    """
    report = []
    for name, root, _ in text_parts(doc):
        changed = True
        while changed:
            changed = False
            paras = paragraphs(root)
            for idx, p in enumerate(paras):
                txt = para_text(p)
                m = IF_RE.search(txt)
                if not m:
                    continue
                cond = m.group(1)
                include = bool(conditions.get(cond, False))
                end_re = re.compile(r"\{\{\s*END_IF_" + re.escape(cond) + r"\s*\}\}")
                e = end_re.search(txt, m.end())
                texts = own_texts(p)
                if e:  # same paragraph
                    if include:
                        splice(texts, e.start(), e.end(), "")
                        splice(own_texts(p), m.start(), m.end(), "")
                    else:
                        splice(texts, m.start(), e.end(), "")
                    report.append({"part": name, "condition": cond, "included": include, "span": "paragraph"})
                    changed = True
                    break
                # multi-paragraph block: find the closing paragraph
                j = None
                for k in range(idx + 1, len(paras)):
                    if end_re.search(para_text(paras[k])):
                        j = k
                        break
                if j is None:
                    raise ValueError("{{IF_%s}} in %s has no matching {{END_IF_%s}}" % (cond, name, cond))
                q = paras[j]
                qe = end_re.search(para_text(q))
                if include:
                    splice(own_texts(q), qe.start(), qe.end(), "")
                    splice(own_texts(p), m.start(), m.end(), "")
                else:
                    splice(own_texts(q), 0, qe.end(), "")
                    splice(texts, m.start(), len(txt), "")
                    for mid in paras[idx + 1:j]:
                        if mid.getparent() is not None:
                            mid.getparent().remove(mid)
                    for edge in (p, q):
                        if not para_text(edge).strip() and edge.getparent() is not None and len(edge.getparent().findall(qn("w:p"))) > 1:
                            edge.getparent().remove(edge)
                report.append({"part": name, "condition": cond, "included": include, "span": "multi-paragraph"})
                changed = True
                break
    return report


# ---------------------------------------------------------------- placeholders

def apply_placeholders(doc, values):
    """Replace every {{KEY}} with values[KEY]. Unknown keys are left in place and reported."""
    replaced, unresolved = {}, []
    for name, root, _ in text_parts(doc):
        for p in paragraphs(root):
            txt = para_text(p)
            matches = list(PLACEHOLDER_RE.finditer(txt))
            for m in reversed(matches):
                key = m.group(1)
                if key not in values:
                    unresolved.append({"part": name, "key": key, "context": txt})
                    continue
                val = to_text(values[key])
                # Template already supplies a currency symbol right before the placeholder:
                # do not double it ("$$185,000").
                if m.start() > 0 and txt[m.start() - 1] in "$€£¥" and val.startswith(txt[m.start() - 1]):
                    val = val[1:]
                splice(own_texts(p), m.start(), m.end(), val)
                replaced[key] = replaced.get(key, 0) + 1
    return replaced, unresolved


# ---------------------------------------------------------------- verification

def all_text(doc):
    chunks = []
    for name, root, _ in text_parts(doc):
        for p in paragraphs(root):
            chunks.append((name, para_text(p)))
    return chunks


def verify(path, values, expected_keys):
    doc = Document(path)
    chunks = all_text(doc)
    problems = []
    for name, txt in chunks:
        if "{{" in txt or "}}" in txt:
            problems.append({"part": name, "issue": "leftover braces", "text": txt})
        if re.search(r"\b(END_)?IF_[A-Z0-9_]+\b", txt):
            problems.append({"part": name, "issue": "conditional marker left behind", "text": txt})
    joined = "\n".join(t for _, t in chunks)
    missing_values = [k for k in expected_keys
                      if to_text(values.get(k)) and to_text(values.get(k)).lstrip("$€£¥") not in joined]
    return problems, missing_values, joined
