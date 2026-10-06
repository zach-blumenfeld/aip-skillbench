#!/usr/bin/env python3
"""Read the xlsx workbook (every sheet, values + formulas, number formats) and
the background PDF (text per page), emit one JSON object to stdout the client
can reason over without opening the files again.

stdin: {"currentState": {"xlsx_path", "pdf_path", "question", ...}, "assets": {}, "expects": {...}}
stdout: merges {"xlsx_data": {...}, "pdf_text": "..."} into state.

xlsx_data layout:
  {
    "sheets": {
      "<sheet>": {
        "dims": [rows, cols],
        "cells": [
          {"ref": "B5", "row": 5, "col": 2,
           "value": <calc value, may be null if never opened in Excel>,
           "formula": "=SUM(...)" | null,
           "number_format": "0.0%" | null,
           "comment": "Source: ..." | null}
        ],
        "named_ranges": {"name": "Sheet1!$A$1:$B$3", ...}
      }
    },
    "workbook_defined_names": {...},
    "notes": ["warnings or caveats encountered during extraction"]
  }

Only cells whose value or formula is non-empty are emitted — a dense dump of a
sparse workbook crushes the context budget.
"""
import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter
from pypdf import PdfReader


def _cell_record(cell, formula_cell):
    value = cell.value
    formula = formula_cell.value if formula_cell is not None else None
    is_formula = isinstance(formula, str) and formula.startswith("=")
    if not is_formula:
        formula = None
    comment = cell.comment.text if cell.comment is not None else None
    if value in (None, "") and formula is None and comment is None:
        return None
    return {
        "ref": cell.coordinate,
        "row": cell.row,
        "col": cell.column,
        "value": value,
        "formula": formula,
        "number_format": cell.number_format if cell.number_format != "General" else None,
        "comment": comment,
    }


def _extract_xlsx(path: Path) -> dict:
    notes: list[str] = []
    try:
        wb_values = openpyxl.load_workbook(path, data_only=True)
    except Exception as exc:  # pragma: no cover
        return {"sheets": {}, "workbook_defined_names": {}, "notes": [f"xlsx open (values) failed: {exc}"]}
    wb_formulas = openpyxl.load_workbook(path, data_only=False)

    sheets: dict = {}
    for sheet_name in wb_values.sheetnames:
        ws_values = wb_values[sheet_name]
        ws_formulas = wb_formulas[sheet_name]
        cells: list = []
        max_row = ws_values.max_row or 0
        max_col = ws_values.max_column or 0
        for row in ws_values.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col):
            for cell in row:
                formula_cell = ws_formulas.cell(row=cell.row, column=cell.column)
                rec = _cell_record(cell, formula_cell)
                if rec is not None:
                    cells.append(rec)
        named_ranges: dict = {}
        if hasattr(ws_formulas, "defined_names"):
            try:
                for dn in getattr(ws_formulas.defined_names, "definedName", []) or []:
                    named_ranges[dn.name] = dn.value
            except Exception:
                pass
        sheets[sheet_name] = {
            "dims": [max_row, max_col],
            "cells": cells,
            "named_ranges": named_ranges,
        }

    workbook_defined_names: dict = {}
    try:
        for dn_name in wb_formulas.defined_names:
            dn = wb_formulas.defined_names[dn_name]
            workbook_defined_names[dn_name] = str(dn.value) if dn is not None else None
    except Exception as exc:
        notes.append(f"workbook defined-names: {exc}")

    formula_cells = [
        rec
        for sheet in sheets.values()
        for rec in sheet["cells"]
        if rec.get("formula") is not None
    ]
    null_formula_values = sum(1 for r in formula_cells if r.get("value") is None)
    if formula_cells and null_formula_values == len(formula_cells):
        notes.append(
            "openpyxl returned no cached values for any formula cell — workbook was likely never "
            "opened in Excel/LibreOffice. Formulas are present; evaluate them yourself from the "
            "dependency chain rather than treating null as 'empty'."
        )
    elif null_formula_values:
        notes.append(
            f"{null_formula_values} of {len(formula_cells)} formula cells have null cached values; "
            "recompute those from the formula chain if the question needs them."
        )
    return {
        "sheets": sheets,
        "workbook_defined_names": workbook_defined_names,
        "notes": notes,
    }


def _extract_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    parts: list[str] = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception as exc:
            text = f"[page {i}: extract failed: {exc}]"
        parts.append(f"=== Page {i} ===\n{text.strip()}")
    return "\n\n".join(parts)


def main() -> None:
    payload = json.loads(sys.stdin.read() or "{}")
    state = payload.get("currentState") or payload
    xlsx_path = Path(state["xlsx_path"]).expanduser()
    pdf_path = Path(state["pdf_path"]).expanduser()
    if not xlsx_path.exists():
        raise SystemExit(f"xlsx_path not found: {xlsx_path}")
    if not pdf_path.exists():
        raise SystemExit(f"pdf_path not found: {pdf_path}")
    out = {
        "xlsx_data": _extract_xlsx(xlsx_path),
        "pdf_text": _extract_pdf(pdf_path),
    }
    json.dump(out, sys.stdout)


if __name__ == "__main__":
    main()
