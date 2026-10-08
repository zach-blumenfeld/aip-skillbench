#!/usr/bin/env python3
"""Map a protein-expression task workbook: task sheet, data sheet, yellow input
regions, sample groups, statistic rows, fold-change table, data scale.

stdin:  {"currentState": {"workbook_path": ...}, ...}
stdout: {"layout": {...}, "scale_hint": {...}, "layout_warnings": [...], "sheet_instructions": [...]}
"""
import json
import re
import statistics
import sys

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter as L

YELLOW = {"FFFFFF00", "00FFFF00"}
STAT_KEYS = [
    ("control_mean", r"control.*mean|mean.*control"),
    ("control_sd", r"control.*(std|sd|stdev|deviation)|(std|sd).*control"),
    ("treated_mean", r"treat.*mean|mean.*treat"),
    ("treated_sd", r"treat.*(std|sd|stdev|deviation)|(std|sd).*treat"),
]


def fail(msg):
    print(json.dumps({"error": msg}))
    sys.exit(1)


def text(v):
    return str(v).strip() if v is not None else ""


def is_yellow(c):
    f = c.fill
    return bool(f and f.fill_type == "solid" and (f.fgColor.rgb or "") in YELLOW)


def find_task_sheet(wb):
    for ws in wb.worksheets:
        for row in ws.iter_rows(max_col=3):
            for c in row:
                if text(c.value).lower().startswith("sample group"):
                    return ws, c.row
    return None, None


def find_data_sheet(wb, task_ws):
    for ws in wb.worksheets:
        if ws is task_ws:
            continue
        hdr = [text(c.value).lower() for c in ws[1]]
        if "protein_id" in hdr:
            return ws
    return None


def main():
    state = json.load(sys.stdin)["currentState"]
    path = state.get("workbook_path")
    try:
        wb = load_workbook(path)
    except Exception as e:
        fail(f"cannot open workbook {path!r}: {e}")
    warnings = []

    ws, group_row = find_task_sheet(wb)
    if ws is None:
        fail("no sheet has a 'Sample Group:' label; set the layout by hand")
    ds = find_data_sheet(wb, ws)
    if ds is None:
        fail("no data sheet with a 'Protein_ID' header in row 1")

    sheet_instructions = [text(ws.cell(r, 1).value) for r in range(1, group_row) if text(ws.cell(r, 1).value)]

    # Header row: 'Protein_ID' in column A at or just below the group row.
    header_row = next((r for r in range(group_row, group_row + 4) if text(ws.cell(r, 1).value).lower() == "protein_id"), None)
    if header_row is None:
        fail(f"no 'Protein_ID' header below the Sample Group row {group_row}")
    gene_col = next((c for c in range(2, 5) if "gene" in text(ws.cell(header_row, c).value).lower()), None)

    control, treated, other = [], [], []
    for c in range(2, ws.max_column + 1):
        g = text(ws.cell(group_row, c).value).lower()
        if not g:
            continue
        if g.startswith("control") or g in ("ctrl", "vehicle", "untreated"):
            control.append(c)
        elif g.startswith("treat"):
            treated.append(c)
        else:
            other.append(f"{L(c)}{group_row}={ws.cell(group_row, c).value!r}")
    if other:
        warnings.append(f"unrecognised sample-group labels: {other}")
    if not control or not treated:
        warnings.append("could not find both Control and Treated sample columns")

    protein_rows = []
    r = header_row + 1
    while text(ws.cell(r, 1).value):
        protein_rows.append(r)
        r += 1
    if not protein_rows:
        fail("no protein rows under the Protein_ID header")
    n = len(protein_rows)

    sample_cols = control + treated
    lookup_cells = [f"{L(c)}{r}" for r in protein_rows for c in sample_cols]
    not_yellow = [x for x in lookup_cells if not is_yellow(ws[x])]
    if not_yellow:
        warnings.append(f"{len(not_yellow)} lookup cells are not yellow, e.g. {not_yellow[:5]}")

    # Statistic rows, labelled in column A below the protein block.
    stats = {}
    for r in range(protein_rows[-1] + 1, ws.max_row + 1):
        lab = text(ws.cell(r, 1).value).lower()
        for key, pat in STAT_KEYS:
            if key not in stats and re.search(pat, lab):
                stats[key] = r
                break
    missing_stats = [k for k, _ in STAT_KEYS if k not in stats]
    if missing_stats:
        warnings.append(f"statistic rows not found: {missing_stats}")
    stat_cols = []
    if stats:
        r0 = min(stats.values())
        stat_cols = [c for c in range(2, ws.max_column + 1) if is_yellow(ws.cell(r0, c))]
        if len(stat_cols) != n:
            warnings.append(f"statistic row {r0} has {len(stat_cols)} yellow cells but there are {n} proteins; "
                            "falling back to one column per protein starting at column B")
            stat_cols = list(range(2, 2 + n))
    stats_header_row = (min(stats.values()) - 1) if stats else None

    # Fold-change table: a second 'Protein_ID' header row below the statistics.
    fc = None
    start = max(stats.values()) + 1 if stats else protein_rows[-1] + 1
    for r in range(start, ws.max_row + 1):
        if text(ws.cell(r, 1).value).lower() == "protein_id":
            heads = {text(ws.cell(r, c).value).lower(): c for c in range(1, ws.max_column + 1) if text(ws.cell(r, c).value)}
            fc_col = next((c for h, c in heads.items() if "fold" in h and "log" not in h), None)
            log_col = next((c for h, c in heads.items() if "log" in h), None)
            gcol = next((c for h, c in heads.items() if "gene" in h), None)
            fc = {"header_row": r, "rows": list(range(r + 1, r + 1 + n)), "id_col": "A",
                  "gene_col": L(gcol) if gcol else None,
                  "fc_col": L(fc_col) if fc_col else None, "log2_col": L(log_col) if log_col else None}
            if not fc_col or not log_col:
                warnings.append(f"fold-change header row {r} lacks a 'Fold Change' or 'Log2 FC' column")
            break
    if fc is None:
        warnings.append("no fold-change table (second Protein_ID header) found")

    # Step 4 (top regulated): use a labelled region if the template has one.
    wants_top = any(re.search(r"top|regulated|rank", s, re.I) for s in sheet_instructions)
    top_region = None
    for r in range((fc["rows"][-1] + 1) if fc else 1, ws.max_row + 1):
        if re.search(r"top|regulated|rank", text(ws.cell(r, 1).value), re.I):
            top_region = {"title_row": r}
            break
    last_used = max(fc["rows"][-1] if fc else 0, ws.max_row)
    top = None
    if wants_top or top_region:
        title_row = top_region["title_row"] if top_region else last_used + 2
        top = {"title_row": title_row, "header_row": title_row + 1, "first_row": title_row + 2,
               "existing_template": bool(top_region)}

    # Instructions often quote stale ranges; compare with the yellow cells.
    yellow = {f"{L(c.column)}{c.row}" for row in ws.iter_rows() for c in row if is_yellow(c)}
    for s in sheet_instructions:
        for a, b in re.findall(r"\b([A-Z]{1,3}\d+):([A-Z]{1,3}\d+)\b", s):
            if a not in yellow or b not in yellow:
                warnings.append(f"sheet instruction range {a}:{b} does not match the yellow cells "
                                f"({s!r}); the yellow cells and row labels are authoritative")

    # Data sheet bounds and checks.
    dhdr = {text(c.value): c.column for c in ds[1] if text(c.value)}
    d_last_row = max(r for r in range(1, ds.max_row + 1) if text(ds.cell(r, 1).value))
    d_last_col = max(dhdr.values())
    first_num_col = next((c for c in range(2, d_last_col + 1)
                          if any(isinstance(ds.cell(r, c).value, (int, float)) for r in range(2, min(d_last_row, 50) + 1))), 4)
    ids = {text(ds.cell(r, 1).value): r for r in range(2, d_last_row + 1)}
    missing_ids = [text(ws.cell(r, 1).value) for r in protein_rows if text(ws.cell(r, 1).value) not in ids]
    if missing_ids:
        warnings.append(f"protein IDs absent from {ds.title}: {missing_ids}")
    missing_samples = [text(ws.cell(header_row, c).value) for c in sample_cols if text(ws.cell(header_row, c).value) not in dhdr]
    if missing_samples:
        warnings.append(f"sample headers absent from {ds.title} row 1: {missing_samples}")
    dup = len(ids) != d_last_row - 1
    if dup:
        warnings.append(f"duplicate Protein_IDs in {ds.title}; MATCH returns the first")

    vals = [v for row in ds.iter_rows(min_row=2, max_row=d_last_row, min_col=first_num_col, max_col=d_last_col, values_only=True)
            for v in row if isinstance(v, (int, float)) and not isinstance(v, bool)]
    cells = (d_last_row - 1) * (d_last_col - first_num_col + 1)
    scale_hint = {
        "n_values": len(vals), "n_blank": cells - len(vals),
        "min": min(vals) if vals else None, "max": max(vals) if vals else None,
        "median": statistics.median(vals) if vals else None,
        "fraction_negative": round(sum(v < 0 for v in vals) / len(vals), 4) if vals else None,
        "note": "Negative values or values centred near 0 mean the data are already log2 (ratios/abundances); "
                "large all-positive values (e.g. intensities in the 1e3-1e9 range) are linear.",
    }

    # Per-protein missingness among the looked-up cells.
    sparse = []
    for r in protein_rows:
        pid = text(ws.cell(r, 1).value)
        if pid not in ids:
            continue
        dr = ids[pid]
        heads = lambda cols: [dhdr[h] for h in (text(ws.cell(header_row, c).value) for c in cols) if h in dhdr]
        cnt = {g: sum(isinstance(ds.cell(dr, dc).value, (int, float)) for dc in heads(cols))
               for g, cols in (("control", control), ("treated", treated))}
        if cnt["control"] < len(control) or cnt["treated"] < len(treated):
            sparse.append({"row": r, "protein": pid, "control_values": cnt["control"], "treated_values": cnt["treated"]})

    layout = {
        "task_sheet": ws.title, "data_sheet": ds.title,
        "group_row": group_row, "header_row": header_row, "id_col": "A", "gene_col": L(gene_col) if gene_col else None,
        "protein_rows": protein_rows,
        "control_cols": [L(c) for c in control], "treated_cols": [L(c) for c in treated],
        "stats_header_row": stats_header_row, "stats_rows": stats, "stat_cols": [L(c) for c in stat_cols],
        "fold_change": fc, "top_table": top,
        "data": {"header_row": 1, "first_row": 2, "last_row": d_last_row, "id_col": "A",
                 "first_value_col": L(first_num_col), "last_col": L(d_last_col)},
        "missing_values": sparse,
    }
    print(json.dumps({"layout": layout, "scale_hint": scale_hint, "layout_warnings": warnings,
                      "sheet_instructions": sheet_instructions}))


if __name__ == "__main__":
    main()
