"""Step `inspect`: map the template's placeholders and conditional blocks against the data.

stdin  {"currentState": {template_path, data_path | data, output_path}, ...}
stdout {values, conditions, condition_sources, placeholders, split_placeholders,
        placeholder_parts, issues, notes, inspect_status}
"""
import os
import sys
import traceback

import docx_template as dt
from docx import Document


def main():
    state = dt.read_stdin()
    issues, notes = [], []
    template = dt.resolve_path(state["template_path"])
    output = dt.resolve_path(state["output_path"])
    if not os.path.isfile(template):
        dt.emit({"inspect_status": "needs_input", "issues": ["template not found: %s" % template],
                 "values": {}, "conditions": {}})
        return
    data = dt.load_data(state)
    values = {k: v for k, v in data.items()}

    if os.path.abspath(output) == os.path.abspath(template):
        issues.append("output_path equals template_path; write the filled letter to a new file, never over the template")
    if not output.lower().endswith(".docx"):
        issues.append("output_path should end in .docx: %s" % output)

    doc = Document(template)
    occ, cond_occ = dt.scan(doc)
    placeholders = sorted({o["key"] for o in occ if not o["key"].startswith(("IF_", "END_IF_"))})
    split = sorted({o["key"] for o in occ if o["runs"] > 1})
    parts = {}
    for o in occ:
        parts.setdefault(o["part"], set()).add(o["key"])

    # Resolve each {{IF_X}} block to include / exclude.
    conditions, sources = {}, {}
    opens = sorted({o["key"][3:] for o in cond_occ if o["key"].startswith("IF_")})
    closes = sorted({o["key"][7:] for o in cond_occ if o["key"].startswith("END_IF_")})
    for c in sorted(set(opens) ^ set(closes)):
        issues.append("conditional %s has an IF_ or END_IF_ marker without its partner" % c)
    for cond in opens:
        key, cands = dt.guess_condition_key(cond, data)
        if key is None:
            issues.append("no data key holds a readable yes/no flag for {{IF_%s}} (candidates: %s); set conditions.%s to true/false"
                          % (cond, ", ".join("%s=%r" % (k, data[k]) for k in cands) or "none", cond))
            conditions[cond] = False
            sources[cond] = {"key": None, "value": None, "candidates": cands}
            continue
        flag = dt.truthiness(data[key])
        sources[cond] = {"key": key, "value": data[key], "candidates": cands}
        if flag is None:
            issues.append("value %r of %s cannot be read as yes/no for {{IF_%s}}; set conditions.%s"
                          % (data[key], key, cond, cond))
            flag = False
        conditions[cond] = flag
        notes.append("{{IF_%s}} decided by %s=%r -> %s" % (cond, key, data[key], "include" if flag else "remove"))

    # Dry run: which placeholders survive the conditionals, and which have no value?
    dry = Document(template)
    try:
        dt.apply_conditionals(dry, conditions)
    except ValueError as e:
        issues.append(str(e))
    surviving = sorted({o["key"] for o in dt.scan(dry)[0]
                        if not o["key"].startswith(("IF_", "END_IF_"))})
    for k in surviving:
        if k not in values:
            toks = set(k.split("_"))
            alias = [d for d in data if d not in placeholders and toks & set(d.split("_")) - {"", "NAME", "DATE"}
                     or (d not in placeholders and (d in k or k in d))]
            issues.append("placeholder {{%s}} has no value in the data%s" % (
                k, "; likely the same field under another name: %s" % ", ".join("%s=%r" % (a, data[a]) for a in alias) if alias else ""))
        elif dt.to_text(values[k]).strip() == "":
            issues.append("placeholder {{%s}} has an empty value" % k)
    for o in occ:
        k = o["key"]
        if k in values and o["prefix_char"] in "$€£¥" and o["prefix_char"] and dt.to_text(values[k]).startswith(o["prefix_char"]):
            notes.append("{{%s}} follows a literal %s in the template and its value already has one; the fill step strips the duplicate"
                         % (k, o["prefix_char"]))
    unused = sorted(k for k in data if k not in placeholders and k not in [s["key"] for s in sources.values()])
    if unused:
        notes.append("data keys not used by the template (fine, informational): %s" % ", ".join(unused))
    if split:
        notes.append("placeholders split across runs (handled): %s" % ", ".join(split))

    dt.emit({
        "values": values,
        "conditions": conditions,
        "condition_sources": sources,
        "placeholders": placeholders,
        "split_placeholders": split,
        "placeholder_parts": {p: sorted(v) for p, v in parts.items()},
        "issues": issues,
        "notes": sorted(set(notes), key=notes.index),
        "inspect_status": "ready" if not issues else "needs_input",
    })


if __name__ == "__main__":
    try:
        main()
    except Exception as e:  # report as JSON, not a traceback
        dt.emit({"error": "%s: %s" % (type(e).__name__, e), "trace": traceback.format_exc(limit=3)})
        sys.exit(1)
