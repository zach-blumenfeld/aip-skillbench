#!/usr/bin/env python3
"""
Financial-model formatting helpers for openpyxl workbooks.

Encodes the color-coding and number-format conventions agents should apply
by default when building or extending a financial model. These are
DEFAULTS only — when editing an existing workbook that already has its own
conventions, study the existing styles and match them; the existing
template wins.

Usage from the calling skill:

    from financial_format import apply_role, NUMBER_FORMATS

    apply_role(ws["B5"], "input",   number_format="currency",  highlight=True)
    apply_role(ws["B6"], "input",   number_format="percent",   highlight=True)
    apply_role(ws["C5"], "formula", number_format="currency")
    apply_role(ws["D5"], "cross_sheet", number_format="multiple")
"""

from openpyxl.styles import Font, PatternFill

FINANCIAL_COLORS = {
    "input":       "0000FF",  # blue   — hardcoded inputs, scenario knobs
    "formula":     "000000",  # black  — all calculated cells
    "cross_sheet": "008000",  # green  — links to other sheets in same workbook
    "external":    "FF0000",  # red    — links to other files
}

ASSUMPTION_HIGHLIGHT_FILL = "FFFF00"  # yellow background — flag cells needing review

NUMBER_FORMATS = {
    "currency":         "$#,##0;($#,##0);-",
    "currency_decimal": "$#,##0.00;($#,##0.00);-",
    "percent":          "0.0%;(0.0%);-",
    "multiple":         "0.0x",
    "year":             "@",                   # text format; keeps "2024" not "2,024"
    "integer":          "#,##0;(#,##0);-",
}

VALID_ROLES = tuple(FINANCIAL_COLORS.keys())


def font_for(role):
    if role not in FINANCIAL_COLORS:
        raise ValueError(f"role must be one of {VALID_ROLES}, got {role!r}")
    return Font(color=FINANCIAL_COLORS[role])


def highlight_fill():
    return PatternFill("solid", start_color=ASSUMPTION_HIGHLIGHT_FILL)


def apply_role(cell, role, number_format=None, highlight=False):
    """
    Apply convention to a single openpyxl cell.

    role: one of "input" | "formula" | "cross_sheet" | "external"
    number_format: optional key from NUMBER_FORMATS, or a raw format string
    highlight: True to set yellow assumption-highlight background
    """
    cell.font = font_for(role)
    if number_format is not None:
        cell.number_format = NUMBER_FORMATS.get(number_format, number_format)
    if highlight:
        cell.fill = highlight_fill()


def apply_range(ws, cell_range, role, number_format=None, highlight=False):
    """Same as apply_role, broadcast over a string range like 'B5:B12'."""
    for row in ws[cell_range]:
        for cell in row:
            apply_role(cell, role, number_format=number_format, highlight=highlight)


if __name__ == "__main__":
    import json
    import sys
    print(json.dumps({
        "colors": FINANCIAL_COLORS,
        "highlight_fill": ASSUMPTION_HIGHLIGHT_FILL,
        "number_formats": NUMBER_FORMATS,
    }, indent=2))
    sys.exit(0)
