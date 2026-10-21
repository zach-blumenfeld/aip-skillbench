#!/usr/bin/env python3
"""Receipt OCR -> date + total -> single-sheet xlsx.

Requires: Python >= 3.8, pytesseract, Pillow, openpyxl, and the `tesseract` binary on PATH.

Subcommands:
  check                                   verify dependencies (exit 1 and name what is missing)
  extract <img_dir> <out.json> [--workers N]
                                          multi-pass OCR every image, apply the date / total rules,
                                          write per-file results + candidates + review flags
  write <extract.json> <out.xlsx> [--overrides overrides.json]
                                          write sheet "results" (filename, date, total_amount),
                                          applying verified corrections, then validate the file
  validate <out.xlsx> [--img-dir DIR]     re-check an existing output file against the spec
"""
import argparse
import datetime as dt
import json
import os
import re
import sys
from concurrent.futures import ThreadPoolExecutor

IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".webp", ".tif", ".tiff", ".bmp"}
HEADER = ["filename", "date", "total_amount"]
SHEET = "results"

# ---- total-amount rules (task hint) -------------------------------------------------------
# Tiers, most specific first. Matched against the line upper-cased with every run of
# non-alphanumerics collapsed to one space, so "TOTAL: RM" and "TOTAL (RM):" both read "TOTAL RM".
TIERS = [
    ("GRAND TOTAL", ["GRAND TOTAL", "GRANDTOTAL"]),
    ("TOTAL RM", ["TOTAL RM"]),
    ("TOTAL AMOUNT", ["TOTAL AMOUNT"]),
    ("TOTAL", ["TOTAL DUE", "AMOUNT DUE", "BALANCE DUE", "NETT TOTAL", "NET TOTAL", "TOTAL", "AMOUNT"]),
]
EXCLUDE = ["SUBTOTAL", "SUB TOTAL", "TAX", "GST", "SST", "DISCOUNT", "CHANGE", "CASH TENDERED",
           # pre-tax totals ("Total (Excluding GST)") survive the tax words when OCR garbles "GST"
           "EXCLUDING", "EXCL"]
TAX_WORDS = ["TAX", "GST", "SST"]
# A rounding-adjusted line ("Total After Adj Incl GST 17.70", "Total Rounded") is the amount paid,
# so the tax words alone do not exclude it (the other exclusions still do).
FINAL_WORDS = ["ROUND", "AFTER ADJ"]
# Column-header lines ("Description Qty Price Amount") carry the AMOUNT/TOTAL keyword but no total.
HEADER_WORDS = ["QTY", "PRICE", "DESCRIPTION", "DESC", "ITEM", "U P", "UNIT"]

MONEY_STRICT = re.compile(r"(?<![\d.,])(\d{1,3}(?:,\d{3})+|\d+)\.(\d{2})(?![\d])")
# OCR-tolerant: "85. 54", "92:80", "5,57", "538 00" (decimal point misread or lost).
# Used only when the strict pattern finds nothing on the line.
MONEY_LOOSE = re.compile(r"(?<![\d.,:])(\d{1,6})(?:\s?[.,:]\s?|\s)(\d{2})(?![\d.,:])")


def norm(line):
    return re.sub(r"[^A-Z0-9]+", " ", line.upper()).strip()


def money_values(line):
    vals = [(m.group(1).replace(",", "") + "." + m.group(2), False) for m in MONEY_STRICT.finditer(line)]
    if vals:
        return vals
    return [(f"{int(m.group(1))}.{m.group(2)}", True) for m in MONEY_LOOSE.finditer(line)]


def fmt_amount(s):
    return f"{float(s):.2f}"


def excluded(n):
    hits = [x for x in EXCLUDE if x in n]
    if "SUBTOTAL" in n.replace(" ", ""):
        hits.append("SUBTOTAL")
    if hits and all(h in TAX_WORDS for h in hits) and any(w in n for w in FINAL_WORDS):
        return False
    return bool(hits)


def find_total(text):
    """Apply the hint rules. Returns (value|None, info dict, list of candidate dicts)."""
    lines = [l for l in text.splitlines() if l.strip()]
    cands = []
    for i, raw in enumerate(lines):
        n = norm(raw)
        if excluded(n):
            continue
        tier = None
        for t_idx, (_, kws) in enumerate(TIERS):
            if any(k in n for k in kws):
                tier = t_idx
                break
        if tier is None:
            continue
        if tier == 3 and any(re.search(rf"\b{h}\b", n) for h in HEADER_WORDS):
            continue
        vals = money_values(raw)
        source, loose = "same_line", False
        if vals:
            value, loose = vals[-1]
        else:
            # Keyword and amount split across two lines: take the last number on the next line.
            # Not when digits follow the keyword: the amount is on this line, just unreadable.
            kw_end = max(n.rfind(k) + len(k) for k in TIERS[tier][1] if k in n)
            if i + 1 >= len(lines) or re.search(r"\d", n[kw_end:]):
                continue
            nxt = lines[i + 1]
            nv = money_values(nxt)
            if not nv or excluded(norm(nxt)):
                continue
            value, loose = nv[-1]
            source = "next_line"
        cands.append({"tier": TIERS[tier][0], "tier_rank": tier, "line_no": i, "line": raw.strip(),
                      "value": fmt_amount(value), "source": source, "ocr_loose": loose,
                      "rounded": "ROUND" in n})
    if not cands:
        return None, {}, []
    best = min(c["tier_rank"] for c in cands)
    tier_c = [c for c in cands if c["tier_rank"] == best]
    # Within a tier: the rounded total is what was paid; otherwise the first (topmost) match,
    # because tax-summary "Total" rows sit below the real total.
    chosen = next((c for c in tier_c if c["rounded"]), tier_c[0])
    return chosen["value"], chosen, cands


# ---- date rules ---------------------------------------------------------------------------
MONTHS = {m: i + 1 for i, m in enumerate(
    ["JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC"])}
SEP = r"[/\-.]"
RE_DMY = re.compile(rf"(?<!\d)(\d{{1,2}})\s?{SEP}\s?(\d{{1,2}})\s?{SEP}\s?(\d{{4}}|\d{{2}})(?!\d)")
RE_YMD = re.compile(rf"(?<!\d)(\d{{4}})\s?{SEP}\s?(\d{{1,2}})\s?{SEP}\s?(\d{{1,2}})(?!\d)")
RE_DMONY = re.compile(r"(?<!\d)(\d{1,2})\s?[-/ .]?\s?(JAN|FEB|MAR|APR|MAY|JUN|JUL|AUG|SEP|OCT|NOV|DEC)[A-Z]*\.?\s?[-/ ,.]?\s?(\d{4}|\d{2})(?!\d)")
RE_MONDY = re.compile(r"\b(JAN|FEB|MAR|APR|MAY|JUN|JUL|AUG|SEP|OCT|NOV|DEC)[A-Z]*\.?\s?(\d{1,2}),?\s?(\d{4})(?!\d)")
YEAR_MIN, YEAR_MAX = 2000, dt.date.today().year + 1


def overlapping(rx, s):
    """All matches starting at every position ("#19- 18/01/2018" must still yield 18/01/2018)."""
    seen = set()
    for start in range(len(s)):
        m = rx.match(s, start)
        if m and m.span() not in seen:
            seen.add(m.span())
            yield m


def mk_date(y, m, d):
    y = int(y)
    if y < 100:
        y += 2000
    if not (YEAR_MIN <= y <= YEAR_MAX):
        return None
    try:
        return dt.date(y, int(m), int(d)).isoformat()
    except ValueError:
        return None


def find_date(text):
    """Returns (iso|None, info dict, candidates). Day-first unless only month-first is valid."""
    cands = []
    for i, raw in enumerate(l for l in text.splitlines() if l.strip()):
        up = raw.upper()
        labelled = bool(re.search(r"\bDATE\b|\bDATA\b|\bTARIKH\b", up)) and "PRINT" not in up
        for m in overlapping(RE_YMD, up):
            iso = mk_date(m.group(1), m.group(2), m.group(3))
            if iso:
                cands.append({"line_no": i, "line": raw.strip(), "date": iso, "order": "YMD", "labelled": labelled})
        for m in overlapping(RE_DMY, up):
            a, b, y = m.group(1), m.group(2), m.group(3)
            iso, order = mk_date(y, b, a), "DMY"
            if not iso:
                iso, order = mk_date(y, a, b), "MDY"
            if iso:
                cands.append({"line_no": i, "line": raw.strip(), "date": iso, "order": order, "labelled": labelled})
        for m in overlapping(RE_DMONY, up):
            iso = mk_date(m.group(3), MONTHS[m.group(2)], m.group(1))
            if iso:
                cands.append({"line_no": i, "line": raw.strip(), "date": iso, "order": "D-MON-Y", "labelled": labelled})
        for m in overlapping(RE_MONDY, up):
            iso = mk_date(m.group(3), MONTHS[m.group(1)], m.group(2))
            if iso:
                cands.append({"line_no": i, "line": raw.strip(), "date": iso, "order": "MON-D-Y", "labelled": labelled})
    if not cands:
        return None, {}, []
    chosen = next((c for c in cands if c["labelled"]), cands[0])
    return chosen["date"], chosen, cands


# ---- OCR ----------------------------------------------------------------------------------
def ocr_passes(path):
    import pytesseract
    from PIL import Image, ImageOps

    img = Image.open(path)
    img = ImageOps.exif_transpose(img)
    if img.mode not in ("RGB", "L"):
        img = img.convert("RGB")
    gray = ImageOps.autocontrast(ImageOps.grayscale(img))
    up = gray
    if max(gray.size) < 2000:  # small scans: upscale 2x for character recognition
        up = gray.resize((gray.width * 2, gray.height * 2), Image.LANCZOS)
    passes = [("raw_psm3", img, ""), ("gray_psm4", gray, "--psm 4"), ("gray_up_psm6", up, "--psm 6")]
    out = []
    for name, im, cfg in passes:
        try:
            out.append((name, pytesseract.image_to_string(im, config=cfg)))
        except Exception as e:  # keep going; other passes may succeed
            out.append((name, ""))
            print(f"warn: {os.path.basename(path)} pass {name} failed: {e}", file=sys.stderr)
    return out


def vote(per_pass, key):
    """Primary = first pass with a value; agreement = how many passes produced that same value."""
    vals = [p[key] for p in per_pass if p[key]]
    if not vals:
        return None, 0
    primary = vals[0]
    counts = {v: vals.count(v) for v in vals}
    top = max(counts.values())
    # A value seen in >= 2 passes beats a lone primary.
    if counts[primary] < top and top >= 2:
        primary = max(counts, key=lambda v: (counts[v], -vals.index(v)))
    return primary, counts[primary]


def pool_total(per_pass):
    """Pool total candidates from every pass. Keep the most specific tier any pass reached; inside
    it prefer same-line amounts, then rounded totals, then cleanly read numbers (a garbled line such
    as "Total (Bxoluding ast): 21,30" slipped past the exclusions), then the value most passes agree
    on, then the topmost line. Returns (value|None, chosen candidate, number of passes that produced the value)."""
    pool = [dict(c, pass_idx=k) for k, p in enumerate(per_pass) for c in p["total_candidates"]]
    if not pool:
        return None, {}, 0
    best = min(c["tier_rank"] for c in pool)
    tier = [c for c in pool if c["tier_rank"] == best]
    passes_with = {v: len({c["pass_idx"] for c in tier if c["value"] == v}) for v in {c["value"] for c in tier}}
    chosen = min(tier, key=lambda c: (c["source"] != "same_line", not c["rounded"], c["ocr_loose"],
                                      -passes_with[c["value"]], c["line_no"], c["pass_idx"]))
    agree = sum(1 for p in per_pass if any(c["value"] == chosen["value"] for c in p["total_candidates"]))
    return chosen["value"], chosen, agree


def process(path):
    name = os.path.basename(path)
    texts = ocr_passes(path)
    per_pass = []
    for pname, text in texts:
        d, d_info, d_c = find_date(text)
        t, t_info, t_c = find_total(text)
        per_pass.append({"pass": pname, "date": d, "date_info": d_info, "total": t, "total_info": t_info,
                         "total_candidates": t_c, "date_candidates": d_c})
    date, d_agree = vote(per_pass, "date")
    total, t_info, t_agree = pool_total(per_pass)
    reasons = []
    if date is None:
        reasons.append("no date found")
    elif d_agree < 2:
        reasons.append("date found in only one OCR pass")
    if len({p["date"] for p in per_pass if p["date"]}) > 1:
        reasons.append("OCR passes disagree on date")
    if total is None:
        reasons.append("no total found")
    else:
        if t_agree < 2:
            reasons.append("total found in only one OCR pass")
        if t_info.get("ocr_loose"):
            reasons.append("total read with the OCR-tolerant number pattern")
        if t_info.get("source") == "next_line":
            reasons.append("total taken from the line after the keyword")
        if len({c["value"] for p in per_pass for c in p["total_candidates"]
                if c["tier_rank"] == t_info["tier_rank"]}) > 1:
            reasons.append("conflicting values in the winning keyword tier")
        if t_info["tier_rank"] == 3 and len(texts) and total in {c["value"] for p in per_pass
                                                                    for c in p["total_candidates"]
                                                                    if c["tier_rank"] == 3 and
                                                                    len(money_values(c["line"])) > 1}:
            reasons.append("total line holds several amounts (possible tax-summary row)")
    return {
        "filename": name,
        "date": date,
        "total_amount": total,
        "total_line": t_info.get("line"),
        "total_tier": t_info.get("tier"),
        "needs_review": bool(reasons),
        "review_reasons": reasons,
        "passes": [{"pass": p["pass"], "date": p["date"], "total": p["total"],
                    "date_line": p["date_info"].get("line"),
                    "total_candidates": [f'[{c["tier"]}] {c["value"]} <- {c["line"]}' for c in p["total_candidates"]],
                    "date_candidates": sorted({c["date"] for c in p["date_candidates"]})}
                   for p in per_pass],
        "ocr_text_primary": texts[0][1] if texts else "",
    }


def list_images(img_dir):
    return sorted(f for f in os.listdir(img_dir)
                  if os.path.splitext(f)[1].lower() in IMAGE_EXTS and not f.startswith("."))


def cmd_check(_):
    missing = []
    for mod, pip in [("pytesseract", "pytesseract"), ("PIL", "Pillow"), ("openpyxl", "openpyxl")]:
        try:
            __import__(mod)
        except ImportError:
            missing.append(f"python package {pip}")
    if not missing:
        import pytesseract
        try:
            print("tesseract", pytesseract.get_tesseract_version())
        except Exception:
            missing.append("tesseract binary (apt: tesseract-ocr)")
    if sys.version_info < (3, 8):
        missing.append("Python >= 3.8")
    if missing:
        print("MISSING: " + "; ".join(missing))
        sys.exit(1)
    print("OK: all dependencies present")


def cmd_extract(a):
    if not os.path.isdir(a.img_dir):
        sys.exit(f"error: image directory not found: {a.img_dir}")
    files = list_images(a.img_dir)
    if not files:
        sys.exit(f"error: no image files ({', '.join(sorted(IMAGE_EXTS))}) in {a.img_dir}")
    os.environ.setdefault("OMP_THREAD_LIMIT", "1")  # we parallelise across images instead
    workers = a.workers or max(1, min(len(files), os.cpu_count() or 2))
    with ThreadPoolExecutor(workers) as ex:
        results = list(ex.map(process, [os.path.join(a.img_dir, f) for f in files]))
    with open(a.out_json, "w") as f:
        json.dump({"img_dir": os.path.abspath(a.img_dir), "results": results}, f, indent=2)
    print(f"{'filename':<16}{'date':<12}{'total':>10}  review")
    for r in results:
        print(f"{r['filename']:<16}{str(r['date']):<12}{str(r['total_amount']):>10}  "
              f"{'; '.join(r['review_reasons']) if r['needs_review'] else '-'}")
    n = sum(r["needs_review"] for r in results)
    print(f"\n{len(results)} images, {n} flagged for review -> {a.out_json}")


DATE_OK = re.compile(r"^\d{4}-\d{2}-\d{2}$")
AMT_OK = re.compile(r"^\d+\.\d{2}$")


def clean_override(fn, field, v):
    if v is None or (isinstance(v, str) and v.strip().lower() in ("", "null", "none")):
        return None
    v = str(v).strip()
    if field == "date":
        if not DATE_OK.match(v) or not mk_date(*v.split("-")):
            sys.exit(f"error: override date for {fn} must be a real YYYY-MM-DD date or null, got {v!r}")
        return v
    v = v.replace(",", "").replace("RM", "").replace("$", "").strip()
    try:
        return fmt_amount(v)
    except ValueError:
        sys.exit(f"error: override total_amount for {fn} must be a number or null, got {v!r}")


def cmd_write(a):
    from openpyxl import Workbook

    data = json.load(open(a.extract_json))
    rows = {r["filename"]: {"date": r["date"], "total_amount": r["total_amount"]} for r in data["results"]}
    if a.overrides:
        ov = json.load(open(a.overrides))
        for fn, fields in ov.items():
            if fn not in rows:
                sys.exit(f"error: override for unknown file {fn!r}; known: {', '.join(sorted(rows))}")
            for field, v in fields.items():
                if field not in ("date", "total_amount"):
                    sys.exit(f"error: override field must be date or total_amount, got {field!r} for {fn}")
                rows[fn][field] = clean_override(fn, field, v)
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append(HEADER)
    for fn in sorted(rows):
        ws.append([fn, rows[fn]["date"], rows[fn]["total_amount"]])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if c.value is not None:
                c.data_type = "s"
                c.number_format = "@"
    os.makedirs(os.path.dirname(os.path.abspath(a.out_xlsx)), exist_ok=True)
    wb.save(a.out_xlsx)
    print(f"wrote {len(rows)} rows -> {a.out_xlsx}")
    a.img_dir = data.get("img_dir")
    cmd_validate(a, path=a.out_xlsx)


def cmd_validate(a, path=None):
    from openpyxl import load_workbook

    path = path or a.out_xlsx
    errs = []
    wb = load_workbook(path)
    if wb.sheetnames != [SHEET]:
        errs.append(f"sheets must be exactly ['{SHEET}'], got {wb.sheetnames}")
    ws = wb[wb.sheetnames[0]]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    if not rows or rows[0][:3] != HEADER or ws.max_column != 3:
        errs.append(f"header must be {HEADER} with exactly 3 columns, got {rows[0] if rows else None}")
    body = rows[1:]
    names = [r[0] for r in body]
    if names != sorted(names):
        errs.append("rows are not ordered by filename")
    img_dir = getattr(a, "img_dir", None)
    if img_dir and os.path.isdir(img_dir) and names != list_images(img_dir):
        errs.append("filenames do not match the image directory one-to-one")
    for r in body:
        if any(v is None for v in r[:1]) or len(r) != 3:
            errs.append(f"bad row {r}")
            continue
        if r[1] is not None and not (isinstance(r[1], str) and DATE_OK.match(r[1])):
            errs.append(f"{r[0]}: date must be a 'YYYY-MM-DD' string or empty, got {r[1]!r}")
        if r[2] is not None and not (isinstance(r[2], str) and AMT_OK.match(r[2])):
            errs.append(f"{r[0]}: total_amount must be a string with 2 decimals or empty, got {r[2]!r}")
    nulls = sum(1 for r in body for v in r[1:] if v is None)
    if errs:
        print("VALIDATION FAILED:\n  " + "\n  ".join(errs))
        sys.exit(1)
    print(f"VALID: 1 sheet '{SHEET}', header + {len(body)} rows, {nulls} null cells")


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest="cmd", required=True)
    sub.add_parser("check").set_defaults(fn=cmd_check)
    e = sub.add_parser("extract")
    e.add_argument("img_dir")
    e.add_argument("out_json")
    e.add_argument("--workers", type=int, default=0)
    e.set_defaults(fn=cmd_extract)
    w = sub.add_parser("write")
    w.add_argument("extract_json")
    w.add_argument("out_xlsx")
    w.add_argument("--overrides")
    w.set_defaults(fn=cmd_write)
    v = sub.add_parser("validate")
    v.add_argument("out_xlsx")
    v.add_argument("--img-dir")
    v.set_defaults(fn=cmd_validate)
    a = p.parse_args()
    a.fn(a)


if __name__ == "__main__":
    main()
