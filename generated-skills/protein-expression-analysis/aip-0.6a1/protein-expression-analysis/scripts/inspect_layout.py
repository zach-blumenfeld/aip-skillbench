#!/usr/bin/env python3
"""Detect the layout of a protein-expression workbook (Task + Data sheets).

stdin:  {"state": {"workbook_path": str}, "assets": {}, "expects": [...]}
stdout: {"detected_layout": {...}, "layout_warnings": [...]}

Finds, without assuming fixed cell addresses: the Data sheet's ID / gene / sample
header columns, the Task sheet's lookup block (protein rows x sample columns plus
the Control/Treated group row), the statistics block (Control Mean / StdDev,
Treated Mean / StdDev) and the fold-change block (Fold Change, Log2 FC). Every
target cell is listed explicitly so fill_workbook.py writes exactly those cells.
"""
import json
import re
import sys

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

STAT_KEYS = {
    "controlmean": "control_mean",
    "controlaverage": "control_mean",
    "controlstddev": "control_sd",
    "controlstdev": "control_sd",
    "controlsd": "control_sd",
    "treatedmean": "treated_mean",
    "treatedaverage": "treated_mean",
    "treatedstddev": "treated_sd",
    "treatedstdev": "treated_sd",
    "treatedsd": "treated_sd",
}


def norm(v):
    return re.sub(r"[^a-z0-9]", "", str(v).lower()) if v is not None else ""


def is_missing(v):
    return v is None or (isinstance(v, str) and v.strip() == "")


def fill_rgb(cell):
    f = cell.fill
    if f is None or f.fill_type is None or f.fgColor is None or f.fgColor.type != "rgb":
        return None
    return str(f.fgColor.rgb)[-6:].upper()


def group_label(v):
    n = norm(v)
    if "control" in n or n in ("ctrl", "ctl", "vehicle", "untreated"):
        return "Control"
    if "treat" in n:
        return "Treated"
    return None


def pick_sheets(wb, warnings):
    names = {s.lower(): s for s in wb.sheetnames}
    data = names.get("data")
    task = names.get("task")
    if data is None:
        data = max(wb.sheetnames, key=lambda s: wb[s].max_row)
        warnings.append(f"No sheet named 'Data'; using '{data}' (most rows) as the data sheet.")
    if task is None:
        others = [s for s in wb.sheetnames if s != data]
        task = others[0] if others else data
        warnings.append(f"No sheet named 'Task'; using '{task}' as the task sheet.")
    return task, data


def find_header(ws, label, max_rows=10):
    for row in ws.iter_rows(min_row=1, max_row=max_rows):
        for c in row:
            if norm(c.value) == label:
                return c.row, c.column
    return None


def inspect(path):
    warnings = []
    wb = load_workbook(path)
    task_name, data_name = pick_sheets(wb, warnings)
    ws, ds = wb[task_name], wb[data_name]

    # ---- Data sheet --------------------------------------------------------
    hit = find_header(ds, "proteinid")
    if hit is None:
        raise SystemExit(f"Data sheet '{data_name}' has no 'Protein_ID' header in its first 10 rows.")
    d_hdr_row, d_id_col = hit
    d_gene_col = None
    headers = {}
    for c in ds[d_hdr_row]:
        if c.value is None:
            continue
        if norm(c.value) in ("genesymbol", "gene"):
            d_gene_col = c.column
        headers[c.column] = str(c.value)
    d_last_col = max(headers)
    d_last_row = d_hdr_row
    for r in range(d_hdr_row + 1, ds.max_row + 1):
        if not is_missing(ds.cell(r, d_id_col).value):
            d_last_row = r
    ids = {}
    genes = {}
    for r in range(d_hdr_row + 1, d_last_row + 1):
        pid = ds.cell(r, d_id_col).value
        if not is_missing(pid):
            key = str(pid).strip()
            if key in ids:
                warnings.append(f"Data protein ID '{key}' appears more than once (rows {ids[key]} and {r}); MATCH uses the first.")
            else:
                ids[key] = r
        if d_gene_col:
            g = ds.cell(r, d_gene_col).value
            if not is_missing(g):
                genes.setdefault(str(g).strip(), []).append(r)
    dup_headers = sorted({h for h in headers.values() if list(headers.values()).count(h) > 1})
    if dup_headers:
        warnings.append(f"Duplicate Data column headers (MATCH uses the first): {dup_headers}")

    # ---- Task lookup block -------------------------------------------------
    t_hdr = None
    for row in ws.iter_rows():
        for c in row:
            if c.column <= 3 and norm(c.value) == "proteinid":
                t_hdr = (c.row, c.column)
                break
        if t_hdr:
            break
    if t_hdr is None:
        raise SystemExit(f"Task sheet '{task_name}' has no 'Protein_ID' header.")
    hdr_row, t_id_col = t_hdr
    t_gene_col = t_id_col + 1 if norm(ws.cell(hdr_row, t_id_col + 1).value) in ("genesymbol", "gene") else None
    first_sample_col = (t_gene_col or t_id_col) + 1
    sample_cols = []
    col = first_sample_col
    while not is_missing(ws.cell(hdr_row, col).value):
        sample_cols.append(col)
        col += 1
    if not sample_cols:
        raise SystemExit(f"No sample names found to the right of the Task header row {hdr_row}.")

    target_rows = []
    r = hdr_row + 1
    while not is_missing(ws.cell(r, t_id_col).value):
        target_rows.append(r)
        r += 1
    if not target_rows:
        raise SystemExit("No target protein IDs found below the Task header row.")

    group_row = None
    groups = {}
    for gr in range(hdr_row - 1, max(0, hdr_row - 4), -1):
        labels = {get_column_letter(c): group_label(ws.cell(gr, c).value) for c in sample_cols}
        if any(labels.values()):
            group_row, groups = gr, labels
            break
    if group_row is None:
        warnings.append("No Control/Treated group row found above the sample header; set groups by override.")
    else:
        unk = [k for k, v in groups.items() if v is None]
        if unk:
            warnings.append(f"Group row {group_row} has no Control/Treated label for columns {unk}; they are left out of the statistics.")
    n_ctrl = sum(1 for v in groups.values() if v == "Control")
    n_trt = sum(1 for v in groups.values() if v == "Treated")
    if n_ctrl < 2 or n_trt < 2:
        warnings.append(f"Groups have {n_ctrl} Control and {n_trt} Treated samples; a StdDev needs at least 2 per group.")

    # resolve sample names against Data headers
    sample_res = {}
    hdr_by_text = {}
    for cidx, h in headers.items():
        hdr_by_text.setdefault(h, cidx)
    for c in sample_cols:
        name = str(ws.cell(hdr_row, c).value)
        L = get_column_letter(c)
        if name in hdr_by_text:
            sample_res[L] = {"name": name, "by": "exact", "data_header": name}
            continue
        if name.strip() in hdr_by_text:
            sample_res[L] = {"name": name, "by": "trim", "data_header": name.strip()}
            continue
        cont = [h for h in hdr_by_text if name.strip().lower() in h.lower()]
        if len(cont) == 1:
            sample_res[L] = {"name": name, "by": "contains", "data_header": cont[0]}
            continue
        sample_res[L] = {"name": name, "by": None, "data_header": None,
                         "candidates": cont[:5]}
        warnings.append(f"Sample '{name}' (column {L}) has {'several' if cont else 'no'} matching Data headers{': ' + str(cont[:5]) if cont else ''}.")

    # resolve proteins
    protein_res = {}
    for r in target_rows:
        pid = str(ws.cell(r, t_id_col).value).strip()
        gene = ws.cell(r, t_gene_col).value if t_gene_col else None
        if pid in ids:
            protein_res[str(r)] = {"id": pid, "gene": gene, "by": "id", "data_row": ids[pid]}
        elif gene is not None and str(gene).strip() in genes and len(genes[str(gene).strip()]) == 1:
            protein_res[str(r)] = {"id": pid, "gene": gene, "by": "gene", "data_row": genes[str(gene).strip()][0]}
            warnings.append(f"Protein '{pid}' (row {r}) not found by ID; matched by gene symbol '{gene}'.")
        else:
            protein_res[str(r)] = {"id": pid, "gene": gene, "by": None, "data_row": None}
            warnings.append(f"Protein '{pid}' (row {r}) not found in Data by ID or unique gene symbol.")

    # ---- statistics block ----------------------------------------------------
    stat_hits = {}
    for row in ws.iter_rows(min_row=hdr_row + 1):
        for c in row:
            k = STAT_KEYS.get(norm(c.value))
            if k and k not in stat_hits:
                stat_hits[k] = (c.row, c.column)
    n = len(target_rows)
    stats_cells = {}
    stats_orientation = None
    if len(stat_hits) == 4:
        cols = {v[1] for v in stat_hits.values()}
        rows = {v[0] for v in stat_hits.values()}
        if len(cols) == 1:
            stats_orientation = "proteins_across_columns"
            for k, (sr, sc) in stat_hits.items():
                stats_cells[k] = [f"{get_column_letter(sc + 1 + i)}{sr}" for i in range(n)]
        elif len(rows) == 1:
            stats_orientation = "proteins_down_rows"
            for k, (sr, sc) in stat_hits.items():
                stats_cells[k] = [f"{get_column_letter(sc)}{sr + 1 + i}" for i in range(n)]
        else:
            warnings.append("Statistics labels are neither in one column nor in one row; set stats_cells by override.")
    else:
        warnings.append(f"Found statistics labels {sorted(stat_hits)}; expected control_mean, control_sd, treated_mean, treated_sd.")

    # ---- fold-change block ---------------------------------------------------
    fc_cells = {}
    fc_hdr = None
    for row in ws.iter_rows(min_row=hdr_row + 1):
        fc_col = log_col = None
        for c in row:
            nv = norm(c.value)
            if nv in ("foldchange", "fc"):
                fc_col = c.column
            elif nv.startswith("log2") or nv in ("logfc", "log2foldchange"):
                log_col = c.column
        if fc_col and log_col:
            fc_hdr = (row[0].row, fc_col, log_col)
            break
    fc_protein_order = []
    if fc_hdr:
        fr, fcc, lgc = fc_hdr
        fc_cells = {
            "fold_change": [f"{get_column_letter(fcc)}{fr + 1 + i}" for i in range(n)],
            "log2_fc": [f"{get_column_letter(lgc)}{fr + 1 + i}" for i in range(n)],
        }
        listed = [ws.cell(fr + 1 + i, t_id_col).value for i in range(n)]
        if any(not is_missing(v) for v in listed):
            fc_protein_order = [None if is_missing(v) else str(v).strip() for v in listed]
            expect = [protein_res[str(r)]["id"] for r in target_rows]
            if fc_protein_order != expect:
                warnings.append(f"Fold-change block lists proteins {fc_protein_order}, not the lookup order {expect}; fix fc_cells by override.")
    else:
        warnings.append("No fold-change header row (Fold Change + Log2 FC) found; set fc_cells by override.")

    # ---- yellow-cell coverage --------------------------------------------------
    targets = {f"{get_column_letter(c)}{r}" for r in target_rows for c in sample_cols}
    for lst in list(stats_cells.values()) + list(fc_cells.values()):
        targets.update(lst)
    fills = {}
    for row in ws.iter_rows():
        for c in row:
            rgb = fill_rgb(c)
            if rgb:
                fills[c.coordinate] = rgb
    yellow = {k for k, v in fills.items() if v == "FFFF00"}
    not_yellow = sorted(targets - yellow)
    uncovered = sorted(yellow - targets)
    prefilled = sorted(t for t in targets if not is_missing(ws[t].value))
    if not_yellow and yellow:
        warnings.append(f"Target cells without the yellow input fill: {not_yellow[:20]}")
    if uncovered:
        warnings.append(f"Yellow cells not covered by the detected targets: {uncovered[:20]}")
    if prefilled:
        warnings.append(f"Target cells already hold values and will be overwritten: {prefilled[:20]}")

    # ---- missing-value profile ---------------------------------------------------
    missing = {}
    for r in target_rows:
        pr = protein_res[str(r)]
        if pr["data_row"] is None:
            continue
        cnt = {"Control": [0, 0], "Treated": [0, 0]}
        for c in sample_cols:
            L = get_column_letter(c)
            g = groups.get(L)
            sr = sample_res[L]
            if g is None or sr["data_header"] is None:
                continue
            v = ds.cell(pr["data_row"], hdr_by_text[sr["data_header"]]).value
            cnt[g][0] += 1
            if is_missing(v) or not isinstance(v, (int, float)):
                cnt[g][1] += 1
        missing[pr["id"]] = {"control_missing": cnt["Control"][1], "control_total": cnt["Control"][0],
                             "treated_missing": cnt["Treated"][1], "treated_total": cnt["Treated"][0]}
    n_missing = sum(m["control_missing"] + m["treated_missing"] for m in missing.values())
    if n_missing:
        sparse = [p for p, m in missing.items()
                  if m["control_total"] - m["control_missing"] < 2 or m["treated_total"] - m["treated_missing"] < 2]
        warnings.append(f"{n_missing} looked-up Data cells are blank (missing measurements)."
                        + (f" Proteins with <2 observed values in a group (StdDev undefined): {sparse}" if sparse else ""))

    layout = {
        "task_sheet": task_name,
        "data_sheet": data_name,
        "data": {
            "header_row": d_hdr_row,
            "id_col": get_column_letter(d_id_col),
            "gene_col": get_column_letter(d_gene_col) if d_gene_col else None,
            "first_row": d_hdr_row + 1,
            "last_row": d_last_row,
            "last_col": get_column_letter(d_last_col),
            "n_proteins": len(ids),
            "n_columns": len(headers),
        },
        "lookup": {
            "header_row": hdr_row,
            "group_row": group_row,
            "id_col": get_column_letter(t_id_col),
            "gene_col": get_column_letter(t_gene_col) if t_gene_col else None,
            "sample_cols": [get_column_letter(c) for c in sample_cols],
            "target_rows": target_rows,
        },
        "groups": groups,
        "sample_resolution": sample_res,
        "protein_resolution": protein_res,
        "stats_orientation": stats_orientation,
        "stats_cells": stats_cells,
        "fc_cells": fc_cells,
        "missing_profile": missing,
    }
    return layout, warnings


def main():
    payload = json.load(sys.stdin)
    state = payload.get("state", payload)
    layout, warnings = inspect(state["workbook_path"])
    print(json.dumps({"detected_layout": layout, "layout_warnings": warnings}))


if __name__ == "__main__":
    main()
