#!/usr/bin/env python3
"""Recalculate the output workbook with LibreOffice (recalc.py), then check it:
zero formula errors, formulas still present, cached values equal the shadow values.

stdin currentState keys: output_path, layout, expected (cell -> value).
stdout: {"verify_status": ok|errors_found|mismatch|recalc_failed|recalc_unavailable, "verify_report": {...}}
"""
import json
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile

from openpyxl import load_workbook

ERRS = ("#VALUE!", "#DIV/0!", "#REF!", "#NAME?", "#NULL!", "#NUM!", "#N/A")


def soffice():
    return shutil.which("soffice") or shutil.which("libreoffice")


def convert_recalc(path):
    """Fallback: open and re-save through LibreOffice; cells without cached values are computed on load."""
    with tempfile.TemporaryDirectory() as td:
        r = subprocess.run([soffice(), "--headless", "--norestore", "--calc", "--convert-to", "xlsx", "--outdir", td, path],
                           capture_output=True, text=True, timeout=180)
        conv = os.path.join(td, os.path.basename(path))
        if r.returncode != 0 or not os.path.exists(conv):
            return f"convert-to failed: {r.stderr.strip() or r.stdout.strip()}"
        shutil.copyfile(conv, path)
    return None


def same(a, b):
    if b is None:
        return a in (None, "")
    if isinstance(b, (int, float)):
        return isinstance(a, (int, float)) and math.isclose(a, b, rel_tol=1e-6, abs_tol=1e-9)
    return str(a) == str(b)


def main():
    st = json.load(sys.stdin)["currentState"]
    path, lay, expected = st["output_path"], st["layout"], st["expected"]
    report = {"recalc": None}

    if not soffice():
        report["recalc"] = "LibreOffice (soffice) not on PATH; formulas have no cached values yet"
        ws_f = load_workbook(path)[lay["task_sheet"]]
        fs = {k: ws_f[k].value for k in expected}
        report["formulas_lost"] = [k for k, f in fs.items() if not (isinstance(f, str) and f.startswith("="))][:30]
        report["static_issues"] = [f"{k}: modern function without _xlfn. prefix" for k, f in fs.items()
                                   if isinstance(f, str) and re.search(r"(?<!_xlfn\.)\b(STDEV|VAR|MODE|RANK)\.[A-Z]", f)][:30]
        print(json.dumps({"verify_status": "recalc_unavailable", "verify_report": report}))
        return

    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    from recalc import recalc

    try:
        res = recalc(path, 90)
    except Exception as e:
        res = {"error": str(e)}
    report["recalc"] = res
    if "error" in res:
        err = convert_recalc(path)
        report["fallback"] = err or "recalculated via soffice --convert-to"
        if err:
            print(json.dumps({"verify_status": "recalc_failed", "verify_report": report}))
            return

    def check():
        wbv, wbf = load_workbook(path, data_only=True), load_workbook(path)
        errors = [f"{ws.title}!{c.coordinate}={c.value}" for ws in wbv.worksheets for row in ws.iter_rows()
                  for c in row if isinstance(c.value, str) and any(e in c.value for e in ERRS)]
        ws_v, ws_f = wbv[lay["task_sheet"]], wbf[lay["task_sheet"]]
        lost = [k for k in expected if not (isinstance(ws_f[k].value, str) and ws_f[k].value.startswith("="))]
        mism = [{"cell": k, "cached": ws_v[k].value, "expected": v} for k, v in expected.items() if not same(ws_v[k].value, v)]
        return errors, lost, mism, [m["cell"] for m in mism if m["cached"] is None]

    errors, lost, mism, uncached = check()
    if uncached and "fallback" not in report:
        # recalc.py can exit cleanly without the macro having run; retry through a plain LibreOffice re-save.
        err = convert_recalc(path)
        report["fallback"] = err or "cells were uncached after recalc.py; recalculated via soffice --convert-to"
        if not err:
            errors, lost, mism, uncached = check()
    report.update({"formula_errors": errors[:30], "n_formula_errors": len(errors),
                   "formulas_lost": lost[:30], "mismatches": mism[:30], "n_mismatches": len(mism),
                   "uncached_cells": uncached[:30], "cells_checked": len(expected)})
    status = "errors_found" if errors else "mismatch" if (mism or lost) else "ok"
    print(json.dumps({"verify_status": status, "verify_report": report}, default=str))


if __name__ == "__main__":
    main()
