#!/usr/bin/env python3
"""Write the scoring sheet as live Excel formulas, cache values, recalculate with
LibreOffice (scripts/recalc.py) and verify every cell against the Python reference.
stdin: {"currentState": {"plan": {...}, ...}}; stdout: JSON with verify_status."""

import copy
import json
import math
import os
import re
import shutil
import sys
import zipfile
from xml.sax.saxutils import escape

import pl_lib as P

P.ensure_openpyxl()

from openpyxl import load_workbook  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402


OUT = {"path": None}


def fail(errors, **extra):
    print(json.dumps({"output_path": OUT["path"], "verify_status": "fail", "verification": {"errors": errors, **extra}},
                     default=str))
    sys.exit(0)


def build(state):
    plan = state["plan"]
    errors = []
    src = P.resolve_path(state.get("workbook_path") or plan.get("workbook_path"))
    out = P.resolve_path(plan.get("output_path") or src)
    OUT["path"] = out
    dname, tname = plan.get("data_sheet"), plan.get("target_sheet")
    columns = plan.get("columns") or []
    copy_mode = plan.get("copy_mode", "formula")
    policy = plan.get("failed_lift_policy", "opl")
    if copy_mode not in ("formula", "value"):
        errors.append(f"copy_mode must be formula|value, got {copy_mode!r}")
    if policy not in ("opl", "sum"):
        errors.append(f"failed_lift_policy must be opl|sum, got {policy!r}")
    wb = load_workbook(src)
    if dname not in wb.sheetnames:
        fail([f"data_sheet {dname!r} not in {wb.sheetnames}"])
    ds = wb[dname]
    dhead = [c.value for c in ds[1]]
    dcol = {h: i + 1 for i, h in enumerate(dhead) if h not in (None, "")}
    nrows = 0
    for r in range(2, ds.max_row + 1):
        if any(c.value not in (None, "") for c in ds[r]):
            nrows = r - 1
    if not columns:
        errors.append("plan.columns is empty")
    headers = [c.get("header") for c in columns]
    if len(set(headers)) != len(headers) or None in headers:
        errors.append(f"column headers must be present and unique: {headers}")
    for c in columns:
        k = c.get("kind")
        if k == "copy" and c.get("source") not in dcol:
            errors.append(f"copy column {c.get('header')!r}: source {c.get('source')!r} not a {dname} header {list(dcol)}")
        elif k == "total":
            bad = [s for s in c.get("sources", []) if s not in dcol and s not in headers]
            if bad or not c.get("sources"):
                errors.append(f"total column {c.get('header')!r}: sources {bad or '[]'} not found")
        elif k == "score":
            if c.get("output", "points") not in ("points", "coefficient"):
                errors.append(f"score column {c.get('header')!r}: output must be points|coefficient")
            if (c.get("system") or state.get("score_system")) not in P.SYSTEMS:
                errors.append(f"score column {c.get('header')!r}: system must be one of {P.SYSTEMS}")
        elif k == "rank":
            if c.get("of") not in headers:
                errors.append(f"rank column {c.get('header')!r}: 'of' must name another target column")
        elif k not in ("copy", "total", "score"):
            errors.append(f"column {c.get('header')!r}: kind must be copy|total|score|rank")
    sort = plan.get("sort_by")
    if sort and (not isinstance(sort, dict) or sort.get("header") not in headers):
        errors.append("sort_by must be {\"header\": <target column>, \"descending\": true|false}")
    if errors:
        fail(errors)

    # target sheet
    if tname in wb.sheetnames:
        ts = wb[tname]
        nonempty = any(c.value not in (None, "") for row in ts.iter_rows() for c in row)
        ours = state.get("written_target") == f"{out}::{tname}"
        if nonempty and not (plan.get("clear_target") or ours):
            fail([f"target sheet {tname!r} is not empty; set clear_target true to overwrite it"])
        if nonempty:
            idx = wb.sheetnames.index(tname)
            wb.remove(ts)
            ts = wb.create_sheet(tname, idx)
    else:
        ts = wb.create_sheet(tname)

    dref = P.sheet_ref(dname)
    tletter = {c["header"]: get_column_letter(i + 1) for i, c in enumerate(columns)}

    # where each canonical field lives: target column first, then data column
    def locate(field, override=None):
        if override:
            if override in tletter:
                return ("t", override)
            if override in dcol:
                return ("d", override)
            fail([f"input header {override!r} not found in target or data sheet"])
        thead = [c["header"] for c in columns if c["kind"] in ("copy", "total")]
        h = P.find_header(thead, field)
        if h:
            return ("t", h)
        tot = [c["header"] for c in columns if c["kind"] == "total"]
        if field == "total" and tot:
            return ("t", tot[0])
        h = P.find_header([x for x in dcol], field)
        return ("d", h) if h else None

    def ref(loc, tr, dr):
        side, h = loc
        return f"{tletter[h]}{tr}" if side == "t" else f"{dref}!{get_column_letter(dcol[h])}{dr}"

    def build_rows(order):
        model, formulas = [], {}
        for i, dr in enumerate(order):
            r = i + 2
            drow = {h: ds.cell(row=dr, column=c).value for h, c in dcol.items()}
            mrow = {}

            def val(loc):
                return mrow[loc[1]] if loc[0] == "t" else drow[loc[1]]

            for c in columns:
                h, k = c["header"], c["kind"]
                if k == "copy":
                    v = drow[c["source"]]
                    if copy_mode == "formula":
                        formulas[(r, h)] = f"={dref}!{get_column_letter(dcol[c['source']])}{dr}"
                        v = 0 if v in (None, "") else v
                    mrow[h] = v
                elif k == "total":
                    locs = [("t", x) if x in tletter and x != h else ("d", x) for x in c["sources"]]
                    refs = ",".join(ref(l, r, dr) for l in locs)
                    vals = [val(l) for l in locs]
                    tot = sum(P.num(v) for v in vals if not isinstance(v, str))
                    if policy == "opl":
                        cond = ",".join(f"{ref(l, r, dr)}<0" for l in locs)
                        formulas[(r, h)] = f"=IF(OR({cond}),0,SUM({refs}))"
                        tot = 0.0 if any(isinstance(v, (int, float)) and v < 0 for v in vals) else tot
                    else:
                        formulas[(r, h)] = f"=SUM({refs})"
                    mrow[h] = _rnd(tot, c.get("round"), formulas, r, h)
                elif k == "score":
                    system = c.get("system") or state.get("score_system")
                    ov = c.get("inputs") or {}
                    coef = c.get("output", "points") == "coefficient"
                    need = ["sex", "bodyweight"] + ([] if coef else ["total"]) + (
                        ["equipment", "event"] if system == "ipf_gl" else [])
                    locs = {}
                    for fld in need:
                        loc = locate(fld, ov.get(fld))
                        if loc is None:
                            fail([f"score column {h!r}: no column for {fld}; add a copy/total column or set inputs.{fld}"])
                        locs[fld] = loc
                    refs = {fld: ref(l, r, dr) for fld, l in locs.items()}
                    if coef:
                        refs["total"] = "1"
                    formulas[(r, h)] = "=" + P.score_formula(system, refs)
                    v = P.points(system, val(locs["sex"]), val(locs["bodyweight"]), 1 if coef else val(locs["total"]),
                                 val(locs["equipment"]) if "equipment" in locs else None,
                                 val(locs["event"]) if "event" in locs else None)
                    mrow[h] = _rnd(v, c.get("round"), formulas, r, h)
            model.append(mrow)
        # rank columns: Excel RANK (1 = highest unless ascending), ties share a rank
        last = len(order) + 1
        for c in columns:
            if c["kind"] != "rank":
                continue
            col, asc = tletter[c["of"]], bool(c.get("ascending"))
            for i, m in enumerate(model):
                r = i + 2
                formulas[(r, c["header"])] = f"=RANK({col}{r},${col}$2:${col}${last},{int(asc)})"
                x = P.num(m[c["of"]])
                m[c["header"]] = 1 + sum(1 for o in model if (P.num(o[c["of"]]) < x if asc else P.num(o[c["of"]]) > x))
        return model, formulas

    order = list(range(2, nrows + 2))
    model, formulas = build_rows(order)
    sort = plan.get("sort_by")
    if sort:
        key = sort["header"]
        desc = sort.get("descending", True)
        ranked = sorted(range(nrows), key=lambda i: P.num(model[i][key]), reverse=desc)  # stable on ties
        order = [order[i] for i in ranked]
        model, formulas = build_rows(order)

    # write
    for j, c in enumerate(columns, start=1):
        cell = ts.cell(row=1, column=j, value=c["header"])
        srch = c.get("source") if c["kind"] == "copy" else None
        hdr_src = ds.cell(row=1, column=dcol[srch]) if srch else ds.cell(row=1, column=1)
        if hdr_src.has_style:
            cell.font, cell.fill = copy.copy(hdr_src.font), copy.copy(hdr_src.fill)
            cell.alignment, cell.border = copy.copy(hdr_src.alignment), copy.copy(hdr_src.border)
        letter = get_column_letter(j)
        if srch and ds.column_dimensions[get_column_letter(dcol[srch])].width:
            ts.column_dimensions[letter].width = ds.column_dimensions[get_column_letter(dcol[srch])].width
        else:
            ts.column_dimensions[letter].width = max(12, len(str(c["header"])) + 2)
        for i in range(nrows):
            r = i + 2
            tc = ts.cell(row=r, column=j)
            tc.value = formulas.get((r, c["header"]), model[i][c["header"]])
            if srch:
                sc = ds.cell(row=order[i], column=dcol[srch])
                tc.number_format = sc.number_format
                if sc.has_style:
                    tc.alignment = copy.copy(sc.alignment)
            elif c.get("number_format"):
                tc.number_format = c["number_format"]
            elif c.get("round") is not None:
                tc.number_format = "0." + "0" * int(c["round"]) if int(c["round"]) > 0 else "0"
    if plan.get("freeze_header"):
        ts.freeze_panes = "A2"
    os.makedirs(os.path.dirname(out) or ".", exist_ok=True)
    wb.save(out)
    return out, tname, columns, model, formulas, nrows, dname, order


def _rnd(v, n, formulas, r, h):
    if n is None:
        return v
    formulas[(r, h)] = f"=ROUND({formulas[(r, h)][1:]},{int(n)})"
    return P.excel_round(v, int(n))


def patch_cache(path, tname, columns, model, formulas):
    """Store the reference values as cached results so readers using
    data_only=True see numbers even if LibreOffice is unavailable."""
    with zipfile.ZipFile(path) as z:
        items = {n: z.read(n) for n in z.namelist()}
    wbx = items["xl/workbook.xml"].decode()
    rid = re.search(r'<sheet[^>]*name="%s"[^>]*r:id="([^"]+)"' % re.escape(escape(tname, {'"': "&quot;"})), wbx)
    rels = items["xl/_rels/workbook.xml.rels"].decode()
    tgt = None
    for m in re.finditer(r"<Relationship\b[^>]*>", rels):
        if rid and f'Id="{rid.group(1)}"' in m.group(0):
            tgt = re.search(r'Target="([^"]+)"', m.group(0)).group(1)
    if not tgt:
        return False
    name = tgt.lstrip("/") if tgt.startswith("/") else "xl/" + tgt
    xml = items[name].decode()
    lookup = {}
    for (r, h), _ in formulas.items():
        j = [c["header"] for c in columns].index(h)
        lookup[f"{get_column_letter(j + 1)}{r}"] = model[r - 2][h]

    def sub(m):
        coord, attrs, ftxt = m.group(1), m.group(2), m.group(3)
        if coord not in lookup:
            return m.group(0)
        v = lookup[coord]
        if isinstance(v, str):
            return f'<c r="{coord}"{attrs} t="str"><f>{ftxt}</f><v>{escape(v)}</v></c>'
        if isinstance(v, bool):
            return f'<c r="{coord}"{attrs} t="b"><f>{ftxt}</f><v>{int(v)}</v></c>'
        if isinstance(v, (int, float)) and math.isfinite(v):
            return f'<c r="{coord}"{attrs}><f>{ftxt}</f><v>{repr(v) if isinstance(v, float) else v}</v></c>'
        return m.group(0)

    xml = re.sub(r'<c r="([A-Z]+[0-9]+)"((?:\s+s="\d+")?)><f>(.*?)</f><v\s*/>(?:</v>)?</c>', sub, xml)
    items[name] = xml.encode()
    tmp = path + ".tmp"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        for n, b in items.items():
            z.writestr(n, b)
    os.replace(tmp, path)
    return True


def close(a, b, n):
    if isinstance(a, str) or isinstance(b, str):
        return str(a) == str(b)
    a, b = P.num(a), P.num(b)
    tol = 1.01 * 10 ** -int(n) if n is not None else 1e-6 * max(1.0, abs(b))
    return abs(a - b) <= tol


def main():
    state = json.load(sys.stdin)["currentState"]
    snapshot_src = P.resolve_path(state.get("workbook_path"))
    before = load_workbook(snapshot_src)
    out, tname, columns, model, formulas, nrows, dname, _ = build(state)
    before_data = [[c.value for c in row] for row in before[dname].iter_rows()]
    cached = patch_cache(out, tname, columns, model, formulas)

    recalc_info = {"ran": False}
    if shutil.which("soffice"):
        import recalc
        try:
            res = recalc.recalc(out, int(state.get("plan", {}).get("recalc_timeout", 60)))
            recalc_info = {"ran": "error" not in res, **res}
        except Exception as e:  # noqa: BLE001
            recalc_info = {"ran": False, "error": str(e)}
        if not recalc_info["ran"]:
            patch_cache(out, tname, columns, model, formulas)
    else:
        recalc_info["error"] = "soffice (LibreOffice) not on PATH; formulas kept, values cached from Python reference"

    errors = []
    vals = load_workbook(out, data_only=True)
    forms = load_workbook(out)
    tv, tf = vals[tname], forms[tname]
    mism = []
    for j, c in enumerate(columns, start=1):
        if tv.cell(row=1, column=j).value != c["header"]:
            errors.append(f"header {get_column_letter(j)}1 is {tv.cell(row=1, column=j).value!r}, expected {c['header']!r}")
        for i in range(nrows):
            r, h = i + 2, c["header"]
            got, exp = tv.cell(row=r, column=j).value, model[i][h]
            if (r, h) in formulas and not str(tf.cell(row=r, column=j).value or "").startswith("="):
                errors.append(f"{get_column_letter(j)}{r} lost its formula")
            if got is None or not close(got, exp, c.get("round")):
                mism.append({"cell": f"{tname}!{get_column_letter(j)}{r}", "got": got, "expected": exp})
    for err, d in (recalc_info.get("error_summary") or {}).items():
        errors.append(f"{err} x{d['count']} at {d['locations'][:5]}")
    after_data = [[c.value for c in row] for row in forms[dname].iter_rows()]
    if after_data != before_data:
        errors.append(f"{dname} sheet values changed")
    if mism:
        errors.append(f"{len(mism)} cells differ from the Python reference")
    status = "fail" if errors else ("pass" if recalc_info.get("ran") else "pass_cached")
    score_cols = [c["header"] for c in columns if c["kind"] == "score"]
    summary = {h: {"min": min((m[h] for m in model), default=None), "max": max((m[h] for m in model), default=None)}
               for h in score_cols}
    print(json.dumps({
        "output_path": out,
        "written_target": f"{out}::{tname}",
        "verify_status": status,
        "verification": {
            "errors": errors, "mismatches": mism[:20], "rows_written": nrows, "cache_patched": cached,
            "recalc": recalc_info, "score_ranges": summary,
            "sample_formulas": {c["header"]: formulas.get((2, c["header"])) for c in columns},
            "preview": model[:5],
        },
    }, default=str))


if __name__ == "__main__":
    main()
