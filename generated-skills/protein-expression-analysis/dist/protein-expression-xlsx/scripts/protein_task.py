#!/usr/bin/env python3
"""Protein expression workbook: detect layout, write lookup/statistics/fold-change
formulas into the Task sheet, and verify the recalculated result.

Requires: Python >= 3.8, openpyxl (3.1.x). No other packages.

Usage:
  python3 protein_task.py inspect <workbook.xlsx>
  python3 protein_task.py fill    <in.xlsx> <out.xlsx> [--missing exclude|zero]
  python3 protein_task.py verify  <original.xlsx> <recalculated.xlsx> [--missing exclude|zero]

All subcommands print JSON. Exit code 0 = ok, 1 = problem found (see "errors").
"""
import argparse
import json
import math
import re
import statistics
import sys

try:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter as L
except ImportError:
    print(json.dumps({"status": "error", "errors": ["openpyxl is not installed (pip install openpyxl)"]}))
    sys.exit(1)


def norm(v):
    return re.sub(r"[^a-z0-9]", "", str(v).lower()) if v is not None else ""


def die(msg):
    print(json.dumps({"status": "error", "errors": [msg]}, indent=2))
    sys.exit(1)


def find_sheet(wb, want):
    for ws in wb.worksheets:
        if norm(ws.title) == want:
            return ws
    die(f"no sheet named like '{want}'; sheets: {wb.sheetnames}")


def match_mode(names, headers):
    """How each Task-side name maps to a Data-side header: exact, trim, or contains."""
    hset = set(h for h in headers if isinstance(h, str))
    if all(n in hset for n in names):
        return "exact", []
    stripped = {h.strip() for h in hset}
    if all(str(n).strip() in stripped for n in names):
        return "trim", []
    bad = []
    for n in names:
        hits = [h for h in hset if str(n).strip().lower() in h.lower()]
        if len(hits) != 1:
            bad.append({"name": n, "candidates": sorted(hits)[:5]})
    return ("contains", []) if not bad else ("unmatched", bad)


def layout(wb):
    task, data = find_sheet(wb, "task"), find_sheet(wb, "data")
    # Task: first 'Protein_ID' header row in column A = lookup grid header
    hdr_rows = [r for r in range(1, task.max_row + 1) if norm(task.cell(r, 1).value) == "proteinid"]
    if len(hdr_rows) < 2:
        die(f"expected two 'Protein_ID' header rows in Task column A, found rows {hdr_rows}")
    hr, fc_hr = hdr_rows[0], hdr_rows[1]
    prot_rows = []
    r = hr + 1
    while task.cell(r, 1).value not in (None, ""):
        prot_rows.append(r)
        r += 1
    samp_cols = []
    c = 3
    while task.cell(hr, c).value not in (None, ""):
        samp_cols.append(c)
        c += 1
    if not prot_rows or not samp_cols:
        die("no target proteins (column A under Protein_ID) or no sample names (header row from column C)")
    # group labels: the row above the header that labels sample columns Control/Treated
    groups, grp_row = {}, None
    for gr in range(hr - 1, 0, -1):
        labs = [norm(task.cell(gr, c).value) for c in samp_cols]
        if all(l.startswith("control") or l.startswith("treat") for l in labs):
            grp_row = gr
            groups = {c: ("control" if norm(task.cell(gr, c).value).startswith("control") else "treated") for c in samp_cols}
            break
    if grp_row is None:
        die("could not find a row above the sample header labelling every sample Control/Treated")
    # statistics rows by label in column A
    stat_rows = {}
    for r in range(prot_rows[-1] + 1, fc_hr):
        k = norm(task.cell(r, 1).value)
        for key in ("controlmean", "controlstd", "treatedmean", "treatedstd"):
            if k.startswith(key):
                stat_rows[key] = r
    missing = [k for k in ("controlmean", "controlstd", "treatedmean", "treatedstd") if k not in stat_rows]
    if missing:
        die(f"statistics label rows not found in Task column A: {missing}")
    stat_cols = [2 + i for i in range(len(prot_rows))]  # B.. one column per protein, in target order
    # fold-change table columns by header text
    fc_col = l2_col = None
    for c in range(1, task.max_column + 1):
        k = norm(task.cell(fc_hr, c).value)
        if "log2" in k:
            l2_col = c
        elif "fold" in k:
            fc_col = c
    if not fc_col or not l2_col:
        die(f"fold-change header row {fc_hr} needs a 'Fold Change' and a 'Log2 FC' column")
    fc_rows = [fc_hr + 1 + i for i in range(len(prot_rows))]
    # Data sheet: header row holding Protein_ID
    dhr = dic = None
    for r in range(1, min(data.max_row, 20) + 1):
        for c in range(1, data.max_column + 1):
            if norm(data.cell(r, c).value) == "proteinid":
                dhr, dic = r, c
                break
        if dhr:
            break
    if not dhr:
        die("Data sheet has no 'Protein_ID' header cell in its first 20 rows")
    dlast_r = max(r for r in range(dhr, data.max_row + 1) if data.cell(r, dic).value not in (None, ""))
    dlast_c = max(c for c in range(1, data.max_column + 1) if data.cell(dhr, c).value not in (None, ""))
    return dict(task=task, data=data, hr=hr, grp_row=grp_row, prot_rows=prot_rows, samp_cols=samp_cols,
                groups=groups, stat_rows=stat_rows, stat_cols=stat_cols, fc_hr=fc_hr, fc_col=fc_col,
                l2_col=l2_col, fc_rows=fc_rows, dhr=dhr, dic=dic, dlast_r=dlast_r, dlast_c=dlast_c)


def raw_values(lo):
    """Python-side lookup of the grid straight from Data: {(row, col): float|None}."""
    task, data = lo["task"], lo["data"]
    hdr = {}
    for c in range(1, lo["dlast_c"] + 1):
        v = data.cell(lo["dhr"], c).value
        if isinstance(v, str):
            hdr[v.strip().lower()] = c
    ids = {}
    for r in range(lo["dhr"] + 1, lo["dlast_r"] + 1):
        v = data.cell(r, lo["dic"]).value
        if v is not None:
            ids[str(v).strip().lower()] = r

    def find(d, name):
        k = str(name).strip().lower()
        if k in d:
            return d[k]
        hits = [v for h, v in d.items() if k in h]
        return hits[0] if len(hits) == 1 else None

    out = {}
    for r in lo["prot_rows"]:
        dr = find(ids, task.cell(r, 1).value)
        for c in lo["samp_cols"]:
            dc = find(hdr, task.cell(lo["hr"], c).value)
            v = data.cell(dr, dc).value if dr and dc else None
            out[(r, c)] = float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None
    return out


def runs(cols, row):
    """Contiguous column runs -> 'C11:G11,I11' style argument list."""
    parts, start = [], None
    for i, c in enumerate(cols):
        if start is None:
            start = c
        if i == len(cols) - 1 or cols[i + 1] != c + 1:
            parts.append(f"{L(start)}{row}" if start == c else f"{L(start)}{row}:{L(c)}{row}")
            start = None
    return ",".join(parts)


def cmd_inspect(path):
    wb = load_workbook(path)
    lo = layout(wb)
    task, data = lo["task"], lo["data"]
    names = [task.cell(lo["hr"], c).value for c in lo["samp_cols"]]
    headers = [data.cell(lo["dhr"], c).value for c in range(1, lo["dlast_c"] + 1)]
    pids = [task.cell(r, 1).value for r in lo["prot_rows"]]
    ids = [data.cell(r, lo["dic"]).value for r in range(lo["dhr"] + 1, lo["dlast_r"] + 1)]
    smode, sbad = match_mode(names, headers)
    pmode, pbad = match_mode(pids, ids)
    vals = raw_values(lo)
    per = []
    for i, r in enumerate(lo["prot_rows"]):
        n = {g: sum(1 for c in lo["samp_cols"] if lo["groups"][c] == g and vals[(r, c)] is not None) for g in ("control", "treated")}
        per.append({"row": r, "protein": task.cell(r, 1).value, "stats_col": L(lo["stat_cols"][i]),
                    "n_control": n["control"], "n_treated": n["treated"],
                    "missing_cells": [f"{L(c)}{r}" for c in lo["samp_cols"] if vals[(r, c)] is None]})
    pr, sc, fr = lo["prot_rows"], lo["samp_cols"], lo["fc_rows"]
    sr = lo["stat_rows"]
    # cells we will write must be empty now
    targets = [(r, c) for r in pr for c in sc] + [(sr[k], c) for k in sr for c in lo["stat_cols"]] + \
              [(r, c) for r in fr for c in (lo["fc_col"], lo["l2_col"])]
    occupied = [f"{L(c)}{r}" for r, c in targets if task.cell(r, c).value not in (None, "")]
    errors = []
    if smode == "unmatched":
        errors.append({"unmatched_samples": sbad})
    if pmode == "unmatched":
        errors.append({"unmatched_proteins": pbad})
    res = {
        "status": "ok" if not errors else "error",
        "errors": errors,
        "lookup_range": f"{L(sc[0])}{pr[0]}:{L(sc[-1])}{pr[-1]}",
        "sample_header_row": lo["hr"], "group_row": lo["grp_row"],
        "control_columns": [L(c) for c in sc if lo["groups"][c] == "control"],
        "treated_columns": [L(c) for c in sc if lo["groups"][c] == "treated"],
        "stats_range": f"{L(lo['stat_cols'][0])}{min(sr.values())}:{L(lo['stat_cols'][-1])}{max(sr.values())}",
        "stats_rows": sr,
        "fold_change_cells": f"{L(lo['fc_col'])}{fr[0]}:{L(lo['fc_col'])}{fr[-1]}",
        "log2fc_cells": f"{L(lo['l2_col'])}{fr[0]}:{L(lo['l2_col'])}{fr[-1]}",
        "data_table": f"Data!{L(1)}{lo['dhr']}:{L(lo['dlast_c'])}{lo['dlast_r']}",
        "sample_match": smode, "protein_match": pmode,
        "already_filled_targets": occupied,
        "total_missing_lookup_cells": sum(len(p["missing_cells"]) for p in per),
        "groups_with_fewer_than_2_values": [f"{p['protein']}:{g}" for p in per for g in ("control", "treated") if p[f"n_{g}"] < 2],
        "proteins": per,
    }
    print(json.dumps(res, indent=2))
    return 0 if not errors else 1


def cmd_fill(src, dst, missing):
    wb = load_workbook(src)
    lo = layout(wb)
    task, data = lo["task"], lo["data"]
    names = [task.cell(lo["hr"], c).value for c in lo["samp_cols"]]
    headers = [data.cell(lo["dhr"], c).value for c in range(1, lo["dlast_c"] + 1)]
    pids = [task.cell(r, 1).value for r in lo["prot_rows"]]
    ids = [data.cell(r, lo["dic"]).value for r in range(lo["dhr"] + 1, lo["dlast_r"] + 1)]
    smode, sbad = match_mode(names, headers)
    pmode, pbad = match_mode(pids, ids)
    if "unmatched" in (smode, pmode):
        die(f"names not found in Data: samples={sbad} proteins={pbad}")
    ds = f"'{data.title}'" if not re.fullmatch(r"\w+", data.title) else data.title
    tbl = f"{ds}!${L(1)}${lo['dhr']}:${L(lo['dlast_c'])}${lo['dlast_r']}"
    idcol = f"{ds}!${L(lo['dic'])}${lo['dhr']}:${L(lo['dic'])}${lo['dlast_r']}"
    hdrrow = f"{ds}!${L(1)}${lo['dhr']}:${L(lo['dlast_c'])}${lo['dhr']}"

    def key(ref, mode):
        return {"exact": ref, "trim": f"TRIM({ref})", "contains": f'"*"&TRIM({ref})&"*"'}[mode]

    hr = lo["hr"]
    for r in lo["prot_rows"]:
        for c in lo["samp_cols"]:
            idx = (f"INDEX({tbl},MATCH({key(f'$A{r}', pmode)},{idcol},0),"
                   f"MATCH({key(f'{L(c)}${hr}', smode)},{hdrrow},0))")
            # a blank Data cell must stay blank (missing), not become 0, unless policy says zero
            task.cell(r, c).value = f"=N({idx})" if missing == "zero" else f'=IF({idx}="","",{idx})'
    sr = lo["stat_rows"]
    ctrl = [c for c in lo["samp_cols"] if lo["groups"][c] == "control"]
    trt = [c for c in lo["samp_cols"] if lo["groups"][c] == "treated"]
    for i, r in enumerate(lo["prot_rows"]):
        sc = lo["stat_cols"][i]
        for cols, mk, sk in ((ctrl, "controlmean", "controlstd"), (trt, "treatedmean", "treatedstd")):
            rng = runs(cols, r)
            task.cell(sr[mk], sc).value = f'=IF(COUNT({rng})>=1,AVERAGE({rng}),"")'
            task.cell(sr[sk], sc).value = f'=IF(COUNT({rng})>=2,STDEV({rng}),"")'
        s = L(sc)
        fr = lo["fc_rows"][i]
        cm, tm = f"{s}{sr['controlmean']}", f"{s}{sr['treatedmean']}"
        l2 = f"{L(lo['l2_col'])}{fr}"
        task.cell(fr, lo["l2_col"]).value = f'=IF(OR({cm}="",{tm}=""),"",{tm}-{cm})'
        task.cell(fr, lo["fc_col"]).value = f'=IF({l2}="","",POWER(2,{l2}))'
    wb.save(dst)
    print(json.dumps({"status": "ok", "written": dst, "missing_policy": missing,
                      "sample_match": smode, "protein_match": pmode,
                      "example_formulas": {
                          "lookup": task.cell(lo["prot_rows"][0], lo["samp_cols"][0]).value,
                          "control_mean": task.cell(sr["controlmean"], lo["stat_cols"][0]).value,
                          "control_std": task.cell(sr["controlstd"], lo["stat_cols"][0]).value,
                          "log2fc": task.cell(lo["fc_rows"][0], lo["l2_col"]).value,
                          "fold_change": task.cell(lo["fc_rows"][0], lo["fc_col"]).value},
                      "next": f"python3 scripts/recalc.py {dst} 60"}, indent=2))
    return 0


def color_sig(col):
    if col is None:
        return None
    return col.rgb if isinstance(col.rgb, str) else f"theme{col.theme}"


def style_sig(c):
    f, fo = c.fill, c.font
    return (f.fill_type, color_sig(f.fgColor) if f.fill_type else None, fo.name, fo.sz, fo.b, fo.i,
            color_sig(fo.color), c.number_format)


def cmd_verify(orig, out, missing, tol=1e-6):
    wo, wf, wv = load_workbook(orig), load_workbook(out), load_workbook(out, data_only=True)
    lo = layout(wo)
    errors, warnings = [], []
    if wo.sheetnames != wf.sheetnames:
        errors.append(f"sheet names changed: {wo.sheetnames} -> {wf.sheetnames}")
    if wf.vba_archive is not None or out.lower().endswith(".xlsm"):
        errors.append("workbook contains macros/VBA")
    t_f, t_v = wf[lo["task"].title], wv[lo["task"].title]
    vals = raw_values(lo)
    sr = lo["stat_rows"]
    expect = {}
    for (r, c), v in vals.items():
        expect[(r, c)] = (0.0 if v is None else v) if missing == "zero" else v
    ctrl = [c for c in lo["samp_cols"] if lo["groups"][c] == "control"]
    trt = [c for c in lo["samp_cols"] if lo["groups"][c] == "treated"]
    for i, r in enumerate(lo["prot_rows"]):
        sc = lo["stat_cols"][i]
        m = {}
        for cols, mk, sk in ((ctrl, "controlmean", "controlstd"), (trt, "treatedmean", "treatedstd")):
            xs = [expect[(r, c)] for c in cols if expect[(r, c)] is not None]
            m[mk] = statistics.mean(xs) if xs else None
            expect[(sr[mk], sc)] = m[mk]
            expect[(sr[sk], sc)] = statistics.stdev(xs) if len(xs) >= 2 else None
        fr = lo["fc_rows"][i]
        l2 = None if None in (m["controlmean"], m["treatedmean"]) else m["treatedmean"] - m["controlmean"]
        expect[(fr, lo["l2_col"])] = l2
        expect[(fr, lo["fc_col"])] = None if l2 is None else 2 ** l2
    nformula, uncached = 0, []
    for (r, c), e in expect.items():
        ref = f"{L(c)}{r}"
        f = t_f.cell(r, c).value
        if not (isinstance(f, str) and f.startswith("=")):
            errors.append(f"{ref}: not a formula ({f!r})")
            continue
        nformula += 1
        v = t_v.cell(r, c).value
        if isinstance(v, str) and v.startswith("#"):
            errors.append(f"{ref}: formula error {v}")
        elif e is None:
            if v not in (None, ""):
                errors.append(f"{ref}: expected blank (missing data), got {v!r}")
        elif v is None or v == "":
            uncached.append(ref)
        elif not isinstance(v, (int, float)) or not math.isclose(v, e, rel_tol=tol, abs_tol=tol):
            errors.append(f"{ref}: value {v!r} != expected {e:.6g}")
    # nothing else changed: values and styles of every other cell, both sheets
    target = set(expect)
    for ws_o in wo.worksheets:
        ws_f = wf[ws_o.title]
        for row in ws_o.iter_rows(min_row=1, max_row=max(ws_o.max_row, ws_f.max_row),
                                  max_col=max(ws_o.max_column, ws_f.max_column)):
            for co in row:
                cf = ws_f.cell(co.row, co.column)
                if ws_o.title == lo["task"].title and (co.row, co.column) in target:
                    if style_sig(co) != style_sig(cf):
                        warnings.append(f"{ws_o.title}!{co.coordinate}: style changed {style_sig(co)} -> {style_sig(cf)}")
                    continue
                vo, vf = co.value, cf.value
                same = vo == vf or (isinstance(vo, float) and isinstance(vf, (int, float)) and math.isclose(vo, vf, rel_tol=1e-9))
                if not same:
                    errors.append(f"{ws_o.title}!{co.coordinate}: non-target cell changed {vo!r} -> {vf!r}")
                if style_sig(co) != style_sig(cf):
                    warnings.append(f"{ws_o.title}!{co.coordinate}: style changed {style_sig(co)} -> {style_sig(cf)}")
        if set(map(str, ws_o.merged_cells.ranges)) != set(map(str, ws_f.merged_cells.ranges)):
            errors.append(f"{ws_o.title}: merged cells changed")
    if uncached:
        errors.append(f"{len(uncached)} target cells have no cached value ({', '.join(uncached[:5])}...): "
                      "the file was not recalculated after the last openpyxl save; run scripts/recalc.py on it")
    res = {"status": "ok" if not errors else "errors_found", "target_formulas": nformula,
           "errors": errors[:40], "error_count": len(errors),
           "style_warnings": warnings[:20], "style_warning_count": len(warnings)}
    print(json.dumps(res, indent=2))
    return 0 if not errors else 1


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("inspect"); p.add_argument("xlsx")
    p = sub.add_parser("fill"); p.add_argument("src"); p.add_argument("dst")
    p.add_argument("--missing", choices=["exclude", "zero"], default="exclude")
    p = sub.add_parser("verify"); p.add_argument("orig"); p.add_argument("out")
    p.add_argument("--missing", choices=["exclude", "zero"], default="exclude")
    a = ap.parse_args()
    if a.cmd == "inspect":
        sys.exit(cmd_inspect(a.xlsx))
    if a.cmd == "fill":
        if a.src == a.dst:
            die("write to a new file first (fill reads the template; keep the original for verify)")
        sys.exit(cmd_fill(a.src, a.dst, a.missing))
    sys.exit(cmd_verify(a.orig, a.out, a.missing))


if __name__ == "__main__":
    main()
