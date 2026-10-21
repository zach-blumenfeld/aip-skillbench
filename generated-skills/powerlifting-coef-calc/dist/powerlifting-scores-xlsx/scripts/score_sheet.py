#!/usr/bin/env python3
"""
Build and verify a powerlifting score sheet (Dots, Wilks, or IPF GOODLIFT) as live
Excel formulas inside an existing workbook.

Requires: python>=3.8, openpyxl (tested with 3.1.5). No other packages.

  build  <xlsx> [--data-sheet Data] [--out-sheet Dots] [--score dots|wilks|goodlift]
                [--digits 3] [--round-coefficient] [--output PATH]
         Copies the columns the score needs (Name + inputs, in their source order and
         with their source header names) from the data sheet into the output sheet,
         then appends a "TotalKg" formula column and a score formula column.
         Saves in place unless --output is given. Prints a JSON plan.

  verify <xlsx> [same options as build]
         Run AFTER recalc.py. Re-reads the workbook, checks headers, copied values,
         that formulas are present, that every formula has a cached value with no
         Excel error, and that each cached value equals an independent Python
         reference (the OpenPowerlifting implementations). Exit 1 on any failure.
"""

import argparse
import json
import math
import sys

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

LIFTS = ["Best3SquatKg", "Best3BenchKg", "Best3DeadliftKg"]
SCORE_INPUTS = {
    "dots": ["Name", "Sex", "BodyweightKg"] + LIFTS,
    "wilks": ["Name", "Sex", "BodyweightKg"] + LIFTS,
    "goodlift": ["Name", "Sex", "Event", "Equipment", "BodyweightKg"] + LIFTS,
}
SCORE_HEADER = {"dots": "Dots", "wilks": "Wilks", "goodlift": "Goodlift"}

# --- coefficients (OpenPowerlifting, crates/coefficients) ---------------------------
DOTS = {  # poly4: A*x^4 + B*x^3 + C*x^2 + D*x + E ; (coeffs, clamp_lo, clamp_hi)
    "M": ((-0.0000010930, 0.0007391293, -0.1918759221, 24.0900756, -307.75076), 40.0, 210.0),
    "F": ((-0.0000010706, 0.0005158568, -0.1126655495, 13.6175032, -57.96288), 40.0, 150.0),
}
WILKS = {  # A + B*x + C*x^2 + D*x^3 + E*x^4 + F*x^5
    "M": ((-216.0475144, 16.2606339, -0.002388645, -0.00113732, 7.01863e-06, -1.291e-08), 40.0, 201.9),
    "F": ((594.31747775582, -27.23842536447, 0.82112226871, -0.00930733913, 0.00004731582, -0.00000009054), 26.51, 154.53),
}
GOODLIFT = {  # (event, sex, equipment) -> (A, B, C); points = total * max(0, 100/(A - B*e^(-C*bw)))
    ("SBD", "M", "Raw"): (1199.72839, 1025.18162, 0.009210),
    ("SBD", "M", "Single"): (1236.25115, 1449.21864, 0.01644),
    ("SBD", "F", "Raw"): (610.32796, 1045.59282, 0.03048),
    ("SBD", "F", "Single"): (758.63878, 949.31382, 0.02435),
    ("B", "M", "Raw"): (320.98041, 281.40258, 0.01008),
    ("B", "M", "Single"): (381.22073, 733.79378, 0.02398),
    ("B", "F", "Raw"): (142.40398, 442.52671, 0.04724),
    ("B", "F", "Single"): (221.82209, 357.00377, 0.02937),
}
RAW_EQ = ("Raw", "Wraps", "Straps")  # everything else (Single-ply, Multi-ply, Unlimited) -> Single


def sex_key(sex):
    return "F" if str(sex).strip() == "F" else "M"  # M and Mx use the men's formula


def num(v):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


# --- Python reference (used only by verify) -------------------------------------------
def ref_total(lifts, digits):
    vals = [x for x in (num(v) for v in lifts) if x is not None]
    if any(x < 0 for x in vals):  # negative Best3 = failed lift -> no total
        return 0.0
    return round(sum(vals), digits)


def ref_coef(score, sex, bw, event=None, equipment=None):
    s = sex_key(sex)
    if score == "dots":
        (a, b, c, d, e), lo, hi = DOTS[s]
        x = min(max(bw, lo), hi)
        return 500.0 / (a * x**4 + b * x**3 + c * x**2 + d * x + e)
    if score == "wilks":
        k, lo, hi = WILKS[s]
        x = min(max(bw, lo), hi)
        return 500.0 / sum(k[i] * x**i for i in range(6))
    eq = "Raw" if str(equipment).strip() in RAW_EQ else "Single"
    p = GOODLIFT.get((str(event).strip(), s, eq))
    if p is None or bw < 35:
        return 0.0
    den = p[0] - p[1] * math.exp(-p[2] * bw)
    return 0.0 if den == 0 else max(0.0, 100.0 / den)


def ref_score(score, sex, bw, total, digits, round_coef, event=None, equipment=None):
    if bw is None or bw <= 0 or not total or total <= 0:
        return 0.0
    c = ref_coef(score, sex, bw, event, equipment)
    if round_coef:
        c = round(c, digits)
    return round(c * total, digits)


# --- Excel formula generation -----------------------------------------------------------
def f(v):
    return repr(float(v)).replace("e", "E")


def poly_formula(coeffs_low_to_high, x):
    terms = []
    for i, k in enumerate(coeffs_low_to_high):
        t = f(k) if i == 0 else f"{f(k)}*{x}" if i == 1 else f"{f(k)}*{x}^{i}"
        terms.append(t)
    return "+".join(terms).replace("+-", "-")


def coef_formula(score, col, r):
    sx, bw = f"{col['Sex']}{r}", f"{col['BodyweightKg']}{r}"
    if score in ("dots", "wilks"):
        table = DOTS if score == "dots" else WILKS
        parts = {}
        for s in ("M", "F"):
            k, lo, hi = table[s]
            if score == "dots":
                k = tuple(reversed(k))  # store low->high
            x = f"MIN(MAX({bw},{f(lo)}),{f(hi)})"
            parts[s] = f"500/({poly_formula(k, x)})"
        return f'IF({sx}="F",{parts["F"]},{parts["M"]})'
    ev, eq = f"{col['Event']}{r}", f"{col['Equipment']}{r}"
    raw = f'OR({eq}="Raw",{eq}="Wraps",{eq}="Straps")'

    def pick(i):
        expr = "0"
        for (e, s, q), p in reversed(list(GOODLIFT.items())):
            sc = f'{sx}="F"' if s == "F" else f'{sx}<>"F"'
            qc = raw if q == "Raw" else f"NOT({raw})"
            cond = f'AND({ev}="{e}",{sc},{qc})'
            expr = f"IF({cond},{f(p[i])},{expr})"
        return expr

    a, b, c = pick(0), pick(1), pick(2)
    den = f"(({a})-({b})*EXP(-({c})*{bw}))"
    return f"IF(OR(({a})=0,{bw}<35),0,IF({den}=0,0,MAX(0,100/{den})))"


def plan(ws, score):
    header = [c.value for c in ws[1]]
    idx = {str(h).strip(): i for i, h in enumerate(header) if h is not None}
    need = SCORE_INPUTS[score]
    missing = [n for n in need if n not in idx and n not in LIFTS]
    if missing:
        sys.exit(f"ERROR: data sheet lacks required column(s) {missing}; headers are {header}")
    present = [n for n in need if n in idx]
    if not any(n in idx for n in LIFTS):
        sys.exit(f"ERROR: no lift columns ({LIFTS}) in data sheet; headers are {header}")
    return sorted(present, key=lambda n: idx[n]), idx


def last_row(ws, name_col):
    r = ws.max_row
    while r > 1 and ws.cell(r, name_col).value in (None, ""):
        r -= 1
    return r


def layout(cols, score):
    col = {n: get_column_letter(i + 1) for i, n in enumerate(cols)}
    col["TotalKg"] = get_column_letter(len(cols) + 1)
    col[SCORE_HEADER[score]] = get_column_letter(len(cols) + 2)
    return col


def formulas(score, col, cols, r, digits, round_coef):
    lifts = [f"{col[n]}{r}" for n in cols if n in LIFTS]
    rng = ",".join(lifts)
    total = f"=IF(MIN({rng})<0,0,ROUND(SUM({rng}),{digits}))"
    coef = coef_formula(score, col, r)
    if round_coef:
        coef = f"ROUND({coef},{digits})"
    g, bw = f"{col['TotalKg']}{r}", f"{col['BodyweightKg']}{r}"
    sc = f'=IF(OR({bw}="",{bw}<=0,{g}<=0),0,ROUND({g}*{coef},{digits}))'
    return total, sc


def build(a):
    wb = load_workbook(a.xlsx)
    if a.data_sheet not in wb.sheetnames:
        sys.exit(f"ERROR: no sheet {a.data_sheet!r}; sheets are {wb.sheetnames}")
    src = wb[a.data_sheet]
    cols, idx = plan(src, a.score)
    n_last = last_row(src, idx["Name"] + 1)
    if a.out_sheet in wb.sheetnames:
        old = wb[a.out_sheet]
        pos = wb.sheetnames.index(a.out_sheet)
        wb.remove(old)
        dst = wb.create_sheet(a.out_sheet, pos)
    else:
        dst = wb.create_sheet(a.out_sheet)
    col = layout(cols, a.score)
    hdr = cols + ["TotalKg", SCORE_HEADER[a.score]]
    for j, h in enumerate(hdr, 1):
        dst.cell(1, j, h)
    negatives = []
    for r in range(2, n_last + 1):
        for j, n in enumerate(cols, 1):
            s = src.cell(r, idx[n] + 1)
            d = dst.cell(r, j, s.value)
            d.number_format = s.number_format
            if n in LIFTS and num(s.value) is not None and num(s.value) < 0:
                negatives.append(f"{a.data_sheet}!{s.coordinate}")
        t, sc = formulas(a.score, col, cols, r, a.digits, a.round_coefficient)
        dst.cell(r, len(cols) + 1, t).number_format = "0.000" if a.digits == 3 else "0." + "0" * a.digits
        dst.cell(r, len(cols) + 2, sc).number_format = dst.cell(r, len(cols) + 1).number_format
    for j, h in enumerate(hdr, 1):
        dst.column_dimensions[get_column_letter(j)].width = max(12, len(str(h)) + 2)
    out = a.output or a.xlsx
    wb.save(out)
    print(json.dumps({
        "saved": out, "sheet": a.out_sheet, "score": a.score, "rows": n_last - 1,
        "copied_columns": cols, "columns": {h: col[h] for h in hdr},
        "dropped_source_TotalKg": "TotalKg" in idx,
        "failed_lift_cells_total_set_0": negatives,
        "sample_formulas": formulas(a.score, col, cols, 2, a.digits, a.round_coefficient) if n_last >= 2 else None,
        "next": f"python scripts/recalc.py {out} 60  then  python scripts/score_sheet.py verify {out}",
    }, indent=2))


def verify(a):
    wbf = load_workbook(a.xlsx)
    wbv = load_workbook(a.xlsx, data_only=True)
    src, dst, dv = wbv[a.data_sheet], wbf[a.out_sheet], wbv[a.out_sheet]
    cols, idx = plan(src, a.score)
    col = layout(cols, a.score)
    hdr = cols + ["TotalKg", SCORE_HEADER[a.score]]
    probs = []
    got = [c.value for c in dst[1]][: len(hdr)]
    if got != hdr or (dst.max_column > len(hdr) and any(c.value is not None for c in dst[1][len(hdr):])):
        probs.append(f"header {[c.value for c in dst[1]]} != expected {hdr}")
    n_last = last_row(src, idx["Name"] + 1)
    rows = []
    for r in range(2, n_last + 1):
        v = {n: src.cell(r, idx[n] + 1).value for n in cols}
        for j, n in enumerate(cols, 1):
            if dv.cell(r, j).value != v[n]:
                probs.append(f"{a.out_sheet}!{get_column_letter(j)}{r} = {dv.cell(r, j).value!r} but data has {v[n]!r}")
        tot_ref = ref_total([v[n] for n in cols if n in LIFTS], a.digits)
        sc_ref = ref_score(a.score, v["Sex"], num(v["BodyweightKg"]), tot_ref, a.digits, a.round_coefficient,
                           v.get("Event"), v.get("Equipment"))
        for name, ref in (("TotalKg", tot_ref), (SCORE_HEADER[a.score], sc_ref)):
            c = f"{col[name]}{r}"
            fv, cv = dst[c].value, dv[c].value
            if not (isinstance(fv, str) and fv.startswith("=")):
                probs.append(f"{c} is not a formula: {fv!r}")
            if cv is None:
                probs.append(f"{c} has no cached value (recalc not run or failed)")
            elif isinstance(cv, str):
                probs.append(f"{c} evaluates to {cv}")
            elif abs(float(cv) - ref) > 1.5 * 10 ** -a.digits:
                probs.append(f"{c} = {cv} but reference = {ref}")
        rows.append({"row": r, "Name": v["Name"], "TotalKg": tot_ref, SCORE_HEADER[a.score]: sc_ref,
                     "cached": [dv[f"{col['TotalKg']}{r}"].value, dv[f"{col[SCORE_HEADER[a.score]]}{r}"].value]})
    print(json.dumps({"status": "ok" if not probs else "failed", "rows_checked": len(rows),
                      "problems": probs[:50], "n_problems": len(probs), "rows": rows}, indent=2, default=str))
    sys.exit(0 if not probs else 1)


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("cmd", choices=["build", "verify"])
    p.add_argument("xlsx")
    p.add_argument("--data-sheet", default="Data")
    p.add_argument("--out-sheet", default=None, help="default: Dots / Wilks / Goodlift per --score")
    p.add_argument("--score", choices=list(SCORE_INPUTS), default="dots")
    p.add_argument("--digits", type=int, default=3)
    p.add_argument("--round-coefficient", action="store_true",
                   help="also round the coefficient itself to --digits before multiplying")
    p.add_argument("--output")
    a = p.parse_args()
    a.out_sheet = a.out_sheet or SCORE_HEADER[a.score]
    build(a) if a.cmd == "build" else verify(a)


if __name__ == "__main__":
    main()
